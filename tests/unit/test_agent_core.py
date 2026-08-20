import sys
import json
import threading
import time
import http.client
from pathlib import Path

import pytest


ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "src"
SCRIPTS = ROOT / "scripts"
for path in (SRC, SCRIPTS):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from hara_agent.domains import (  # noqa: E402
    AVPDomainPolicy, DomainRegistry, default_domain_registry, load_domain_profile,
)
from hara_agent.compatibility import V13ScoringAdapter  # noqa: E402
from hara_agent.infrastructure.llm import (  # noqa: E402
    LLMJSONContractError, LLMOutputLimitError, LLMSchemaContractError,
    LLMTimeoutError, LLMTransportError,
    LLMResponse,
    OpenAICompatibleClient,
)
from hara_agent.infrastructure.llm import LLMRequest, create_llm_client  # noqa: E402
from hara_agent.config import LLMConfig  # noqa: E402
from hara_agent.config import RunConfig  # noqa: E402
from hara_agent.application import HARAApplication  # noqa: E402
from hara_agent.models import (  # noqa: E402
    EvidenceValue, FunctionDefinition, GuidewordAssessment, MalfunctionCandidate,
    ReviewStatus, SafetyGoal, ScenarioCandidate, SourceRef,
)
from hara_agent.services.analysis import (  # noqa: E402
    DomainScoringService,
    AVPScenarioCandidateService,
    FTTIService,
    RiskAggregationService,
    SafetyGoalCatalogService,
    TemplateASILService,
)
from hara_agent.workflow import (  # noqa: E402
    CheckpointRepository, HARAState, SemanticWorkflowAgents, SemanticWorkflowInputs,
    RiskWorkflowServices, WorkflowGraph, WorkflowStage, build_hara_agent_graph,
    build_semantic_frontend_graph,
)
from hara_agent.workflow.nodes import (  # noqa: E402
    aggregate_safety_goals, assess_guidewords, import_v13_scoring, pass_quality_gate,
    score_structured_scenarios,
)
from hara_agent.workflow.nodes import extract_functions, extract_item_artifacts  # noqa: E402
from hara_agent.workflow.nodes import read_item_document  # noqa: E402
from hara_agent.workflow.nodes.parallel import ordered_parallel_map  # noqa: E402
from hara_agent.services.semantic import FunctionExtractionAgent  # noqa: E402
from hara_agent.services.semantic import ItemDefinitionExtractionAgent  # noqa: E402
from hara_agent.services.semantic import ItemArtifactExtractionAgent  # noqa: E402
from hara_agent.services.semantic import ItemEvidenceRouter  # noqa: E402
from hara_agent.services.semantic import ItemSupplementAgent  # noqa: E402
from hara_agent.services.semantic import GuidewordApplicabilityAgent  # noqa: E402
from hara_agent.services.semantic import MalfunctionHazardAgent  # noqa: E402
from hara_agent.services.semantic import ScenarioFeasibilityAgent  # noqa: E402
from hara_agent.services.semantic.parsing import parse_confidence  # noqa: E402
from hara_agent.services.semantic.scenario_batching import (  # noqa: E402
    ScenarioAdaptiveBatchError, ScenarioBatchSizeError, ScenarioFactConsistencyError,
    ScenarioSchemaContractError,
    build_scenario_batches, build_scenario_user_prompt,
    estimate_scenario_prompt_chars,
)
from hara_agent.services.semantic.scenario_contract import (  # noqa: E402
    RISK_DIMENSION_VALUES, SCENARIO_CONTRACT_VERSION,
)
from hara_agent.services.semantic.scenario_evidence import (  # noqa: E402
    FactRegistry, ScenarioEvidenceContractError, build_fact_registry,
    validate_evidence_contract,
)
from hara_agent.services.reporting import HARAExcelRenderer  # noqa: E402
from hara_agent.services.validation import DownstreamPreflightService  # noqa: E402
from hara_agent.services.extraction import TemplateInputReader  # noqa: E402
from hara_agent.services.extraction import ValidatedArtifactCache  # noqa: E402
from scoring_sg_engine import (  # noqa: E402
    AVPSafetyGoalCatalog,
    ControllabilityReasoningEngine,
    SeverityReasoningEngine,
)
from logic_checkers.scoring_checker import ScoringChecker  # noqa: E402
import scenario_smoke  # noqa: E402


TEMPLATE = ROOT / "references" / "HARA_Template_AI_20260327.xlsx"


def _scoring_service(policy):
    standards = TemplateInputReader().read(TEMPLATE).scoring_standards
    return DomainScoringService(policy, standards)


def _apply_driver_context(facts, profile, context_id="driver_outside_remote_monitoring"):
    contexts = profile.section("scenario_policy")["driver_context_candidates"]
    context = next(item for item in contexts if item["context_id"] == context_id)
    facts.update(context)
    return facts


def test_finalized_value_requires_traceability():
    with pytest.raises(ValueError, match="source or approved rule version"):
        EvidenceValue(value="E3", status=ReviewStatus.FINALIZED)

    value = EvidenceValue(
        value="E3",
        status=ReviewStatus.FINALIZED,
        sources=[SourceRef("project_input", "SCENARIO-001", "Scenarios_Library!A12")],
    )
    assert value.to_dict()["status"] == "FINALIZED"


@pytest.mark.parametrize(
    ("value", "expected"),
    [(0.82, 0.82), (1, 1.0), ("0.82", 0.82), (None, 0.0)],
)
def test_parse_confidence_accepts_shared_numeric_contract(value, expected):
    assert parse_confidence(value) == expected


@pytest.mark.parametrize("value", ["high", "medium", "low", True, -0.1, 1.01, {}, []])
def test_parse_confidence_rejects_qualitative_boolean_and_out_of_range(value):
    with pytest.raises(ValueError, match=r"expected=number\[0\.0,1\.0\]"):
        parse_confidence(value, field_name="test.confidence")


def test_semantic_agents_reject_qualitative_confidence_with_context():
    function = FunctionDefinition("FUN-1", "制动控制", "制动请求")
    with pytest.raises(ValueError, match="Function confidence.*actual='high'"):
        FunctionExtractionAgent._parse({
            "function_id": "FUN-1", "name": "制动控制", "output": "制动请求",
            "confidence": "high",
        }, "ItemDef.docx")
    with pytest.raises(ValueError, match="function=FUN-1 guideword=loss.*actual='high'"):
        GuidewordApplicabilityAgent._parse(function, {
            "guideword": "loss", "applicable": True, "rationale": "请求丢失",
            "confidence": "high",
        })
    with pytest.raises(ValueError, match="function=FUN-1 malfunction=MF-1.*actual='high'"):
        MalfunctionHazardAgent._parse(function, {
            "malfunction_id": "MF-1", "guideword": "loss", "description": "请求丢失",
            "functional_effect": "无法减速", "vehicle_level_hazard": "车辆继续运动",
            "causal_chain": ["请求丢失", "车辆继续运动"], "confidence": "high",
        })
    malfunction = MalfunctionCandidate(
        "MF-1", "FUN-1", "loss", "请求丢失", "无法减速",
        "车辆继续运动", ["请求丢失", "车辆继续运动"],
    )
    with pytest.raises(ValueError, match="malfunction=MF-1 scenario=SCN-1.*actual='high'"):
        item = _v9_fixture({
            "scenario_id": "SCN-1", "physically_feasible": False,
            "functionally_relevant": False, "causally_relevant": False,
            "risk_dimensions_changed": [], "rationale": "不相关", "confidence": "high",
        })
        ScenarioFeasibilityAgent._parse(
            malfunction, item,
            scenario=ScenarioCandidate("SCN-1", "Parking", "fixture", "fixture", {}),
        )


def test_pending_review_blocks_publication():
    state = HARAState(run_id="run-001")
    assert state.can_publish
    state.pending_reviews.append({"field": "exposure", "reason": "missing IFFT"})
    assert not state.can_publish


def test_migration_profile_is_not_silently_treated_as_approved():
    with pytest.raises(ValueError, match="未批准"):
        load_domain_profile("avp")

    profile = load_domain_profile("avp", require_approved=False)
    assert profile.name == "avp"
    assert profile.approval_status == "migration_baseline"
    assert "asil_table" not in profile.data
    assert "asil_matrix" not in profile.data
    assert "scoring_policy" not in profile.data
    assert "risk_mapping_policy" in profile.data
    assert "driver_state" not in profile.section("scenario_policy")


def test_domain_registry_is_explicit_and_fail_closed():
    registry = default_domain_registry()
    assert registry.names == ("avp",)
    assert registry.create("AVP", require_approved=False).profile.name == "avp"
    with pytest.raises(ValueError, match="未注册Domain"):
        registry.create("wiper", require_approved=False)
    with pytest.raises(ValueError, match="重复注册"):
        DomainRegistry().register("avp", AVPDomainPolicy).register("AVP", AVPDomainPolicy)


def test_avp_safety_goal_catalog_comes_from_domain_profile():
    profile = load_domain_profile("avp", require_approved=False)
    catalog = AVPSafetyGoalCatalog(profile)

    assert catalog.classify("输出制动扭矩") == "SG_AVP_02"
    assert catalog.classify("车辆转向扭矩请求") == "SG_AVP_03"
    assert catalog.definitions["SG_AVP_07"]["safety_goal"].startswith("应及时、正确地")

    catalog.register("SG_AVP_02", "输出制动扭矩", "失效A", "loss", "SCN-1", "事件A", "A", {
        "ftti_value_s": 0.5, "ftti_status": "NEEDS_REVIEW"
    })
    catalog.register("SG_AVP_02", "输出制动扭矩", "失效B", "late", "SCN-2", "事件B", "B", {
        "ftti_value_s": 0.3, "ftti_status": "NEEDS_REVIEW"
    })
    aggregated = catalog.to_dict()["SG_AVP_02"]
    assert aggregated["max_asil"] == "B"
    assert aggregated["ftti_value_s"] == 0.3
    assert catalog.classify("未定义的AVP功能") is None


def test_unapproved_domain_profile_is_draft_only():
    data = {
        "metadata": {
            "subsystem": "avp",
            "domain_profile": {
                "name": "avp",
                "version": "1.0.0-migration",
                "approval_status": "migration_baseline",
            },
        }
    }

    formal = ScoringChecker("unused.json", allow_draft=False)
    formal.data = data
    formal._check_domain_profile()
    assert formal.errors[0]["type"] == "domain_profile_not_approved"

    draft = ScoringChecker("unused.json", allow_draft=True)
    draft.data = data
    draft._check_domain_profile()
    assert draft.warnings[0]["type"] == "domain_profile_not_approved"


def test_avp_candidate_math_and_review_metadata_come_from_profile():
    profile = load_domain_profile("avp", require_approved=False)
    policy = AVPDomainPolicy(profile)

    pedestrian = policy.candidate("near_pedestrian", ego_speed_kph=5.0)
    vehicle = policy.candidate("near_vehicle_10", ego_speed_kph=5.0)

    assert pedestrian["relative_speed_kph"] == 5.0
    assert vehicle["relative_speed_kph"] == 15.0
    assert vehicle["engineering_status"] == "PENDING"
    assert vehicle["rule_version"] == profile.version


def test_avp_scoring_decisions_come_from_versioned_domain_policy():
    profile = load_domain_profile("avp", require_approved=False)
    policy = AVPDomainPolicy(profile)

    pedestrian = policy.candidate("near_pedestrian", ego_speed_kph=5.0)
    vehicle = policy.candidate("near_vehicle_30", ego_speed_kph=5.0)
    _apply_driver_context(vehicle, profile)

    assert policy.severity_decision(pedestrian)["score"] == "S1"
    assert policy.severity_decision(vehicle)["score"] == "S2"
    assert policy.exposure_decision(vehicle)["score"] == "E1"
    assert policy.controllability_decision(vehicle)["score"] == "C3"
    assert policy.severity_decision(vehicle)["rule_version"] == profile.version


def test_structured_avp_scoring_cannot_silently_fall_back_to_keywords():
    profile = load_domain_profile("avp", require_approved=False)
    policy = AVPDomainPolicy(profile)
    incomplete = {"scenario_variant": "unknown", "relative_speed_kph": 5}

    with pytest.raises(ValueError, match="禁止回退"):
        SeverityReasoningEngine(policy).analyze("collision", incomplete)

    controllability = ControllabilityReasoningEngine(policy)
    assert controllability._ctrl_ref_loaded is False
    with pytest.raises(ValueError, match="禁止回退"):
        controllability.analyze("collision", refined_scenario=incomplete)
    assert controllability._ctrl_ref_loaded is False


def test_domain_scoring_service_owns_structured_scoring_path():
    profile = load_domain_profile("avp", require_approved=False)
    policy = AVPDomainPolicy(profile)
    scenario = policy.candidate("near_pedestrian", ego_speed_kph=5.0)
    _apply_driver_context(scenario, profile)

    result = _scoring_service(policy).score(scenario, "车辆可能碰撞行人")

    assert result["severity"]["severity_score"] == "S1"
    assert result["exposure"]["exposure_score"] == "E4"
    assert result["exposure"]["exposure_method"] == "F"
    assert result["controllability"]["controllability_score"] == "C3"
    assert result["exposure"]["template_standard_sheet"] == "Exposure"
    assert all(
        item["_source"] == "domain_policy_structured_scoring"
        for item in result.values()
    )


def test_exposure_method_comes_from_evidence_not_from_e_level():
    profile = load_domain_profile("avp", require_approved=False)
    policy = AVPDomainPolicy(profile)
    time_based = _apply_driver_context(
        policy.candidate("controlled_standstill", ego_speed_kph=5.0), profile
    )
    frequency_based = _apply_driver_context(
        policy.candidate("near_pedestrian", ego_speed_kph=5.0), profile
    )

    time_result = _scoring_service(policy).score(time_based, "车辆保持静止")["exposure"]
    frequency_result = _scoring_service(policy).score(
        frequency_based, "车辆可能碰撞行人"
    )["exposure"]

    assert time_result["exposure_score"] == frequency_result["exposure_score"] == "E4"
    assert time_result["exposure_method"] == "T"
    assert frequency_result["exposure_method"] == "F"
    assert time_result["template_standard_criterion"].startswith(">10")
    assert "almost every drive" in frequency_result["template_standard_criterion"].lower()


def test_checker_flags_legacy_text_scoring_path():
    checker = ScoringChecker("unused.json", allow_draft=False)
    checker.data = {
        "severity_results": {
            "f": [{
                "scenario_id": "SCN-1",
                "_source": "legacy_text_scoring_compatibility",
                "engineering_rule_id": "LEGACY-TEXT-SEVERITY",
            }]
        }
    }
    checker._check_legacy_scoring_compatibility()
    assert checker.errors[0]["type"] == "legacy_text_scoring_used"


def test_asil_service_requires_run_template_and_reads_its_matrix():
    with pytest.raises(ValueError, match="显式提供"):
        TemplateASILService("")

    service = TemplateASILService(str(ROOT / "references" / "HARA_Template_AI_20260327.xlsx"))
    assert len(service.matrix) == 80
    assert service.determine("S0", "E4", "C3") == "QM"
    assert service.determine("S1", "E4", "C3") == "B"
    assert service.source.endswith("#ASIL_Table")


def test_ftti_service_keeps_ttc_as_screening_and_candidate_pending():
    result = FTTIService().evaluate({
        "scenario_id": "SCN-FTTI",
        "relative_distance": "0.5 m",
        "relative_speed_kph": 5,
        "collision_geometry": "frontal/rear vehicle",
    }, "车辆碰撞", "B")

    assert result["formula_id"] == "longitudinal_t1_t2_draft"
    assert result["ftti_status"] == "NEEDS_REVIEW"
    assert result["ttc_screening_s"] != result["ftti_value_s"]


def test_risk_aggregation_preserves_object_class_and_audits_coverage():
    scenarios = [
        {"scenario_id": "P1", "object_type": "行人", "collision_geometry": "vehicle-pedestrian"},
        {"scenario_id": "P2", "object_type": "行人", "collision_geometry": "vehicle-pedestrian"},
        {"scenario_id": "V1", "object_type": "车辆", "collision_geometry": "frontal/rear vehicle"},
    ]
    scores = [{"severity_score": "S1"}] * 3
    exposure = [{"exposure_score": "E4"}] * 3
    controllability = [{"controllability_score": "C3"}] * 3
    asil = [{"ASIL": "B"}] * 3
    goals = [{"sg_id": "SG_AVP_01"}] * 3
    ftti = [{"ftti_status": "NEEDS_REVIEW", "ftti_value_s": 0.3, "formula_id": "draft"}] * 3

    selected, groups = RiskAggregationService().select(
        scenarios, scores, exposure, controllability, asil, goals, ftti
    )

    assert selected == [0, 2]
    assert groups[0]["covered_scenario_ids"] == ["P1", "P2"]
    assert groups[1]["covered_scenario_ids"] == ["V1"]


def test_v13_adapter_builds_typed_agent_state_from_phase3_output():
    scenario = {
        "scenario_id": "SCN-1", "operating_scenario": "Parking",
        "situational_description": "停车场低速泊车", "situational_detailing": "儿童位于轨迹前方",
        "engineering_status": "PENDING", "rule_version": "1.1", "parameter_sources": {"speed": "DOC-1"},
    }
    data = {
        "metadata": {
            "asil_matrix_source": "template.xlsx#ASIL_Table",
            "domain_profile": {"approval_status": "migration_baseline"},
        },
        "scenario_catalog": {"SCN-1": scenario},
        "severity_results": {"f": [{"scenario_id": "SCN-1", "severity_score": "S1", "engineering_status": "PENDING", "engineering_rule_version": "1.1"}]},
        "exposure_results": {"f": [{"scenario_id": "SCN-1", "exposure_score": "E4", "engineering_status": "PENDING", "engineering_rule_version": "1.1"}]},
        "controllability_results": {"f": [{"scenario_id": "SCN-1", "controllability_score": "C3", "engineering_status": "PENDING", "engineering_rule_version": "1.1"}]},
        "asil_results": {"f": [{"scenario_id": "SCN-1", "ASIL": "B"}]},
        "ftti_results": {"f": [{"scenario_id": "SCN-1", "ftti_value_s": 0.3, "ftti_status": "NEEDS_REVIEW", "formula_id": "draft"}]},
        "safety_goals": {"f": [{"scenario_id": "SCN-1", "sg_id": "SG-1"}]},
        "safety_goal_catalog": {"SG-1": {
            "safety_goal": "避免碰撞", "safety_state": "受控制动至静止", "max_asil": "B",
            "ftti_value_s": 0.3, "ftti_status": "NEEDS_REVIEW", "associations": [{"scenario_id": "SCN-1"}],
        }},
    }
    adapter = V13ScoringAdapter(data)

    scenarios = adapter.scenarios()
    risks = adapter.risks()
    goals = adapter.safety_goals()

    assert len(scenarios) == 1
    assert len(risks) == 1
    assert len(goals) == 1
    assert all(risk.asil.status is ReviewStatus.FINALIZED for risk in risks)
    assert all(goal.status is ReviewStatus.PENDING for goal in goals)
    assert {risk.scenario_id for risk in risks}.issubset({item.scenario_id for item in scenarios})
    assert len({risk.assessment_id for risk in risks}) == len(risks)

    state = import_v13_scoring(HARAState(run_id="adapter-test"), data)
    assert state.stage.value == "quality_gate"
    assert not state.can_publish
    assert state.audit_trail[-1]["event"] == "v13_scoring_imported"


def test_function_agent_uses_llm_semantics_but_enforces_typed_contract():
    class FakeClient:
        def complete_json(self, request):
            assert request.prompt_version == "function-extraction-v3"
            assert request.max_tokens == 12288
            return LLMResponse(data={"functions": [{
                "function_id": "FUN-01", "name": "输出制动扭矩", "output": "目标制动扭矩请求",
                "description": "AVP向车辆制动系统请求减速度", "preconditions": ["AVP已激活"],
                "triggers": ["检测到停车需求或障碍物"], "odd_constraints": ["停车场低速运行"],
                "fallback_behavior": "进入受控静止", "consequences": ["制动不足可能导致碰撞"],
                "source_location": "paragraph 3", "source_excerpt": "输出目标制动扭矩",
                "confidence": 0.92, "status": "PENDING",
            }]}, model="fake-model", request_id="req-1")

    state = extract_functions(
        HARAState(run_id="llm-test"), FunctionExtractionAgent(FakeClient()),
        "AVP输出目标制动扭矩。", "ItemDef.docx",
    )
    assert state.functions[0]["name"] == "输出制动扭矩"
    assert state.functions[0]["output"] == "目标制动扭矩请求"
    assert state.audit_trail[-1]["model"] == "fake-model"
    json.dumps(state.to_dict(), ensure_ascii=False)


def test_function_agent_rejects_function_output_copy():
    class InvalidClient:
        def complete_json(self, request):
            return LLMResponse(data={"functions": [{
                "function_id": "FUN-X", "name": "同一段文本", "output": "同一段文本",
                "source_location": "p1", "confidence": 0.5,
            }]}, model="fake-model")

    with pytest.raises(ValueError, match="不得完全相同"):
        FunctionExtractionAgent(InvalidClient()).extract("输入文本", "ItemDef.docx")


def test_function_agent_rejects_section_heading_as_function():
    class HeadingClient:
        def complete_json(self, request):
            return LLMResponse(data={"functions": [{
                "function_id": "FUN-H", "name": "已识别的潜在后果", "output": "风险说明",
                "source_location": "paragraph 8", "source_excerpt": "已识别的潜在后果",
                "confidence": 0.8,
            }]}, model="fake-model")

    with pytest.raises(ValueError, match="section_as_function"):
        FunctionExtractionAgent(HeadingClient()).extract("输入文本", "ItemDef.docx")


def test_openai_compatible_provider_is_configurable_and_returns_json():
    captured = {}

    def transport(url, headers, body, timeout):
        captured.update({"url": url, "headers": headers, "body": json.loads(body), "timeout": timeout})
        return {
            "id": "req-provider", "model": "test-model",
            "choices": [{"message": {"content": "```json\n{\"functions\": []}\n```"}}],
            "usage": {"total_tokens": 12},
        }

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret", timeout_seconds=9,
    ), transport=transport)
    response = client.complete_json(LLMRequest(
        task="test", system_prompt="system", user_prompt="user",
        schema_name="test", prompt_version="v1",
    ))

    assert response.data == {"functions": []}
    assert captured["url"] == "https://llm.example/v1/chat/completions"
    assert captured["headers"]["Authorization"] == "Bearer secret"
    assert "response_format" not in captured["body"]
    assert captured["body"]["max_tokens"] == 32768


def test_agent_plan_provider_alias_is_accepted_as_openai_compatible():
    client = create_llm_client(LLMConfig(
        provider="volcengine-agent-plan", base_url="https://llm.example/v1",
        model="test-model", api_key="secret", timeout_seconds=9,
    ))

    assert client is not None


def test_openai_compatible_extracts_json_surrounded_by_model_prose():
    content = "分析完成。\n```json\n{\"functions\": [{\"name\": \"AVP\"}]}\n```\n以上为结果。"

    parsed = OpenAICompatibleClient._parse_json_content(content)

    assert parsed == {"functions": [{"name": "AVP"}]}


def test_openai_compatible_json_extraction_handles_braces_inside_strings():
    content = 'result: {"description": "保留 {占位符}", "items": []} done'

    parsed = OpenAICompatibleClient._parse_json_content(content)

    assert parsed["description"] == "保留 {占位符}"


def test_openai_compatible_invalid_response_reports_safe_preview():
    with pytest.raises(ValueError, match="响应摘要"):
        OpenAICompatibleClient._parse_json_content("模型只返回了普通说明，没有结构化结果")


def test_item_definition_agent_accepts_unwrapped_schema_fields():
    payload = {
        "system_description": "AVP系统",
        "item_boundary": "停车场内",
        "operating_modes": ["泊车"],
        "odd": {"locations": ["停车场"], "speed_range_kph": [0, 10]},
    }

    raw = ItemDefinitionExtractionAgent._unwrap_payload(payload)

    assert raw is payload


def test_item_definition_agent_rejects_unrelated_json_object():
    with pytest.raises(ValueError, match="顶层字段"):
        ItemDefinitionExtractionAgent._unwrap_payload({"answer": "done"})


def test_item_definition_agent_normalizes_nullable_lists_and_numeric_confidence():
    raw = {
        "system_description": "AVP系统",
        "item_boundary": "停车场内",
        "operating_modes": None,
        "odd": {"locations": None, "road_types": None, "speed_range_kph": [None, None]},
        "performance_parameters": None,
        "driver_contexts": None,
        "exposure_inputs": None,
        "confidence": "0.8",
    }

    facts = ItemDefinitionExtractionAgent._parse(raw, "ItemDef.docx")

    assert facts.operating_modes == []
    assert facts.odd_road_types == []
    assert facts.speed_min_kph is None
    assert facts.performance_parameters == []
    assert facts.confidence == 0.8


def test_item_definition_agent_rejects_qualitative_confidence():
    with pytest.raises(ValueError, match="Item Definition confidence.*actual='high'"):
        ItemDefinitionExtractionAgent._parse({
            "system_description": "AVP系统",
            "item_boundary": "停车场内",
            "odd": {"speed_range_kph": [0, 5]},
            "confidence": "high",
        }, "ItemDef.docx")


def test_item_definition_agent_normalizes_reversed_speed_and_forces_review():
    raw = {
        "system_description": "AVP系统",
        "item_boundary": "停车场内",
        "odd": {"speed_range_kph": [10, 0]},
        "status": "FINALIZED",
    }

    facts, warnings = ItemDefinitionExtractionAgent._parse_with_warnings(
        raw, "ItemDef.docx"
    )

    assert (facts.speed_min_kph, facts.speed_max_kph) == (0.0, 10.0)
    assert facts.status is ReviewStatus.PENDING
    assert "顺序颠倒" in warnings[0]


def test_item_definition_agent_invalid_speed_degrades_to_pending_without_crash():
    raw = {
        "system_description": "AVP系统",
        "item_boundary": "停车场内",
        "odd": {"speed_range_kph": [-1, 10]},
        "status": "FINALIZED",
    }

    facts, warnings = ItemDefinitionExtractionAgent._parse_with_warnings(
        raw, "ItemDef.docx"
    )

    assert facts.speed_min_kph is None and facts.speed_max_kph is None
    assert facts.status is ReviewStatus.PENDING
    assert warnings


def test_item_definition_agent_accepts_unambiguous_string_speed_range():
    assert ItemDefinitionExtractionAgent._speed_range("0-10 km/h") == (0.0, 10.0, "")


def test_openai_compatible_reports_finish_reason_for_empty_content():
    def transport(url, headers, body, timeout):
        return {"choices": [{"finish_reason": "length", "message": {"content": ""}}]}

    client = create_llm_client(LLMConfig(
        provider="volcengine-agent-plan", base_url="https://llm.example/v1",
        model="test-model", api_key="secret",
    ), transport=transport)

    with pytest.raises(ValueError, match="finish_reason=length"):
        client.complete_json(LLMRequest(
            task="test", system_prompt="system", user_prompt="user",
            schema_name="test", prompt_version="v1",
        ))


def test_openai_compatible_length_rejects_salvageable_inner_object():
    truncated = '''{
      "assessments": [
        {"guideword": "No", "applicable": true},
        {"guideword":
    '''

    def transport(url, headers, body, timeout):
        return {"choices": [{
            "finish_reason": "length", "message": {"content": truncated},
        }]}

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret",
    ), transport=transport)

    with pytest.raises(LLMOutputLimitError, match="结构化结果不完整"):
        client.complete_json(LLMRequest(
            task="assess_guideword_applicability",
            system_prompt="system", user_prompt="user",
            schema_name="GuidewordAssessmentList", prompt_version="v1",
        ))


def test_openai_compatible_stop_salvages_complete_schema_from_prose():
    content = 'Here is the JSON:\n{"assessments": []}\nDone.'

    def transport(url, headers, body, timeout):
        return {"choices": [{
            "finish_reason": "stop", "message": {"content": content},
        }]}

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret",
    ), transport=transport)
    response = client.complete_json(LLMRequest(
        task="assess_guideword_applicability",
        system_prompt="system", user_prompt="user",
        schema_name="GuidewordAssessmentList", prompt_version="v1",
    ))

    assert response.data == {"assessments": []}


def test_openai_compatible_stop_rejects_wrong_guideword_top_level_schema():
    def transport(url, headers, body, timeout):
        return {"choices": [{
            "finish_reason": "stop",
            "message": {"content": '{"guideword": "No", "applicable": true}'},
        }]}

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret",
    ), transport=transport)

    with pytest.raises(ValueError, match="GuidewordAssessmentList必须包含assessments"):
        client.complete_json(LLMRequest(
            task="assess_guideword_applicability",
            system_prompt="system", user_prompt="user",
            schema_name="GuidewordAssessmentList", prompt_version="v1",
        ))


@pytest.mark.parametrize("payload", [
    [{"guideword": "loss", "applicable": True}],
    {"guideword_assessments": [{"guideword": "loss", "applicable": True}]},
    {"data": {"assessments": [{"guideword": "loss", "applicable": True}]}},
    {"result": {"assessments": [{"guideword": "loss", "applicable": True}]}},
])
def test_openai_compatible_normalizes_safe_guideword_envelopes(payload):
    def transport(url, headers, body, timeout):
        return {"choices": [{
            "finish_reason": "stop",
            "message": {"content": json.dumps(payload)},
        }]}

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret",
    ), transport=transport)
    response = client.complete_json(LLMRequest(
        task="assess_guideword_applicability",
        system_prompt="system", user_prompt="user",
        schema_name="GuidewordAssessmentList", prompt_version="v1",
    ))

    assert isinstance(response.data["assessments"], list)
    assert response.data["assessments"][0]["guideword"] == "loss"


def test_openai_compatible_does_not_wrap_single_guideword_object():
    data = {"guideword": "loss", "applicable": True}
    normalized = OpenAICompatibleClient._normalize_schema_envelope(
        data, "GuidewordAssessmentList",
    )
    with pytest.raises(ValueError, match="必须包含assessments"):
        OpenAICompatibleClient._validate_schema_envelope(
            normalized, "GuidewordAssessmentList",
        )


@pytest.mark.parametrize("payload", [
    {"candidates": []},
    {"candidates": [], "warnings": []},
])
def test_malfunction_schema_envelope_accepts_candidates_list_and_extensions(payload):
    OpenAICompatibleClient._validate_schema_envelope(
        payload, "MalfunctionHazardCandidateList",
    )


@pytest.mark.parametrize(("payload", "message"), [
    ({"malfunctions": []}, "必须包含candidates"),
    ([], "顶层必须为JSON object"),
    ({"candidates": {}}, "candidates类型必须为list"),
])
def test_malfunction_schema_envelope_rejects_drift_without_alias_guessing(payload, message):
    with pytest.raises(LLMSchemaContractError, match=message):
        OpenAICompatibleClient._validate_schema_envelope(
            payload, "MalfunctionHazardCandidateList",
        )


def _malfunction_candidate_payload(description="request lost"):
    return {
        "malfunction_id": "MF-1", "guideword": "loss", "description": description,
        "functional_effect": "braking unavailable",
        "vehicle_level_hazard": "vehicle continues moving",
        "causal_chain": ["request lost", "vehicle does not decelerate"],
        "confidence": 0.8, "status": "PENDING",
    }


def _malfunction_schema_request():
    return LLMRequest(
        task="derive_malfunctions_and_hazards",
        system_prompt="system", user_prompt="user",
        schema_name="MalfunctionHazardCandidateList", prompt_version="v1",
        metadata={"function_id": "F04"}, max_tokens=8192,
    )


def test_malfunction_schema_mismatch_repairs_once_and_preserves_candidates(capsys):
    candidate = _malfunction_candidate_payload()
    responses = iter([
        {"malfunctions": [candidate]},
        {"candidates": [candidate]},
    ])
    calls = []

    def transport(url, headers, body, timeout):
        calls.append(json.loads(body))
        payload = next(responses)
        return {"choices": [{
            "finish_reason": "stop", "message": {"content": json.dumps(payload)},
        }]}

    client = create_llm_client(LLMConfig(
        provider="volcengine-agent-plan", base_url="https://llm.example/v1",
        model="test-model", api_key="secret",
    ), transport=transport)
    response = client.complete_json(_malfunction_schema_request())

    assert response.data == {"candidates": [candidate]}
    assert response.usage["schema_repair_count"] == 1
    assert len(calls) == 2
    assert calls[1]["thinking"] == {"type": "disabled"}
    logs = capsys.readouterr().err
    assert "function=F04" in logs
    assert "top_level_keys=['malfunctions']" in logs
    assert "schema repair completed" in logs


def test_malfunction_schema_repair_failure_is_fail_closed_after_one_attempt():
    candidate = _malfunction_candidate_payload()
    responses = iter([
        {"malfunctions": [candidate]},
        {"results": [candidate]},
    ])
    calls = 0

    def transport(url, headers, body, timeout):
        nonlocal calls
        calls += 1
        return {"choices": [{
            "finish_reason": "stop",
            "message": {"content": json.dumps(next(responses))},
        }]}

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret",
    ), transport=transport)

    with pytest.raises(LLMSchemaContractError, match="必须包含candidates"):
        client.complete_json(_malfunction_schema_request())
    assert calls == 2


def test_malfunction_schema_repair_rejects_semantic_candidate_changes():
    candidate = _malfunction_candidate_payload()
    changed = _malfunction_candidate_payload(description="changed by repair")
    responses = iter([
        {"malfunctions": [candidate]},
        {"candidates": [changed]},
    ])

    def transport(url, headers, body, timeout):
        return {"choices": [{
            "finish_reason": "stop",
            "message": {"content": json.dumps(next(responses))},
        }]}

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret",
    ), transport=transport)

    with pytest.raises(LLMSchemaContractError, match="changed semantic candidates"):
        client.complete_json(_malfunction_schema_request())


def test_malfunction_length_failure_never_enters_schema_repair():
    calls = 0

    def transport(url, headers, body, timeout):
        nonlocal calls
        calls += 1
        return {"choices": [{
            "finish_reason": "length",
            "message": {"content": '{"malfunctions": ['},
        }]}

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret",
    ), transport=transport)

    with pytest.raises(LLMOutputLimitError):
        client.complete_json(_malfunction_schema_request())
    assert calls == 1


def test_openai_compatible_diagnostic_logs_lengths_and_safe_usage(capsys):
    def transport(url, headers, body, timeout):
        return {
            "choices": [{
                "finish_reason": "stop",
                "message": {"content": '{"ok": true}', "reasoning_content": "hidden"},
            }],
            "usage": {"prompt_tokens": 10, "completion_tokens": 5, "total_tokens": 15},
        }

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="super-secret",
    ), transport=transport)
    client.complete_json(LLMRequest(
        task="diagnostic_task", system_prompt="system", user_prompt="user",
        schema_name="DiagnosticSchema", prompt_version="v1", max_tokens=4096,
    ))

    log = capsys.readouterr().err
    assert "task=diagnostic_task" in log
    assert "schema_name=DiagnosticSchema" in log
    assert "reasoning_chars=6" in log
    assert "total_tokens" in log
    assert "hidden" not in log
    assert "super-secret" not in log


def test_openai_compatible_retries_transient_timeout_then_succeeds():
    calls = []

    def transport(url, headers, body, timeout):
        calls.append(timeout)
        if len(calls) == 1:
            raise TimeoutError("read timed out")
        return {"choices": [{"message": {"content": '{"ok": true}'}}]}

    client = create_llm_client(LLMConfig(
        provider="volcengine-agent-plan", base_url="https://llm.example/v1",
        model="test-model", api_key="secret", max_retries=1,
        retry_backoff_seconds=0,
    ), transport=transport)

    response = client.complete_json(LLMRequest(
        task="test", system_prompt="system", user_prompt="user",
        schema_name="test", prompt_version="v1",
    ))

    assert response.data == {"ok": True}
    assert len(calls) == 2
    assert response.usage["transport_attempts"] == 2
    assert response.usage["transient_error_counts"] == {"TimeoutError": 1}


def test_transient_failure_log_includes_type_without_credentials(capsys):
    calls = 0

    def transport(url, headers, body, timeout):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise TimeoutError("secret-token must not be logged")
        return {"choices": [{"message": {"content": '{"ok": true}'}}]}

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="super-secret", max_retries=1,
        retry_backoff_seconds=0,
    ), transport=transport)
    client.complete_json(LLMRequest(
        task="test", system_prompt="system", user_prompt="user",
        schema_name="test", prompt_version="v1",
    ))

    log = capsys.readouterr().err
    assert "type=TimeoutError" in log
    assert "secret-token" not in log
    assert "super-secret" not in log


def test_openai_compatible_stops_after_configured_retry_limit():
    def transport(url, headers, body, timeout):
        raise TimeoutError("read timed out")

    client = create_llm_client(LLMConfig(
        provider="volcengine-agent-plan", base_url="https://llm.example/v1",
        model="test-model", api_key="secret", max_retries=1,
        retry_backoff_seconds=0,
    ), transport=transport)

    with pytest.raises(RuntimeError, match="2次尝试"):
        client.complete_json(LLMRequest(
            task="test", system_prompt="system", user_prompt="user",
            schema_name="test", prompt_version="v1",
        ))


def test_scenario_incomplete_read_retries_same_payload_then_succeeds():
    calls = []

    def transport(url, headers, body, timeout):
        calls.append(body)
        if len(calls) == 1:
            raise http.client.IncompleteRead(b'{"partial":', 2445)
        return {"choices": [{"message": {"content": '{"assessments": []}'}}]}

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret", max_retries=2, retry_backoff_seconds=0,
    ), transport=transport)
    result = client.complete_json(LLMRequest(
        task="assess_scenario_feasibility", system_prompt="system", user_prompt="user",
        schema_name="ScenarioFeasibilityAssessmentList", prompt_version="v1",
        metadata={"malfunction_id": "MF-F01-003", "parent_batch": "1/5",
                  "split_path": "R", "split_depth": 1, "scenario_count": 6},
    ))
    assert result.usage["transport_attempts"] == 2
    assert result.usage["transient_error_counts"] == {"incomplete_read": 1}
    assert calls[0] == calls[1]


def test_scenario_incomplete_read_exhaustion_is_transport_error_not_split():
    def transport(url, headers, body, timeout):
        raise http.client.IncompleteRead(b'{"choices":[', 10)

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret", max_retries=2, retry_backoff_seconds=0,
    ), transport=transport)
    with pytest.raises(LLMTransportError) as error:
        client.complete_json(LLMRequest(
            task="assess_scenario_feasibility", system_prompt="system", user_prompt="user",
            schema_name="ScenarioFeasibilityAssessmentList", prompt_version="v1",
        ))
    assert error.value.category == "incomplete_read"
    assert error.value.attempts == 2


@pytest.mark.parametrize("transport_error, category", [
    (http.client.RemoteDisconnected("closed"), "remote_disconnect"),
    (ConnectionResetError("reset"), "connection_reset"),
])
def test_transport_disconnect_categories_are_typed(transport_error, category):
    def transport(url, headers, body, timeout):
        raise transport_error

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret", max_retries=0,
    ), transport=transport)
    with pytest.raises(LLMTransportError) as error:
        client.complete_json(LLMRequest(
            task="test", system_prompt="system", user_prompt="user",
            schema_name="test", prompt_version="v1",
        ))
    assert error.value.category == category


def _scenario_json_request():
    return LLMRequest(
        task="assess_scenario_feasibility", system_prompt="system", user_prompt="user",
        schema_name="ScenarioFeasibilityAssessmentList", prompt_version="v7",
        metadata={"malfunction_id": "MF-1", "parent_batch": "1/1",
                  "split_path": "root", "split_depth": 0, "scenario_count": 1},
    )


@pytest.mark.parametrize(("content", "normalized"), [
    ('{"assessments": []}', 0),
    ('```json\n{"assessments": []}\n```', 1),
    ('  ```JSON\n\n{"assessments": []}\n\n```  ', 1),
])
def test_scenario_json_accepts_raw_or_single_outer_fence(content, normalized):
    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret", max_retries=0,
    ), transport=lambda *args: {
        "choices": [{"finish_reason": "stop", "message": {"content": content}}],
    })
    result = client.complete_json(_scenario_json_request())
    assert result.data == {"assessments": []}
    assert result.usage["markdown_fence_normalizations"] == normalized
    assert result.usage["format_retry_calls"] == 0


@pytest.mark.parametrize("content", [
    'Here is the result:\n{"assessments": []}',
    '{"assessments": []}\nHope this helps.',
    '{"assessments": [{"scenario_id":',
    '{"assessments": []',
    '{"rationale": "abc "def"}',
])
def test_scenario_strict_json_parser_rejects_explanation_and_malformed(content):
    with pytest.raises(LLMJSONContractError):
        OpenAICompatibleClient._parse_scenario_json_content(
            content, request=_scenario_json_request(), finish_reason="stop",
        )


def test_scenario_malformed_stop_json_retries_format_once_then_succeeds():
    calls = []

    def transport(url, headers, body, timeout):
        calls.append(json.loads(body))
        content = '{"assessments": [' if len(calls) == 1 else '{"assessments": []}'
        return {"choices": [{"finish_reason": "stop", "message": {"content": content}}]}

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret", max_retries=0,
    ), transport=transport)
    result = client.complete_json(_scenario_json_request())
    assert len(calls) == 2
    assert "FORMAT RETRY" in calls[1]["messages"][0]["content"]
    assert result.usage["json_contract_errors"] == 1
    assert result.usage["format_retry_calls"] == 1
    assert result.usage["format_retry_successes"] == 1
    assert result.usage["format_retry_failures"] == 0


def test_scenario_malformed_stop_json_twice_fails_typed_without_split():
    calls = []

    def transport(url, headers, body, timeout):
        calls.append(body)
        return {"choices": [{"finish_reason": "stop", "message": {
            "content": '```json\n{"assessments": [\n```',
        }}]}

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret", max_retries=0,
    ), transport=transport)
    with pytest.raises(LLMJSONContractError) as error:
        client.complete_json(_scenario_json_request())
    assert len(calls) == 2
    assert error.value.diagnostics["format_retry_calls"] == 1
    assert error.value.diagnostics["format_retry_failures"] == 1
    assert error.value.diagnostics["outer_fence_removed"] is True


def test_scenario_machine_output_rule_requires_raw_json_envelope():
    rule = ScenarioFeasibilityAgent.MACHINE_OUTPUT_RULE
    assert "Do not use Markdown" in rule
    assert "first non-whitespace character must be {" in rule
    assert "last non-whitespace character must be }" in rule
    assert '{"assessments": [...]}' in rule


def test_ordered_parallel_map_runs_concurrently_and_preserves_order():
    thread_names = set()
    lock = threading.Lock()

    def worker(value):
        with lock:
            thread_names.add(threading.current_thread().name)
        time.sleep(0.02 if value % 2 == 0 else 0.01)
        return value * 10

    progress = []
    result = ordered_parallel_map(
        [0, 1, 2, 3], worker, max_workers=4,
        on_progress=lambda done, total: progress.append((done, total)),
    )

    assert result == [0, 10, 20, 30]
    assert len(thread_names) > 1
    assert progress[-1] == (4, 4)


def test_ordered_parallel_map_serial_mode_uses_same_contract():
    assert ordered_parallel_map([1, 2], lambda value: value + 1, max_workers=1) == [2, 3]


def test_ordered_parallel_map_propagates_worker_error():
    def worker(value):
        if value == 2:
            raise RuntimeError("provider failed")
        return value

    with pytest.raises(RuntimeError, match="provider failed"):
        ordered_parallel_map([1, 2, 3], worker, max_workers=3)


def test_llm_max_tokens_is_configurable_from_environment(monkeypatch):
    monkeypatch.setenv("HARA_LLM_PROVIDER", "volcengine-agent-plan")
    monkeypatch.setenv("HARA_LLM_BASE_URL", "https://llm.example/v1")
    monkeypatch.setenv("HARA_LLM_MODEL", "test-model")
    monkeypatch.setenv("HARA_LLM_API_KEY", "secret")
    monkeypatch.setenv("HARA_LLM_MAX_TOKENS", "8192")

    assert LLMConfig.from_env().max_tokens == 8192


def test_openai_compatible_request_token_budget_overrides_global_default():
    captured = {}

    def transport(url, headers, body, timeout):
        captured.update(json.loads(body))
        return {"choices": [{"message": {"content": '{"ok": true}'}}]}

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret", max_tokens=32768,
    ), transport=transport)
    client.complete_json(LLMRequest(
        task="test", system_prompt="system", user_prompt="user",
        schema_name="test", prompt_version="v1", max_tokens=4096,
    ))

    assert captured["max_tokens"] == 4096


def _captured_body_for_task(
    task, provider="volcengine-agent-plan", thinking="disabled",
    malfunction_thinking="default", scenario_thinking="default",
):
    captured = {}

    def transport(url, headers, body, timeout):
        captured.update(json.loads(body))
        return {"choices": [{"finish_reason": "stop", "message": {"content": '{"ok": true}'}}]}

    client = create_llm_client(LLMConfig(
        provider=provider, base_url="https://llm.example/v1",
        model="test-model", api_key="secret", extraction_thinking=thinking,
        malfunction_thinking=malfunction_thinking,
        scenario_thinking=scenario_thinking,
    ), transport=transport)
    client.complete_json(LLMRequest(
        task=task, system_prompt="system", user_prompt="user",
        schema_name="test", prompt_version="v1",
    ))
    return captured


def test_core_extraction_disables_thinking_for_supported_provider():
    body = _captured_body_for_task("extract_core_item_artifacts")
    assert body["thinking"] == {"type": "disabled"}


def test_project_evidence_supplement_disables_thinking_for_supported_provider():
    body = _captured_body_for_task("supplement_project_evidence")
    assert body["thinking"] == {"type": "disabled"}


def test_odd_repair_supplement_disables_thinking_for_supported_provider():
    body = _captured_body_for_task("supplement_odd_repair")
    assert body["thinking"] == {"type": "disabled"}


def test_guideword_applicability_disables_thinking_for_supported_provider():
    body = _captured_body_for_task("assess_guideword_applicability")
    assert body["thinking"] == {"type": "disabled"}


def test_reasoning_task_keeps_provider_default_thinking_behavior():
    body = _captured_body_for_task("hazard_analysis")
    assert "thinking" not in body


def test_malfunction_task_keeps_provider_default_thinking_behavior():
    body = _captured_body_for_task("derive_malfunctions_and_hazards")
    assert "thinking" not in body


def test_malfunction_thinking_can_be_disabled_independently():
    body = _captured_body_for_task(
        "derive_malfunctions_and_hazards", malfunction_thinking="disabled",
    )
    assert body["thinking"] == {"type": "disabled"}


def test_scenario_task_keeps_provider_default_thinking_behavior():
    body = _captured_body_for_task("assess_scenario_feasibility")
    assert "thinking" not in body


def test_scenario_thinking_can_be_disabled_independently():
    body = _captured_body_for_task(
        "assess_scenario_feasibility", scenario_thinking="disabled",
    )
    assert body["thinking"] == {"type": "disabled"}


def test_generic_openai_compatible_does_not_receive_thinking_parameter():
    body = _captured_body_for_task(
        "extract_core_item_artifacts", provider="openai-compatible",
    )
    assert "thinking" not in body


def test_extraction_thinking_is_configurable_from_environment(monkeypatch):
    monkeypatch.setenv("HARA_LLM_PROVIDER", "volcengine-agent-plan")
    monkeypatch.setenv("HARA_LLM_BASE_URL", "https://llm.example/v1")
    monkeypatch.setenv("HARA_LLM_MODEL", "test-model")
    monkeypatch.setenv("HARA_LLM_API_KEY", "secret")
    monkeypatch.setenv("HARA_LLM_EXTRACTION_THINKING", "auto")

    assert LLMConfig.from_env().extraction_thinking == "auto"


def test_guideword_thinking_is_configurable_from_environment(monkeypatch):
    monkeypatch.setenv("HARA_LLM_PROVIDER", "volcengine-agent-plan")
    monkeypatch.setenv("HARA_LLM_BASE_URL", "https://llm.example/v1")
    monkeypatch.setenv("HARA_LLM_MODEL", "test-model")
    monkeypatch.setenv("HARA_LLM_API_KEY", "secret")
    monkeypatch.setenv("HARA_LLM_GUIDEWORD_THINKING", "auto")

    assert LLMConfig.from_env().guideword_thinking == "auto"


def test_malfunction_thinking_is_configurable_from_environment(monkeypatch):
    monkeypatch.setenv("HARA_LLM_MALFUNCTION_THINKING", "auto")

    assert LLMConfig.from_env().malfunction_thinking == "auto"


def test_scenario_batch_policy_is_configurable_from_environment(monkeypatch):
    monkeypatch.setenv("HARA_LLM_SCENARIO_THINKING", "auto")
    monkeypatch.setenv("HARA_LLM_SCENARIO_BATCH_MAX_CHARS", "9000")
    monkeypatch.setenv("HARA_LLM_SCENARIO_BATCH_MAX_ITEMS", "7")
    monkeypatch.setenv("HARA_LLM_SCENARIO_MAX_SPLIT_DEPTH", "6")

    config = LLMConfig.from_env()
    assert config.scenario_thinking == "auto"
    assert config.scenario_batch_max_chars == 9000
    assert config.scenario_batch_max_items == 7
    assert config.scenario_max_split_depth == 6


def test_validated_artifact_cache_round_trip_contains_no_credentials():
    cache_dir = ROOT / "runtime" / ".test-agent-application" / "artifact-cache"
    cache = ValidatedArtifactCache(cache_dir, "readwrite")
    material = {
        "document_text": "AVP item",
        "prompt_version": "v1",
        "model": "test-model",
    }
    key = cache.key(material)
    payload = {"facts": {"system_description": "AVP"}, "functions": []}

    cache.save(key, payload)
    cache_text = (cache_dir / f"{key}.json").read_text(encoding="utf-8")

    assert cache.load(key) == payload
    assert "secret" not in cache_text


def test_validated_artifact_cache_key_changes_with_prompt_version():
    cache = ValidatedArtifactCache(ROOT / "runtime" / ".test-agent-application", "off")
    first = cache.key({"document_text": "AVP item", "prompt_version": "v1"})
    second = cache.key({"document_text": "AVP item", "prompt_version": "v2"})

    assert first != second


def test_artifact_cache_key_changes_with_extraction_thinking_mode():
    class RecordingCache:
        def __init__(self):
            self.keys = []
            self.delegate = ValidatedArtifactCache(ROOT / "runtime" / "unused", "off")

        def key(self, material):
            key = self.delegate.key(material)
            self.keys.append((key, dict(material)))
            return key

        def load_with_status(self, key):
            return None, "cache_file_missing"

        def save(self, key, payload):
            return None

    class CoreClient:
        def __init__(self, thinking):
            self.config = LLMConfig(
                "volcengine-agent-plan", "https://llm.example/v1", "model", "secret",
                extraction_thinking=thinking,
            )

        def complete_json(self, request):
            return LLMResponse(data=_core_artifact_response(), model="model")

    def key_for(thinking):
        state = HARAState(run_id=f"cache-{thinking}", stage=WorkflowStage.EXTRACT)
        state.item_definition = {
            "text": "same document", "source_id": "ItemDef.docx",
            "blocks": [{"block_id": "B1", "location": "p1", "text": "basic"}],
        }
        cache = RecordingCache()
        client = CoreClient(thinking)
        extract_item_artifacts(
            state, ItemArtifactExtractionAgent(client), ItemSupplementAgent(client),
            ItemEvidenceRouter(), cache=cache,
        )
        return cache.keys[0]

    disabled_key, disabled_material = key_for("disabled")
    enabled_key, enabled_material = key_for("enabled")

    assert disabled_material["extraction_thinking"] == "disabled"
    assert enabled_material["extraction_thinking"] == "enabled"
    assert disabled_key != enabled_key


def test_llm_config_fails_closed_when_credentials_are_missing():
    with pytest.raises(ValueError, match="base_url、model和api_key"):
        create_llm_client(LLMConfig("openai-compatible", "", "", ""))


def test_item_document_reader_preserves_traceable_blocks():
    state = read_item_document(
        HARAState(run_id="document-test"), str(ROOT / "input" / "ItemDef.docx")
    )
    assert state.stage.value == "extract"
    assert state.item_definition["blocks"]
    assert all(block["location"] for block in state.item_definition["blocks"])
    assert state.audit_trail[-1]["character_count"] == len(state.item_definition["text"])


def _core_artifact_response():
    return {
        "item_definition": {
            "system_description": "AVP low-speed parking system",
            "item_boundary": "From parking request to vehicle motion-control output",
            "operating_modes": ["space search", "parking"],
            "odd": {
                "locations": ["parking facility"],
                "road_types": ["parking aisle"],
                "weather_conditions": ["normal"],
                "road_surfaces": ["paved"],
                "speed_range_kph": [0, 5],
            },
            "source_location": "paragraph 1",
            "source_excerpt": "AVP low-speed parking",
            "confidence": 0.9,
            "status": "PENDING",
        },
        "functions": [{
            "function_id": "FUN-1",
            "name": "Provide braking control",
            "output": "Target braking request",
            "description": "Control vehicle deceleration during parking",
            "preconditions": None,
            "triggers": ["deceleration requested"],
            "odd_constraints": ["parking facility"],
            "fallback_behavior": "request stop",
            "consequences": ["vehicle may not decelerate"],
            "source_location": "paragraph 2",
            "source_excerpt": "provide braking control",
            "confidence": 0.9,
            "status": "PENDING",
        }],
    }


def test_item_artifact_agent_uses_one_bounded_full_document_request(monkeypatch):
    monkeypatch.delenv("HARA_ITEM_ARTIFACT_MAX_TOKENS", raising=False)

    class CapturingClient:
        def __init__(self):
            self.requests = []

        def complete_json(self, request):
            self.requests.append(request)
            return LLMResponse(data=_core_artifact_response(), model="fake-model")

    client = CapturingClient()
    facts, functions, _ = ItemArtifactExtractionAgent(client).extract(
        "FULL DOCUMENT SENTINEL", "ItemDef.docx",
    )

    assert len(client.requests) == 1
    assert client.requests[0].task == "extract_core_item_artifacts"
    assert client.requests[0].max_tokens == 16384
    assert "FULL DOCUMENT SENTINEL" in client.requests[0].user_prompt
    assert facts.confidence == 0.9
    assert functions[0].preconditions == []
    assert functions[0].confidence == 0.9


def test_item_artifact_agent_retries_length_once_at_configured_limit(monkeypatch):
    monkeypatch.delenv("HARA_ITEM_ARTIFACT_MAX_TOKENS", raising=False)

    class LengthThenSuccessClient:
        def __init__(self):
            self.config = LLMConfig(
                provider="openai-compatible",
                base_url="https://llm.example/v1",
                model="test-model",
                api_key="secret",
                max_tokens=32768,
            )
            self.requests = []

        def complete_json(self, request):
            self.requests.append(request)
            if len(self.requests) == 1:
                raise LLMOutputLimitError("finish_reason=length")
            return LLMResponse(data=_core_artifact_response(), model="fake-model")

    client = LengthThenSuccessClient()
    _, _, audit = ItemArtifactExtractionAgent(client).extract(
        "FULL DOCUMENT SENTINEL", "ItemDef.docx",
    )

    assert [request.max_tokens for request in client.requests] == [16384, 32768]
    assert audit["output_limit_retry"] is True
    assert audit["max_tokens"] == 32768


def _project_evidence_route():
    return ItemEvidenceRouter().route([{
        "block_id": "B1", "location": "p1", "text": "performance latency 100 ms",
    }], "project_evidence")


def test_item_supplement_retries_output_limit_with_larger_budget(monkeypatch):
    monkeypatch.delenv("HARA_ITEM_SUPPLEMENT_MAX_TOKENS", raising=False)

    class LengthThenSuccessClient:
        def __init__(self):
            self.config = LLMConfig(
                "openai-compatible", "https://llm.example/v1", "model", "secret",
                max_tokens=32768,
            )
            self.requests = []

        def complete_json(self, request):
            self.requests.append(request)
            if len(self.requests) == 1:
                raise LLMOutputLimitError("finish_reason=length")
            return LLMResponse(data={
                "performance_parameters": [], "driver_contexts": [], "exposure_inputs": [],
            }, model="model")

    client = LengthThenSuccessClient()
    _, audit = ItemSupplementAgent(client).extract(_project_evidence_route(), "ItemDef.docx")

    assert [item.max_tokens for item in client.requests] == [4096, 8192]
    assert audit["output_limit_retry"] is True
    assert audit["max_tokens"] == 8192


def test_item_supplement_does_not_retry_at_configured_limit(monkeypatch):
    monkeypatch.delenv("HARA_ITEM_SUPPLEMENT_MAX_TOKENS", raising=False)

    class LimitedClient:
        config = LLMConfig(
            "openai-compatible", "https://llm.example/v1", "model", "secret",
            max_tokens=4096,
        )

        def __init__(self):
            self.calls = 0

        def complete_json(self, request):
            self.calls += 1
            raise LLMOutputLimitError("finish_reason=length")

    client = LimitedClient()
    with pytest.raises(LLMOutputLimitError):
        ItemSupplementAgent(client).extract(_project_evidence_route(), "ItemDef.docx")
    assert client.calls == 1


def test_item_supplement_success_has_no_extra_call_and_bounds_arrays(monkeypatch):
    monkeypatch.delenv("HARA_ITEM_SUPPLEMENT_MAX_TOKENS", raising=False)

    class SuccessfulClient:
        config = LLMConfig(
            "openai-compatible", "https://llm.example/v1", "model", "secret",
        )

        def __init__(self):
            self.calls = 0

        def complete_json(self, request):
            self.calls += 1
            return LLMResponse(data={
                "performance_parameters": [{"id": index} for index in range(20)],
                "driver_contexts": [{"id": index} for index in range(10)],
                "exposure_inputs": [{"id": index} for index in range(15)],
            }, model="model")

    client = SuccessfulClient()
    data, audit = ItemSupplementAgent(client).extract(_project_evidence_route(), "ItemDef.docx")

    assert client.calls == 1
    assert audit["output_limit_retry"] is False
    assert len(data["performance_parameters"]) == 12
    assert len(data["driver_contexts"]) == 8
    assert len(data["exposure_inputs"]) == 12


def test_item_supplement_normalizes_driver_control_without_semantic_guessing(monkeypatch):
    monkeypatch.delenv("HARA_ITEM_SUPPLEMENT_MAX_TOKENS", raising=False)

    class DriverContextClient:
        config = LLMConfig(
            "openai-compatible", "https://llm.example/v1", "model", "secret",
        )

        def complete_json(self, request):
            return LLMResponse(data={
                "performance_parameters": [],
                "driver_contexts": [
                    {
                        "context_id": "DC-01", "driver_position": "inside",
                        "direct_vehicle_control": "true",
                    },
                    {
                        "context_id": "DC-02", "driver_position": "outside",
                        "direct_vehicle_control": "系统接管横控纵控权限",
                    },
                ],
                "exposure_inputs": [],
            }, model="model")

    data, audit = ItemSupplementAgent(DriverContextClient()).extract(
        _project_evidence_route(), "ItemDef.docx",
    )

    assert data["driver_contexts"][0]["direct_vehicle_control"] is True
    assert data["driver_contexts"][1]["direct_vehicle_control"] is None
    assert len(audit["normalization_warnings"]) == 2
    assert "无法安全推断boolean" in audit["normalization_warnings"][1]


def test_item_definition_driver_context_normalization_marks_pending():
    raw = {
        "system_description": "AVP系统",
        "item_boundary": "从请求到控制输出",
        "odd": {"speed_range_kph": [0, 5]},
        "driver_contexts": [{
            "context_id": "DC-01", "driver_position": "inside",
            "direct_vehicle_control": "驾驶员干预导致任务取消",
        }],
        "source_location": "p1", "source_excerpt": "驾驶员可干预",
        "status": "FINALIZED",
    }

    facts, warnings = ItemDefinitionExtractionAgent._parse_with_warnings(raw, "ItemDef.docx")

    assert facts.driver_contexts[0]["direct_vehicle_control"] is None
    assert facts.status is ReviewStatus.PENDING
    assert any("描述性文本" in warning for warning in warnings)


def test_item_evidence_router_sends_only_relevant_blocks_with_neighbors():
    blocks = [
        {"block_id": "B1", "location": "p1", "text": "unrelated introduction"},
        {"block_id": "B2", "location": "p2", "text": "maximum speed is 5 km/h"},
        {"block_id": "B3", "location": "p3", "text": "nearby supporting sentence"},
        {"block_id": "B4", "location": "p4", "text": "unrelated appendix content"},
    ]

    routed = ItemEvidenceRouter().route(blocks, "odd_repair")

    assert routed is not None
    assert routed.block_ids == ["B1", "B2", "B3"]
    assert "unrelated appendix content" not in routed.text


def test_validated_artifact_cache_hit_skips_llm_extraction():
    class CountingClient:
        def __init__(self):
            self.calls = 0

        def complete_json(self, request):
            self.calls += 1
            return LLMResponse(data=_core_artifact_response(), model="fake-model")

    client = CountingClient()
    cache_dir = ROOT / "runtime" / ".test-agent-application" / "validated-artifacts"

    def state():
        value = HARAState(run_id="artifact-cache-test", stage=WorkflowStage.EXTRACT)
        value.item_definition = {
            "text": "artifact-cache-skip-llm-sentinel",
            "source_id": "ItemDef.docx",
            "blocks": [{
                "block_id": "B1", "location": "p1", "text": "basic capability",
            }],
        }
        return value

    agent = ItemArtifactExtractionAgent(client)
    supplement = ItemSupplementAgent(client)
    extract_item_artifacts(
        state(), agent, supplement, ItemEvidenceRouter(),
        cache=ValidatedArtifactCache(cache_dir, "refresh"),
    )
    cached_state = extract_item_artifacts(
        state(), agent, supplement, ItemEvidenceRouter(),
        cache=ValidatedArtifactCache(cache_dir, "readwrite"),
    )

    assert client.calls == 1
    assert cached_state.audit_trail[-1]["cache_hit"] is True
    assert cached_state.audit_trail[-1]["event"] == "extract_timing"
    assert cached_state.audit_trail[-1]["llm_calls"] == 0


def test_core_artifact_cache_survives_failed_supplement(monkeypatch):
    monkeypatch.delenv("HARA_ITEM_SUPPLEMENT_MAX_TOKENS", raising=False)
    config = LLMConfig(
        "openai-compatible", "https://llm.example/v1", "model", "secret",
        max_tokens=8192,
    )

    def state():
        value = HARAState(run_id="partial-cache-test", stage=WorkflowStage.EXTRACT)
        value.item_definition = {
            "text": "core extraction must be reused",
            "source_id": "ItemDef.docx",
            "blocks": [{
                "block_id": "B1", "location": "p1", "text": "performance latency 100 ms",
            }],
        }
        return value

    class FailingSupplementClient:
        def __init__(self):
            self.config = config
            self.tasks = []

        def complete_json(self, request):
            self.tasks.append(request.task)
            if request.task == "extract_core_item_artifacts":
                return LLMResponse(data=_core_artifact_response(), model="model")
            raise LLMOutputLimitError("finish_reason=length")

    cache_dir = ROOT / "runtime" / ".test-agent-application" / "partial-cache"
    cache = ValidatedArtifactCache(cache_dir, "refresh")
    first_client = FailingSupplementClient()
    with pytest.raises(LLMOutputLimitError):
        extract_item_artifacts(
            state(), ItemArtifactExtractionAgent(first_client),
            ItemSupplementAgent(first_client), ItemEvidenceRouter(), cache=cache,
        )
    assert first_client.tasks.count("extract_core_item_artifacts") == 1

    class SuccessfulSupplementClient:
        def __init__(self):
            self.config = config
            self.tasks = []

        def complete_json(self, request):
            self.tasks.append(request.task)
            assert request.task == "supplement_project_evidence"
            return LLMResponse(data={
                "performance_parameters": [], "driver_contexts": [], "exposure_inputs": [],
            }, model="model")

    second_client = SuccessfulSupplementClient()
    extract_item_artifacts(
        state(), ItemArtifactExtractionAgent(second_client),
        ItemSupplementAgent(second_client), ItemEvidenceRouter(),
        cache=ValidatedArtifactCache(cache_dir, "readwrite"),
    )
    assert second_client.tasks == ["supplement_project_evidence"]


def test_extract_timing_proves_supplements_run_in_parallel():
    class TimedClient:
        def complete_json(self, request):
            if request.task == "extract_core_item_artifacts":
                payload = json.loads(json.dumps(_core_artifact_response()))
                payload["item_definition"]["operating_modes"] = []
                payload["item_definition"]["odd"] = {
                    "locations": [], "road_types": [], "weather_conditions": [],
                    "road_surfaces": [], "speed_range_kph": [None, None],
                }
                return LLMResponse(data=payload, model="model")
            time.sleep(0.06)
            if request.task == "supplement_odd_repair":
                return LLMResponse(data={
                    "operating_modes": ["parking"],
                    "odd": {
                        "locations": ["parking lot"], "road_types": ["aisle"],
                        "weather_conditions": ["normal"], "road_surfaces": ["paved"],
                        "speed_range_kph": [0, 5],
                    },
                }, model="model")
            return LLMResponse(data={
                "performance_parameters": [], "driver_contexts": [], "exposure_inputs": [],
            }, model="model")

    def run(max_workers):
        state = HARAState(run_id=f"timing-{max_workers}", stage=WorkflowStage.EXTRACT)
        state.item_definition = {
            "text": "speed and performance evidence",
            "source_id": "ItemDef.docx",
            "blocks": [{
                "block_id": "B1", "location": "p1",
                "text": "speed limit 5 km/h and performance latency 100 ms",
            }],
        }
        client = TimedClient()
        result = extract_item_artifacts(
            state, ItemArtifactExtractionAgent(client), ItemSupplementAgent(client),
            ItemEvidenceRouter(), max_workers=max_workers,
        )
        return result.audit_trail[-1]

    serial = run(1)
    parallel = run(2)
    serial_tasks = sum(serial["supplement_task_elapsed_seconds"].values())
    parallel_tasks = sum(parallel["supplement_task_elapsed_seconds"].values())

    assert serial["supplement_parallel_wall_seconds"] >= serial_tasks * 0.9
    assert parallel["supplement_parallel_wall_seconds"] < parallel_tasks * 0.75
    assert parallel["supplement_parallel_wall_seconds"] < serial["supplement_parallel_wall_seconds"]


def test_guideword_agent_covers_all_but_only_marks_applicable_candidates():
    class GuidewordClient:
        def complete_json(self, request):
            assert '顶层必须且只能使用assessments字段' in request.user_prompt
            assert '{"assessments":[...]}' in request.user_prompt
            assert "confidence必须为0.0到1.0范围内的JSON number" in request.user_prompt
            assert "禁止使用high、medium、low" in request.user_prompt
            return LLMResponse(data={"assessments": [
                {"guideword": "loss", "applicable": True, "rationale": "制动请求可能完全丢失", "confidence": 0.9, "status": "FINALIZED"},
                {"guideword": "too large", "applicable": True, "rationale": "扭矩请求可能超过目标值", "confidence": 0.9, "status": "FINALIZED"},
                {"guideword": "too long", "applicable": False, "rationale": "持续时间偏差已由持续控制需求限定", "confidence": 0.7, "status": "FINALIZED"},
            ]}, model="fake-model")

    function = FunctionDefinition(
        "FUN-1", "输出制动扭矩", "目标制动扭矩请求",
        sources=[SourceRef("item_definition", "ItemDef.docx", "p3", "输出制动扭矩")],
    )
    assessments, audit = GuidewordApplicabilityAgent(GuidewordClient()).assess(
        function, ["loss", "too large", "too long"]
    )
    assert len(assessments) == 3
    assert sum(item.applicable for item in assessments) == 2
    assert audit["coverage_count"] == 3
    assert audit["applicable_count"] == 2


def test_guideword_agent_rejects_incomplete_coverage():
    class IncompleteClient:
        def complete_json(self, request):
            return LLMResponse(data={"assessments": [
                {"guideword": "loss", "applicable": True, "rationale": "可能丢失", "confidence": 0.8},
            ]}, model="fake-model")

    function = FunctionDefinition(
        "FUN-1", "输出制动扭矩", "目标制动扭矩请求",
        sources=[SourceRef("item_definition", "ItemDef.docx", "p3", "输出制动扭矩")],
    )
    with pytest.raises(ValueError, match="覆盖不完整"):
        GuidewordApplicabilityAgent(IncompleteClient()).assess(function, ["loss", "too large"])


def test_guideword_agent_rejects_missing_assessments_array():
    class MissingArrayClient:
        def complete_json(self, request):
            return LLMResponse(data={"guideword": "loss"}, model="fake-model")

    function = FunctionDefinition(
        "FUN-1", "输出制动扭矩", "目标制动扭矩请求",
        sources=[SourceRef("item_definition", "ItemDef.docx", "p3", "制动扭矩")],
    )
    with pytest.raises(ValueError, match="缺少assessments数组"):
        GuidewordApplicabilityAgent(MissingArrayClient()).assess(function, ["loss"])


def test_guideword_agent_rejects_duplicate_guideword_assessments():
    class DuplicateClient:
        def complete_json(self, request):
            return LLMResponse(data={"assessments": [
                {"guideword": "loss", "applicable": True, "rationale": "丢失"},
                {"guideword": "loss", "applicable": False, "rationale": "重复"},
            ]}, model="fake-model")

    function = FunctionDefinition(
        "FUN-1", "输出制动扭矩", "目标制动扭矩请求",
        sources=[SourceRef("item_definition", "ItemDef.docx", "p3", "制动扭矩")],
    )
    with pytest.raises(ValueError, match="重复项"):
        GuidewordApplicabilityAgent(DuplicateClient()).assess(
            function, ["loss", "too large"],
        )


def test_guideword_domain_contract_finalized_requires_rationale_but_pending_does_not():
    finalized = GuidewordAssessment(
        "FUN-1", "loss", True, "request may be lost",
        status=ReviewStatus.FINALIZED, confidence=0.8,
    )
    pending = GuidewordAssessment(
        "FUN-1", "too late", True, "",
        status=ReviewStatus.PENDING, confidence=0.8,
    )

    assert finalized.status is ReviewStatus.FINALIZED
    assert pending.rationale == ""
    with pytest.raises(ValueError, match="FINALIZED GuidewordAssessment缺少rationale"):
        GuidewordAssessment(
            "FUN-1", "loss", True, " ",
            status=ReviewStatus.FINALIZED, confidence=0.8,
        )


@pytest.mark.parametrize(("kwargs", "message"), [
    ({"guideword": "", "applicable": True}, "缺少guideword"),
    ({"guideword": "loss", "applicable": None}, "applicable必须为boolean"),
    ({"guideword": "loss", "applicable": "yes"}, "applicable必须为boolean"),
])
def test_guideword_domain_contract_rejects_missing_identity_or_boolean(kwargs, message):
    with pytest.raises(ValueError, match=message):
        GuidewordAssessment(
            function_id="FUN-1", rationale="reason", status=ReviewStatus.PENDING,
            confidence=0.8, **kwargs,
        )


def test_guideword_agent_degrades_only_missing_rationale_and_reports_coverage(capsys):
    class Client:
        def complete_json(self, request):
            return LLMResponse(data={"assessments": [
                {"guideword": "loss", "applicable": True, "rationale": "lost",
                 "confidence": 0.8, "status": "FINALIZED"},
                {"guideword": "too late", "applicable": True,
                 "confidence": 0.7, "status": "FINALIZED"},
            ]}, model="fake-model")

    function = FunctionDefinition("FUN-1", "brake", "torque request")
    assessments, audit = GuidewordApplicabilityAgent(Client()).assess(
        function, ["loss", "too late"],
    )

    assert assessments[0].status is ReviewStatus.FINALIZED
    assert assessments[1].status is ReviewStatus.PENDING
    assert assessments[1].applicable is True
    assert assessments[1].rationale == ""
    assert audit["coverage_count"] == 2
    assert audit["finalized_count"] == 1
    assert audit["pending_count"] == 1
    assert audit["complete_count"] == 1
    assert audit["incomplete_count"] == 1
    assert audit["parse_error_count"] == 1
    assert audit["parse_errors"][0]["field"] == "rationale"
    logs = capsys.readouterr().err
    assert "guideword item validation failed" in logs
    assert "index=1" in logs
    assert "missing_fields=['rationale']" in logs
    assert "complete=1 incomplete=1 review_pending=1 review_finalized=1 parse_errors=1" in logs


@pytest.mark.parametrize("item", [
    {"applicable": True, "rationale": "reason", "confidence": 0.8, "status": "PENDING"},
    {"guideword": "loss", "rationale": "reason", "confidence": 0.8, "status": "PENDING"},
    {"guideword": "loss", "applicable": "yes", "rationale": "reason",
     "confidence": 0.8, "status": "PENDING"},
])
def test_guideword_agent_does_not_degrade_missing_identity_or_applicable(item):
    class Client:
        def complete_json(self, request):
            return LLMResponse(data={"assessments": [item]}, model="fake-model")

    function = FunctionDefinition("FUN-1", "brake", "torque request")
    with pytest.raises(ValueError):
        GuidewordApplicabilityAgent(Client()).assess(function, ["loss"])


def test_guideword_agent_rejects_unknown_guideword_instead_of_fuzzy_matching():
    class Client:
        def complete_json(self, request):
            return LLMResponse(data={"assessments": [{
                "guideword": "late timing", "applicable": True, "rationale": "reason",
                "confidence": 0.8, "status": "FINALIZED",
            }]}, model="fake-model")

    function = FunctionDefinition("FUN-1", "brake", "torque request")
    with pytest.raises(ValueError, match=r"missing=\['too late'\].*extra=\['late timing'\]"):
        GuidewordApplicabilityAgent(Client()).assess(function, ["too late"])


@pytest.mark.parametrize(("status", "applicable", "rationale", "expected_calls", "skip_reason"), [
    (ReviewStatus.PENDING, True, "reason", 1, None),
    (ReviewStatus.FINALIZED, True, "reason", 1, None),
    (ReviewStatus.PENDING, True, "", 0, "no_complete_applicable_guidewords"),
    (ReviewStatus.PENDING, False, "reason", 0, "no_applicable_guidewords"),
    (ReviewStatus.FINALIZED, False, "reason", 0, "no_applicable_guidewords"),
])
def test_semantic_completeness_and_applicability_control_malfunction_generation(
    status, applicable, rationale, expected_calls, skip_reason,
):
    class Client:
        def __init__(self):
            self.calls = 0

        def complete_json(self, request):
            self.calls += 1
            return LLMResponse(data={"candidates": [{
                "malfunction_id": "MF-1", "guideword": "loss", "description": "lost",
                "functional_effect": "braking unavailable",
                "vehicle_level_hazard": "vehicle continues moving",
                "causal_chain": ["request lost", "vehicle not decelerating"],
                "confidence": 0.8,
            }]}, model="fake-model")

    client = Client()
    function = FunctionDefinition(
        "FUN-1", "brake", "torque request",
        sources=[SourceRef("item_definition", "item.docx", "p1", "brake")],
    )
    assessment = GuidewordAssessment(
        "FUN-1", "loss", applicable, rationale, status=status, confidence=0.8,
    )
    candidates, audit = MalfunctionHazardAgent(client).generate(function, [assessment])

    assert client.calls == expected_calls
    assert len(candidates) == expected_calls
    if expected_calls == 0:
        assert audit["skip_reason"] == skip_reason


def test_malfunction_agent_only_uses_applicable_guidewords_and_builds_causal_chain():
    class MalfunctionClient:
        def complete_json(self, request):
            return LLMResponse(data={"candidates": [{
                "malfunction_id": "MF-1", "guideword": "loss",
                "description": "目标制动扭矩请求丢失", "functional_effect": "车辆无法按需建立制动力",
                "vehicle_level_hazard": "车辆在泊车轨迹中继续运动并接近障碍物",
                "causal_chain": ["制动请求丢失", "车辆未减速", "车辆接近障碍物"],
                "source_location": "paragraph 3", "source_excerpt": "输出目标制动扭矩",
                "confidence": 0.88,
            }]}, model="fake-model")

    function = FunctionDefinition(
        "FUN-1", "输出制动扭矩", "目标制动扭矩请求",
        sources=[SourceRef("item_definition", "ItemDef.docx", "p3", "输出制动扭矩")],
    )
    assessments = [
        GuidewordAssessment("FUN-1", "loss", True, "请求可能丢失", status=ReviewStatus.FINALIZED),
        GuidewordAssessment("FUN-1", "too long", False, "不适用", status=ReviewStatus.FINALIZED),
    ]
    candidates, audit = MalfunctionHazardAgent(MalfunctionClient()).generate(function, assessments)
    assert len(candidates) == 1
    assert candidates[0].guideword == "loss"
    assert len(candidates[0].causal_chain) == 3
    assert audit["applicable_guidewords"] == 1


def test_guideword_sources_are_propagated_from_function_and_llm_source_is_ignored():
    trusted = SourceRef("item_definition", "ItemDef.docx", "p3", "trusted")
    function = FunctionDefinition(
        "FUN-1", "输出制动扭矩", "目标制动扭矩请求", sources=[trusted],
    )
    assessment = GuidewordApplicabilityAgent._parse(function, {
        "guideword": "loss", "applicable": True, "rationale": "请求可能丢失",
        "source_location": "fabricated p99", "source_excerpt": "fabricated",
        "confidence": 0.8,
    })

    assert assessment.sources == [trusted]
    assert assessment.sources[0].location == "p3"


def test_malfunction_sources_prefer_guideword_and_are_stably_deduplicated():
    function_source = SourceRef("item_definition", "ItemDef.docx", "p2", "function")
    guideword_source = SourceRef("item_definition", "ItemDef.docx", "p3", "guideword")
    function = FunctionDefinition(
        "FUN-1", "输出制动扭矩", "目标制动扭矩请求", sources=[function_source],
    )
    assessment = GuidewordAssessment(
        "FUN-1", "loss", True, "请求可能丢失",
        sources=[guideword_source, guideword_source],
    )
    candidate = MalfunctionHazardAgent._parse(function, {
        "malfunction_id": "MF-1", "guideword": "loss", "description": "请求丢失",
        "functional_effect": "车辆无法建立制动力",
        "vehicle_level_hazard": "车辆继续运动并接近障碍物",
        "causal_chain": ["请求丢失", "车辆未减速"],
        "source_location": "fabricated p99", "source_excerpt": "fabricated",
        "confidence": 0.8,
    }, assessment)

    assert candidate.sources == [guideword_source]
    assert candidate.sources[0].location == "p3"
    assert candidate.vehicle_level_hazard == "车辆继续运动并接近障碍物"


def test_malfunction_sources_fall_back_to_function_when_guideword_has_none():
    trusted = SourceRef("item_definition", "ItemDef.docx", "p4", "function")
    function = FunctionDefinition(
        "FUN-1", "输出制动扭矩", "目标制动扭矩请求", sources=[trusted],
    )
    assessment = GuidewordAssessment(
        "FUN-1", "loss", True, "请求可能丢失", status=ReviewStatus.FINALIZED,
    )
    candidate = MalfunctionHazardAgent._parse(function, {
        "malfunction_id": "MF-1", "guideword": "loss", "description": "请求丢失",
        "functional_effect": "车辆无法建立制动力",
        "vehicle_level_hazard": "车辆继续运动并接近障碍物",
        "causal_chain": ["请求丢失", "车辆未减速"], "confidence": 0.8,
    }, assessment)

    assert candidate.sources == [trusted]


def test_malfunction_agent_fails_closed_when_no_upstream_source_exists():
    class SourceFabricatingClient:
        def complete_json(self, request):
            return LLMResponse(data={"candidates": [{
                "malfunction_id": "MF-1", "guideword": "loss", "description": "请求丢失",
                "functional_effect": "车辆无法建立制动力",
                "vehicle_level_hazard": "车辆继续运动并接近障碍物",
                "causal_chain": ["请求丢失", "车辆未减速"],
                "source_location": "fabricated p99", "source_excerpt": "fabricated",
                "confidence": 0.8,
            }]}, model="fake-model")

    function = FunctionDefinition("FUN-1", "输出制动扭矩", "目标制动扭矩请求")
    assessment = GuidewordAssessment(
        "FUN-1", "loss", True, "请求可能丢失", status=ReviewStatus.FINALIZED,
    )

    with pytest.raises(ValueError, match="Malfunction缺少来源位置"):
        MalfunctionHazardAgent(SourceFabricatingClient()).generate(function, [assessment])


def test_malfunction_agent_rejects_injury_as_vehicle_level_hazard():
    class InvalidHazardClient:
        def complete_json(self, request):
            return LLMResponse(data={"candidates": [{
                "malfunction_id": "MF-X", "guideword": "loss", "description": "制动丢失",
                "functional_effect": "车辆继续运动", "vehicle_level_hazard": "导致行人死亡",
                "causal_chain": ["制动丢失", "车辆继续运动"], "source_location": "p3",
                "confidence": 0.8,
            }]}, model="fake-model")

    function = FunctionDefinition("FUN-1", "输出制动扭矩", "目标制动扭矩请求")
    assessment = GuidewordAssessment(
        "FUN-1", "loss", True, "请求可能丢失", status=ReviewStatus.FINALIZED,
    )
    with pytest.raises(ValueError, match="Hazard不得直接写伤害"):
        MalfunctionHazardAgent(InvalidHazardClient()).generate(function, [assessment])


def _batching_malfunction():
    return MalfunctionCandidate(
        "MF-BATCH", "FUN-1", "loss", "brake request lost",
        "vehicle cannot decelerate", "vehicle continues toward obstacle",
        ["request lost", "vehicle not decelerating"], confidence=0.8,
    )


def _batching_scenarios(count=4, padding=80):
    return [
        ScenarioCandidate(
            f"SCN-{index}", "Parking", f"scenario {index}", f"detail {index}",
            {"mode": "parking", "context": "x" * padding, "index": index},
        )
        for index in range(count)
    ]


def test_scenario_batch_builder_is_budgeted_complete_unique_and_ordered():
    malfunction = _batching_malfunction()
    scenarios = _batching_scenarios(5, padding=100)
    single_chars = estimate_scenario_prompt_chars(
        malfunction, scenarios[:1], system_prompt=ScenarioFeasibilityAgent.SYSTEM_PROMPT,
    )
    batches = build_scenario_batches(
        malfunction, scenarios, max_chars=single_chars + 10, max_items=3,
        system_prompt=ScenarioFeasibilityAgent.SYSTEM_PROMPT,
    )

    flattened = [item.scenario_id for batch in batches for item in batch]
    assert len(batches) == len(scenarios)
    assert flattened == [item.scenario_id for item in scenarios]
    assert len(flattened) == len(set(flattened))


def test_scenario_batch_builder_keeps_small_inputs_together_and_honors_item_cap():
    malfunction = _batching_malfunction()
    scenarios = _batching_scenarios(4, padding=5)

    assert len(build_scenario_batches(
        malfunction, scenarios, max_chars=50000, max_items=10,
        system_prompt=ScenarioFeasibilityAgent.SYSTEM_PROMPT,
    )) == 1
    capped = build_scenario_batches(
        malfunction, scenarios, max_chars=50000, max_items=2,
        system_prompt=ScenarioFeasibilityAgent.SYSTEM_PROMPT,
    )
    assert [len(batch) for batch in capped] == [2, 2]


def test_single_oversize_scenario_fails_before_provider_call():
    class NeverCalledClient:
        def __init__(self):
            self.calls = 0

        def complete_json(self, request):
            self.calls += 1
            raise AssertionError("provider must not be called")

    client = NeverCalledClient()
    agent = ScenarioFeasibilityAgent(client, batch_max_chars=1000, batch_max_items=12)
    with pytest.raises(
        ScenarioBatchSizeError,
        match="malfunction_id=MF-BATCH.*scenario_id=SCN-0.*max_chars=1000",
    ):
        agent.assess(_batching_malfunction(), _batching_scenarios(1, padding=3000))
    assert client.calls == 0


def _scenario_response_for_prompt(request):
    payload_text = request.user_prompt.split("Scenarios=", 1)[1].split("\n这里只评估", 1)[0]
    payload = json.loads(payload_text)
    return [{
        "scenario_id": item["scenario_id"],
        "physically_feasible": True,
        "functionally_relevant": True,
        "causally_relevant": True,
        "breakpoint": "NONE",
        "causal_chain": {
            hop: {
                "claim": f"supported {hop}",
                "basis_type": "DIRECT_FACT",
                "evidence_refs": ["MF.description"],
            }
            for hop in ("m_to_b", "b_to_i", "i_to_h", "h_to_harm")
        },
        "risk_dimension_changes": [{
            "dimension": "collision_object",
            "evidence_refs": ["SCN.object_type"] if "SCN.object_type" in item["fact_registry"] else ["MF.description"],
            "reason": "explicit input changes downstream review context",
        }],
        "rationale": "risk context differs",
        "hazardous_event": f"hazard {item['scenario_id']}",
        "potential_harm": f"harm {item['scenario_id']}",
        "confidence": 0.8,
        "status": "PENDING",
    } for item in payload]


def _v9_fixture(item):
    """Upgrade an intentionally concise test result to the production v9 envelope."""
    result = dict(item)
    dimensions = result.pop("risk_dimensions_changed", [])
    causal = result["causally_relevant"]
    result["breakpoint"] = "NONE" if causal else "I_TO_H"
    result["causal_chain"] = {
        "m_to_b": {"claim": "supported", "basis_type": "DIRECT_FACT", "evidence_refs": ["MF.description"]},
        "b_to_i": {"claim": "supported", "basis_type": "DIRECT_FACT", "evidence_refs": ["MF.functional_effect"]},
        "i_to_h": {"claim": "supported" if causal else "unsupported", "basis_type": "DIRECT_FACT" if causal else "ASSUMPTION", "evidence_refs": ["MF.vehicle_level_hazard"] if causal else []},
    }
    if causal:
        result["causal_chain"]["h_to_harm"] = {
            "claim": "supported", "basis_type": "DIRECT_FACT",
            "evidence_refs": ["MF.vehicle_level_hazard"],
        }
    result["risk_dimension_changes"] = [
        {"dimension": dimension, "evidence_refs": ["MF.description"], "reason": "fixture"}
        for dimension in dimensions
    ]
    return result


@pytest.mark.parametrize(("mutation", "message"), [
    ("missing", "覆盖不完整"),
    ("duplicate", "重复项"),
    ("unknown", "未知scenario_id"),
])
def test_scenario_batch_coverage_rejects_missing_duplicate_and_unknown(mutation, message):
    class Client:
        def complete_json(self, request):
            assessments = _scenario_response_for_prompt(request)
            if mutation == "missing":
                assessments.pop()
            elif mutation == "duplicate":
                assessments.append(dict(assessments[-1]))
            else:
                assessments[-1]["scenario_id"] = "SCN-UNKNOWN"
            return LLMResponse(data={"assessments": assessments}, model="fake-model")

    with pytest.raises(ValueError, match=message):
        ScenarioFeasibilityAgent(
            Client(), batch_max_chars=50000, batch_max_items=12,
        ).assess(_batching_malfunction(), _batching_scenarios(3, padding=5))


def test_scenario_batching_preserves_domain_results_and_global_coverage():
    class Client:
        def __init__(self):
            self.calls = 0

        def complete_json(self, request):
            self.calls += 1
            return LLMResponse(
                data={"assessments": _scenario_response_for_prompt(request)},
                model="fake-model", usage={},
            )

    malfunction = _batching_malfunction()
    scenarios = _batching_scenarios(4, padding=80)
    single_chars = estimate_scenario_prompt_chars(
        malfunction, scenarios[:1], system_prompt=(
            ScenarioFeasibilityAgent.SYSTEM_PROMPT + "\n"
            + ScenarioFeasibilityAgent.MACHINE_OUTPUT_RULE
        ),
    )
    single_client, batched_client = Client(), Client()
    single, single_audit = ScenarioFeasibilityAgent(
        single_client, batch_max_chars=50000, batch_max_items=12,
    ).assess(malfunction, scenarios)
    batched, batched_audit = ScenarioFeasibilityAgent(
        batched_client, batch_max_chars=single_chars + 10, batch_max_items=12,
    ).assess(malfunction, scenarios)

    fields = lambda item: (
        item.scenario_id, item.retain, item.risk_dimensions_changed,
        item.hazardous_event, item.potential_harm, item.confidence,
    )
    assert [fields(item) for item in batched] == [fields(item) for item in single]
    assert [item.scenario_id for item in batched] == [item.scenario_id for item in scenarios]
    assert single_audit["batch_count"] == 1
    assert batched_audit["batch_count"] == len(scenarios)
    assert batched_audit["llm_calls"] == len(scenarios)


@pytest.mark.parametrize(("limit", "expected_splits", "expected_leaves", "expected_calls"), [
    (4, 1, 2, 3),
    (2, 3, 4, 7),
])
def test_scenario_output_limit_adaptively_splits_and_preserves_coverage(
    limit, expected_splits, expected_leaves, expected_calls,
):
    class Client:
        def __init__(self):
            self.requests = []

        def complete_json(self, request):
            assessments = _scenario_response_for_prompt(request)
            self.requests.append(dict(request.metadata))
            if len(assessments) > limit:
                raise LLMOutputLimitError("length", diagnostics={
                    "reasoning_characters": 123,
                    "completion_tokens": 8192,
                })
            return LLMResponse(data={"assessments": assessments}, model="fake-model")

    client = Client()
    scenarios = _batching_scenarios(8, padding=5)
    assessments, audit = ScenarioFeasibilityAgent(
        client, batch_max_chars=50000, batch_max_items=12,
    ).assess(_batching_malfunction(), scenarios)

    assert [item.scenario_id for item in assessments] == [item.scenario_id for item in scenarios]
    assert len({item.scenario_id for item in assessments}) == 8
    assert audit["adaptive_split_count"] == expected_splits
    assert audit["leaf_batch_count"] == expected_leaves
    assert audit["actual_llm_calls"] == expected_calls
    assert audit["output_limit_count"] == expected_splits
    assert client.requests[0]["split_path"] == "root"
    assert client.requests[0]["scenario_count"] == 8


def test_scenario_timeout_adaptively_splits():
    class Client:
        def __init__(self):
            self.calls = 0

        def complete_json(self, request):
            self.calls += 1
            assessments = _scenario_response_for_prompt(request)
            if len(assessments) > 2:
                raise LLMTimeoutError("timeout", attempts=2)
            return LLMResponse(data={"assessments": assessments}, model="fake-model")

    client = Client()
    assessments, audit = ScenarioFeasibilityAgent(
        client, batch_max_chars=50000, batch_max_items=12,
    ).assess(_batching_malfunction(), _batching_scenarios(4, padding=5))

    assert len(assessments) == 4
    assert audit["adaptive_split_count"] == 1
    assert audit["timeout_count"] == 1
    assert audit["retry_calls"] == 1
    assert client.calls == 3


def test_scenario_boolean_risk_dimensions_fails_with_typed_contract_error():
    class Client:
        def complete_json(self, request):
            item = _scenario_response_for_prompt(request)[0]
            item["risk_dimensions_changed"] = True
            return LLMResponse(data={"assessments": [item]}, model="fake-model")

    with pytest.raises(ScenarioSchemaContractError, match="risk_dimensions_changed"):
        ScenarioFeasibilityAgent(Client()).assess(
            _batching_malfunction(), _batching_scenarios(1, padding=5),
        )


def _scenario_dimensions_result(dimensions):
    class Client:
        def complete_json(self, request):
            item = _scenario_response_for_prompt(request)[0]
            if isinstance(dimensions, list):
                item["risk_dimension_changes"] = [
                    {"dimension": value, "evidence_refs": ["MF.description"], "reason": "fixture"}
                    for value in dimensions
                ]
            else:
                item["risk_dimension_changes"] = dimensions
            if dimensions == []:
                item.update({
                    "causally_relevant": False,
                    "breakpoint": "I_TO_H",
                    "causal_chain": {
                        "m_to_b": {"claim": "supported", "basis_type": "DIRECT_FACT", "evidence_refs": ["MF.description"]},
                        "b_to_i": {"claim": "supported", "basis_type": "DIRECT_FACT", "evidence_refs": ["MF.functional_effect"]},
                        "i_to_h": {"claim": "unsupported", "basis_type": "ASSUMPTION", "evidence_refs": []},
                    },
                    "hazardous_event": "",
                    "potential_harm": "",
                })
            return LLMResponse(data={"assessments": [item]}, model="fake-model")

    return ScenarioFeasibilityAgent(Client()).assess(
        _batching_malfunction(), _batching_scenarios(1, padding=5),
    )[0][0]


def test_scenario_prompt_and_validator_share_canonical_risk_dimensions():
    prompt = build_scenario_user_prompt(
        _batching_malfunction(), _batching_scenarios(1, padding=5),
    )
    assert json.dumps(list(RISK_DIMENSION_VALUES), ensure_ascii=False) in prompt
    assert ScenarioFeasibilityAgent.ALLOWED_RISK_DIMENSIONS == frozenset(RISK_DIMENSION_VALUES)
    assert ScenarioFeasibilityAgent.PROMPT_VERSION == "scenario-feasibility-v9"


def test_scenario_canonical_risk_dimensions_accept_valid_empty_and_multiple():
    single = _scenario_dimensions_result([RISK_DIMENSION_VALUES[0]])
    assert single.risk_dimensions_changed == [RISK_DIMENSION_VALUES[0]]
    multiple = _scenario_dimensions_result(list(RISK_DIMENSION_VALUES[:2]))
    assert multiple.risk_dimensions_changed == list(RISK_DIMENSION_VALUES[:2])
    empty = _scenario_dimensions_result([])
    assert empty.risk_dimensions_changed == []
    assert not empty.retain


@pytest.mark.parametrize("dimensions", [
    "S", None, {}, [RISK_DIMENSION_VALUES[0], 3], ["__UNKNOWN_DIMENSION__"],
])
def test_scenario_risk_dimensions_reject_invalid_types_and_values(dimensions):
    with pytest.raises(ValueError, match="risk_dimension_changes|dimension"):
        _scenario_dimensions_result(dimensions)


def test_scenario_real_abbreviation_drift_fails_closed_when_not_canonical():
    if "S" in RISK_DIMENSION_VALUES or "C" in RISK_DIMENSION_VALUES:
        pytest.skip("S/C are canonical in this contract")
    with pytest.raises(
        ValueError,
        match="dimension",
    ):
        _scenario_dimensions_result(["S", "C"])


def test_scenario_duplicate_risk_dimension_fails_closed():
    value = RISK_DIMENSION_VALUES[0]
    with pytest.raises(ValueError, match="duplicate dimension"):
        _scenario_dimensions_result([value, value])


def _cross_field_item(causal, dimensions, hazard, harm):
    item = {
        "scenario_id": "SCN-0", "physically_feasible": True,
        "functionally_relevant": True, "causally_relevant": causal,
        "breakpoint": "NONE" if causal else "I_TO_H",
        "causal_chain": {
            "m_to_b": {"claim": "supported", "basis_type": "DIRECT_FACT", "evidence_refs": ["MF.description"]},
            "b_to_i": {"claim": "supported", "basis_type": "DIRECT_FACT", "evidence_refs": ["MF.functional_effect"]},
            "i_to_h": {"claim": "supported" if causal else "unsupported", "basis_type": "DIRECT_FACT" if causal else "ASSUMPTION", "evidence_refs": ["MF.vehicle_level_hazard"] if causal else []},
        },
        "risk_dimension_changes": [
            {"dimension": dimension, "evidence_refs": ["MF.description"], "reason": "fixture"}
            for dimension in dimensions
        ], "rationale": "contract fixture",
        "hazardous_event": hazard, "potential_harm": harm,
        "confidence": 0.8, "status": "PENDING",
    }
    if causal:
        item["causal_chain"]["h_to_harm"] = {
            "claim": "supported", "basis_type": "DIRECT_FACT",
            "evidence_refs": ["MF.vehicle_level_hazard"],
        }
    return item


@pytest.mark.parametrize(("item", "field"), [
    (_cross_field_item(False, ["distance"], "", ""), "risk_dimension_changes"),
    (_cross_field_item(False, [], "unexpected hazard", ""), "hazardous_event"),
    (_cross_field_item(False, [], "", "unexpected harm"), "potential_harm"),
    (_cross_field_item(True, [], "hazard", "harm"), "risk_dimension_changes"),
    (_cross_field_item(True, ["distance"], "", "harm"), "hazardous_event"),
    (_cross_field_item(True, ["distance"], "hazard", ""), "potential_harm"),
])
def test_scenario_cross_field_invariants_fail_closed_with_execution_context(item, field):
    expected_error = ValueError if field == "risk_dimension_changes" else ScenarioSchemaContractError
    with pytest.raises(expected_error) as error:
        ScenarioFeasibilityAgent._parse(
            _batching_malfunction(), item,
            scenario=_batching_scenarios(1, padding=5)[0],
            batch="2/4", split_path="LR", split_depth=2,
        )
    message = str(error.value)
    expected_marker = f"hop={field}" if field == "risk_dimension_changes" else f"field={field}"
    assert expected_marker in message
    assert "malfunction_id=MF-BATCH" in message
    assert "scenario_id=SCN-0" in message
    assert "batch=2/4 split_path=LR split_depth=2" in message


def test_scenario_cross_field_invariant_accepts_complete_negative_and_positive():
    negative = ScenarioFeasibilityAgent._parse(
        _batching_malfunction(), _cross_field_item(False, [], "", ""),
        scenario=_batching_scenarios(1, padding=5)[0],
    )
    positive = ScenarioFeasibilityAgent._parse(
        _batching_malfunction(),
        _cross_field_item(True, ["distance"], "hazard", "harm"),
        scenario=_batching_scenarios(1, padding=5)[0],
    )
    assert not negative.causally_relevant
    assert positive.causally_relevant


def test_scenario_v8_prompt_enforces_fact_precedence_and_batch_independence():
    prompt = ScenarioFeasibilityAgent.SYSTEM_PROMPT
    assert "Scenario labels and names are descriptive only" in prompt
    assert "MUST NOT override structured facts" in prompt
    assert "Do NOT compare scenarios with one another" in prompt
    assert "Do NOT use another Scenario in the same batch as evidence" in prompt
    assert "severity ladder" in prompt
    assert "does not make ego_speed_kph zero" in prompt


def _evidence_context():
    malfunction = _batching_malfunction()
    scenario = ScenarioCandidate(
        "SCN-EVIDENCE", "Parking", "atomic", "atomic",
        {"relative_distance": "0.5 m", "relative_speed_kph": 5.0,
         "object_type": "pedestrian"},
        semantic_fingerprint="fixture-fingerprint",
        scenario_contract_version=SCENARIO_CONTRACT_VERSION,
    )
    return malfunction, scenario, build_fact_registry(malfunction, scenario)


def _validate_fixture(item, registry=None):
    malfunction, scenario, actual_registry = _evidence_context()
    return validate_evidence_contract(
        malfunction=malfunction, scenario=scenario, item=item,
        registry=registry or actual_registry,
        prompt_version="scenario-feasibility-v9", batch="1/1",
        split_path="root", split_depth=0,
    )


def test_fact_registry_resolves_valid_ref_rejects_unknown_and_is_pair_local():
    _, _, registry = _evidence_context()
    assert registry.resolve("SCN.relative_distance")
    assert registry.resolve("SCN.driver_panic") is None
    assert registry.resolve("SCN-OTHER.relative_distance") is None
    assert registry.resolve("DERIVED.ttc_s")["derivation_type"] == "TTC"


def test_evidence_contract_rejects_unknown_ref_and_assumption_positive():
    item = _v9_fixture({
        "scenario_id": "SCN-EVIDENCE", "physically_feasible": True,
        "functionally_relevant": True, "causally_relevant": True,
        "risk_dimensions_changed": ["distance"], "rationale": "fixture",
        "hazardous_event": "hazard", "potential_harm": "harm", "confidence": 0.8,
    })
    item["causal_chain"]["i_to_h"]["evidence_refs"] = ["SCN.driver_panic"]
    with pytest.raises(ScenarioEvidenceContractError, match="SCN.driver_panic"):
        _validate_fixture(item)
    item["causal_chain"]["i_to_h"] = {
        "claim": "unsupported human response", "basis_type": "ASSUMPTION", "evidence_refs": [],
    }
    with pytest.raises(ScenarioEvidenceContractError, match="cannot depend on assumption"):
        _validate_fixture(item)


def test_derived_physics_is_bounded_to_registered_derivation():
    item = _v9_fixture({
        "scenario_id": "SCN-EVIDENCE", "physically_feasible": True,
        "functionally_relevant": True, "causally_relevant": True,
        "risk_dimensions_changed": ["distance"], "rationale": "fixture",
        "hazardous_event": "hazard", "potential_harm": "harm", "confidence": 0.8,
    })
    item["causal_chain"]["i_to_h"] = {
        "claim": "bounded time-to-collision relation",
        "basis_type": "DERIVED_PHYSICS", "evidence_refs": ["DERIVED.ttc_s"],
    }
    assert _validate_fixture(item)[0] == ["distance"]
    item["causal_chain"]["i_to_h"]["evidence_refs"] = ["SCN.relative_speed_kph"]
    with pytest.raises(ScenarioEvidenceContractError, match="bounded derivation"):
        _validate_fixture(item)


def test_approved_rule_requires_registry_approved_rule_kind():
    item = _v9_fixture({
        "scenario_id": "SCN-EVIDENCE", "physically_feasible": True,
        "functionally_relevant": True, "causally_relevant": True,
        "risk_dimensions_changed": ["distance"], "rationale": "fixture",
        "hazardous_event": "hazard", "potential_harm": "harm", "confidence": 0.8,
    })
    item["causal_chain"]["i_to_h"] = {
        "claim": "approved mechanism", "basis_type": "APPROVED_RULE",
        "evidence_refs": ["DOMAIN_RULE.RULE-1"],
    }
    _, _, base = _evidence_context()
    facts = dict(base.facts)
    facts["DOMAIN_RULE.RULE-1"] = {
        "value": "rule", "kind": "APPROVED_RULE",
        "provenance": "DOMAIN_POLICY", "approval_status": "FINALIZED",
    }
    assert _validate_fixture(item, FactRegistry(facts))[0] == ["distance"]
    facts["DOMAIN_RULE.RULE-1"] = {
        "value": "migration rule", "kind": "ASSUMPTION",
        "provenance": "LEGACY_MIGRATION", "approval_status": "PENDING",
    }
    with pytest.raises(ScenarioEvidenceContractError, match="approved rule ref"):
        _validate_fixture(item, FactRegistry(facts))


def test_scenario_v6_prompt_enforces_engineering_causal_fact_boundary():
    prompt = ScenarioFeasibilityAgent.SYSTEM_PROMPT
    assert "Malfunction(M) → direct system/vehicle behavior(B)" in prompt
    assert "interaction with explicit Scenario facts(I)" in prompt
    assert "Hazardous Event(H) → direct Potential Harm" in prompt
    assert "COUNTERFACTUAL TEST" in prompt
    assert "禁止加入任何输入未提供的新事件" in prompt
    assert "rear vehicle" in prompt
    assert "driver panic/steering error" in prompt
    assert "新增故障" in prompt
    assert "每个选择的dimension都必须" in prompt
    assert "负面结论是正常且预期的输出" in prompt
    assert "confidence表示“当前分类判断由给定证据支持”" in prompt


def test_scenario_allows_physical_functional_but_causally_irrelevant_result():
    class Client:
        def complete_json(self, request):
            item = _scenario_response_for_prompt(request)[0]
            item.update({
                "physically_feasible": True,
                "functionally_relevant": True,
                "causally_relevant": False,
                "breakpoint": "I_TO_H",
                "risk_dimension_changes": [],
                "causal_chain": {
                    "m_to_b": {"claim": "supported", "basis_type": "DIRECT_FACT", "evidence_refs": ["MF.description"]},
                    "b_to_i": {"claim": "supported", "basis_type": "DIRECT_FACT", "evidence_refs": ["MF.functional_effect"]},
                    "i_to_h": {"claim": "unsupported", "basis_type": "ASSUMPTION", "evidence_refs": []},
                },
                "rationale": "M→B成立，但缺少B与明确Scenario事实形成危险交互的条件，B→I断裂。",
                "hazardous_event": "",
                "potential_harm": "",
            })
            return LLMResponse(data={"assessments": [item]}, model="fake-model")

    result, _ = ScenarioFeasibilityAgent(Client()).assess(
        _batching_malfunction(), _batching_scenarios(1, padding=5),
    )
    assert result[0].physically_feasible
    assert result[0].functionally_relevant
    assert not result[0].causally_relevant
    assert result[0].risk_dimensions_changed == []
    assert result[0].hazardous_event == ""
    assert result[0].potential_harm == ""
    assert not result[0].retain


def test_scenario_allows_direct_evidence_causal_positive_result():
    values = list(RISK_DIMENSION_VALUES[:2])
    result = _scenario_dimensions_result(values)
    assert result.causally_relevant
    assert result.risk_dimensions_changed == values
    assert result.hazardous_event
    assert result.potential_harm
    assert result.retain


def test_single_scenario_runtime_pressure_fails_closed_without_recursion():
    class Client:
        def __init__(self):
            self.calls = 0

        def complete_json(self, request):
            self.calls += 1
            raise LLMOutputLimitError("length")

    client = Client()
    with pytest.raises(
        ScenarioAdaptiveBatchError,
        match="scenario_id=SCN-0.*failure_type=output_limit.*split_depth=0",
    ):
        ScenarioFeasibilityAgent(
            client, batch_max_chars=50000, batch_max_items=12,
        ).assess(_batching_malfunction(), _batching_scenarios(1, padding=5))
    assert client.calls == 1


def test_scenario_adaptive_split_stops_at_configured_max_depth():
    class Client:
        def __init__(self):
            self.calls = 0

        def complete_json(self, request):
            self.calls += 1
            raise LLMTimeoutError("timeout", attempts=2)

    client = Client()
    with pytest.raises(
        ScenarioAdaptiveBatchError,
        match="scenario_id=<multiple>.*failure_type=timeout.*split_depth=0",
    ):
        ScenarioFeasibilityAgent(
            client, batch_max_chars=50000, batch_max_items=12, max_split_depth=0,
        ).assess(_batching_malfunction(), _batching_scenarios(4, padding=5))
    assert client.calls == 1


def test_scenario_schema_and_coverage_errors_do_not_trigger_adaptive_split():
    class SchemaClient:
        def __init__(self):
            self.calls = 0

        def complete_json(self, request):
            self.calls += 1
            raise LLMSchemaContractError("wrong envelope")

    schema_client = SchemaClient()
    with pytest.raises(LLMSchemaContractError):
        ScenarioFeasibilityAgent(
            schema_client, batch_max_chars=50000, batch_max_items=12,
        ).assess(_batching_malfunction(), _batching_scenarios(4, padding=5))
    assert schema_client.calls == 1

    class CoverageClient:
        def __init__(self):
            self.calls = 0

        def complete_json(self, request):
            self.calls += 1
            assessments = _scenario_response_for_prompt(request)
            assessments.pop()
            return LLMResponse(data={"assessments": assessments}, model="fake-model")

    coverage_client = CoverageClient()
    with pytest.raises(ValueError, match="覆盖不完整"):
        ScenarioFeasibilityAgent(
            coverage_client, batch_max_chars=50000, batch_max_items=12,
        ).assess(_batching_malfunction(), _batching_scenarios(4, padding=5))
    assert coverage_client.calls == 1


def test_scenario_timeout_retry_budget_is_task_specific():
    calls = []

    def transport(url, headers, body, timeout):
        calls.append(json.loads(body))
        raise TimeoutError("read timeout")

    client = create_llm_client(LLMConfig(
        provider="openai-compatible", base_url="https://llm.example/v1",
        model="test-model", api_key="secret", max_retries=2,
        retry_backoff_seconds=0,
    ), transport=transport)
    with pytest.raises(LLMTimeoutError) as error:
        client.complete_json(LLMRequest(
            task="assess_scenario_feasibility", system_prompt="system", user_prompt="user",
            schema_name="ScenarioFeasibilityAssessmentList", prompt_version="v1",
        ))
    assert error.value.attempts == 2
    assert len(calls) == 2


def test_scenario_agent_retains_by_risk_difference_not_fixed_count():
    class ScenarioClient:
        def complete_json(self, request):
            return LLMResponse(data={"assessments": [_v9_fixture(item) for item in [
                {"scenario_id": "SCN-P", "physically_feasible": True, "functionally_relevant": True,
                 "causally_relevant": True, "risk_dimensions_changed": ["collision_object", "severity"],
                 "rationale": "行人对象改变伤害后果", "hazardous_event": "车辆未减速并接近行人",
                 "potential_harm": "可能碰撞行人并致伤", "confidence": 0.9},
                {"scenario_id": "SCN-V", "physically_feasible": True, "functionally_relevant": True,
                 "causally_relevant": True, "risk_dimensions_changed": ["relative_speed", "exposure"],
                 "rationale": "接近车辆改变相对速度和暴露度", "hazardous_event": "车辆未减速并接近车辆",
                 "potential_harm": "可能发生车辆碰撞并致伤", "confidence": 0.85},
                {"scenario_id": "SCN-X", "physically_feasible": False, "functionally_relevant": False,
                 "causally_relevant": False, "risk_dimensions_changed": [],
                 "rationale": "该运动关系在当前ODD内不可实现", "confidence": 0.8},
            ]]}, model="fake-model")

    malfunction = MalfunctionCandidate(
        "MF-1", "FUN-1", "loss", "制动请求丢失", "车辆无法减速",
        "车辆继续运动并接近障碍物", ["请求丢失", "车辆继续运动"],
    )
    scenarios = [
        ScenarioCandidate("SCN-P", "Parking", "行人场景", "行人在轨迹前方", {"object_type": "行人"}),
        ScenarioCandidate("SCN-V", "Parking", "车辆场景", "车辆接近", {"object_type": "车辆"}),
        ScenarioCandidate("SCN-X", "Parking", "无效场景", "物理关系冲突", {}),
    ]
    assessments, audit = ScenarioFeasibilityAgent(ScenarioClient()).assess(malfunction, scenarios)
    assert [item.scenario_id for item in assessments if item.retain] == ["SCN-P", "SCN-V"]
    assert audit["candidate_count"] == 3
    assert audit["retained_count"] == 2


def test_scenario_agent_rejects_unreported_candidate():
    class MissingScenarioClient:
        def complete_json(self, request):
            return LLMResponse(data={"assessments": []}, model="fake-model")

    malfunction = MalfunctionCandidate(
        "MF-1", "FUN-1", "loss", "制动请求丢失", "车辆无法减速",
        "车辆继续运动", ["请求丢失", "车辆继续运动"],
    )
    scenario = ScenarioCandidate("SCN-1", "Parking", "场景", "详情")
    with pytest.raises(ValueError, match="覆盖不完整"):
        ScenarioFeasibilityAgent(MissingScenarioClient()).assess(malfunction, [scenario])


def test_checkpoint_round_trip_preserves_typed_state():
    state = HARAState(run_id="run-checkpoint", stage=WorkflowStage.QUALITY_GATE, domain="avp")
    state.scenarios.append(ScenarioCandidate(
        "SCN-1", "Parking", "停车场低速泊车", "车辆接近行人", {"object_type": "pedestrian"},
        sources=[SourceRef("scenario_library", "Scenarios_Library", "row 8")],
    ))
    state.pending_reviews.append({"field": "exposure", "reason": "待工程确认"})
    repository = CheckpointRepository(ROOT / "output" / ".test-checkpoints" / "round-trip")

    path = repository.save(state)
    restored = repository.load(state.run_id)

    assert path.is_file()
    assert restored.to_dict() == state.to_dict()
    assert isinstance(restored.scenarios[0], ScenarioCandidate)


def test_workflow_graph_progresses_and_checkpoints():
    repository = CheckpointRepository(ROOT / "output" / ".test-checkpoints" / "progress")
    graph = WorkflowGraph(repository)

    def initialize(state):
        state.record("initialized")
        state.stage = WorkflowStage.EXTRACT
        return state

    def extract(state):
        state.record("extracted")
        state.stage = WorkflowStage.COMPLETE
        return state

    graph.add_node(WorkflowStage.INITIALIZE, initialize)
    graph.add_node(WorkflowStage.EXTRACT, extract)
    result = graph.run(HARAState(run_id="run-graph"))

    assert not result.interrupted
    assert result.state.stage is WorkflowStage.COMPLETE
    assert repository.load("run-graph").stage is WorkflowStage.COMPLETE


def test_workflow_graph_interrupts_at_pending_quality_gate():
    state = HARAState(run_id="run-review", stage=WorkflowStage.QUALITY_GATE)
    state.pending_reviews.append({"field": "f_tti", "reason": "缺少批准来源"})
    repository = CheckpointRepository(ROOT / "output" / ".test-checkpoints" / "review")
    result = WorkflowGraph(repository).run(state)

    assert result.interrupted
    assert result.reason == "pending_engineering_review"
    assert result.state.stage is WorkflowStage.QUALITY_GATE


def test_workflow_graph_blocks_node_that_makes_no_progress():
    graph = WorkflowGraph().add_node(WorkflowStage.INITIALIZE, lambda state: state)
    result = graph.run(HARAState(run_id="run-stalled"))

    assert result.state.stage is WorkflowStage.BLOCKED
    assert result.state.errors[-1]["type"] == "workflow_no_progress"


def test_semantic_frontend_graph_reaches_scoring_boundary():
    class RoutedClient:
        def complete_json(self, request):
            if request.task.startswith("extract_targeted_project_facts:"):
                fact_types = {
                    "speed": ["speed.search", "speed.control", "speed.parking"],
                    "performance": [
                        "performance.normal_braking", "performance.emergency_braking",
                        "performance.brake_response", "performance.steering_error",
                        "performance.steering_response",
                    ],
                    "driver": ["driver.inside", "driver.outside"],
                }[request.task.rsplit(":", 1)[1]]
                return LLMResponse(
                    data={"results": [
                        {"fact_type": fact_type, "status": "NOT_FOUND"}
                        for fact_type in fact_types
                    ]},
                    model="fake-model",
                )
            responses = {
                "extract_core_item_artifacts": {"item_definition": {
                    "system_description": "AVP低速泊车系统",
                    "item_boundary": "从泊车请求到车辆运动控制输出",
                    "operating_modes": ["搜索车位", "泊入"],
                    "odd": {
                        "locations": ["停车场"], "road_types": ["停车场通道"],
                        "weather_conditions": ["正常天气"], "road_surfaces": ["铺装路面"],
                        "speed_range_kph": [0, 5],
                    },
                    "performance_parameters": [], "source_location": "paragraph 1",
                    "source_excerpt": "AVP低速泊车", "confidence": 0.9, "status": "PENDING",
                }, "functions": [{
                    "function_id": "FUN-1", "name": "输出制动扭矩", "output": "目标制动扭矩请求",
                    "description": "控制车辆低速制动", "preconditions": ["AVP激活"],
                    "triggers": ["规划请求减速"], "odd_constraints": ["停车场"],
                    "fallback_behavior": "请求停车", "consequences": ["车辆不能按需减速"],
                    "source_location": "paragraph 1", "source_excerpt": "AVP制动控制",
                    "confidence": 0.9, "status": "PENDING",
                }]},
                "supplement_project_evidence": {
                    "performance_parameters": [], "driver_contexts": [], "exposure_inputs": [],
                },
                "assess_guideword_applicability": {"assessments": [{
                    "guideword": "loss", "applicable": True,
                    "rationale": "制动请求可能丢失", "confidence": 0.9, "status": "PENDING",
                }]},
                "derive_malfunctions_and_hazards": {"candidates": [{
                    "malfunction_id": "MF-1", "guideword": "loss", "description": "制动请求丢失",
                    "functional_effect": "车辆不能按需减速",
                    "vehicle_level_hazard": "车辆继续运动并接近障碍物",
                    "causal_chain": ["制动请求丢失", "车辆未减速", "车辆接近障碍物"],
                    "source_location": "paragraph 1", "source_excerpt": "AVP制动控制",
                    "confidence": 0.88, "status": "PENDING",
                }]},
                    "assess_scenario_feasibility": {"assessments": [_v9_fixture({
                    "scenario_id": "SCN-1", "physically_feasible": True,
                    "functionally_relevant": True, "causally_relevant": True,
                    "risk_dimensions_changed": ["collision_object", "severity"],
                    "rationale": "行人对象改变潜在伤害", "confidence": 0.87,
                    "hazardous_event": "车辆未减速并接近行人",
                    "potential_harm": "可能碰撞行人并致伤",
                    "status": "PENDING",
                    })]},
            }
            return LLMResponse(data=responses[request.task], model="fake-model")

    client = RoutedClient()
    profile = load_domain_profile("avp", require_approved=False)
    policy = AVPDomainPolicy(profile)
    scenario_facts = policy.candidate("near_pedestrian", ego_speed_kph=5.0)
    _apply_driver_context(scenario_facts, profile)
    workflow_inputs = SemanticWorkflowInputs(
        item_path=ROOT / "input" / "ItemDef.docx",
        guidewords=["loss"],
        scenario_candidates=[ScenarioCandidate(
            "SCN-1", "Parking", "停车场低速泊车", "车辆接近行人", scenario_facts,
        )],
    )
    workflow_agents = SemanticWorkflowAgents(
        item_definition=ItemDefinitionExtractionAgent(client),
        functions=FunctionExtractionAgent(client),
        guidewords=GuidewordApplicabilityAgent(client),
        malfunctions=MalfunctionHazardAgent(client),
        scenarios=ScenarioFeasibilityAgent(client),
    )
    graph = build_semantic_frontend_graph(
        workflow_inputs,
        workflow_agents,
    )
    result = graph.run(
        HARAState(run_id="semantic-front-end"),
        stop_before={WorkflowStage.SCORING},
    )

    assert result.interrupted
    assert result.reason == "planned_stop:scoring"
    assert result.state.stage is WorkflowStage.SCORING
    assert len(result.state.functions) == 1
    assert len(result.state.malfunctions) == 1
    assert result.state.malfunctions[0]["malfunction_id"] == "MF-FUN-1-001"
    assert result.state.malfunctions[0]["model_local_id"] == "MF-1"
    assert [item.scenario_id for item in result.state.scenarios] == ["SCN-1"]

    template = ROOT / "references" / "HARA_Template_AI_20260327.xlsx"
    full_graph = build_hara_agent_graph(
        workflow_inputs,
        workflow_agents,
        RiskWorkflowServices(
            scoring=_scoring_service(policy),
            asil_table=TemplateASILService(str(template)),
            ftti=FTTIService(),
            safety_goals=SafetyGoalCatalogService(profile),
        ),
    )
    full_result = full_graph.run(HARAState(run_id="full-agent-graph"))
    assert full_result.interrupted
    assert full_result.reason == "pending_engineering_review"
    assert full_result.state.stage is WorkflowStage.QUALITY_GATE
    assert len(full_result.state.risk_results) == 1
    assert len(full_result.state.safety_goals) == 1


def test_scoring_node_uses_domain_policy_and_input_template_matrix():
    profile = load_domain_profile("avp", require_approved=False)
    policy = AVPDomainPolicy(profile)
    facts = policy.candidate("near_pedestrian", ego_speed_kph=5.0)
    _apply_driver_context(facts, profile)
    state = HARAState(run_id="score-node", stage=WorkflowStage.SCORING)
    state.scenarios = [ScenarioCandidate(
        "SCN-1", "Parking", "停车场内低速泊车", "车辆以5km/h接近行人", facts,
    )]
    state.malfunctions = [{
        "malfunction_id": "MF-1", "function_id": "FUN-1", "guideword": "loss",
        "description": "制动请求丢失", "functional_effect": "车辆不能按需减速",
        "vehicle_level_hazard": "车辆继续运动并接近行人",
        "causal_chain": ["制动请求丢失", "车辆未减速"],
    }]
    state.item_definition["scenario_assessments"] = [{
        "malfunction_id": "MF-1", "scenario_id": "SCN-1",
            "physically_feasible": True, "functionally_relevant": True,
            "causally_relevant": True, "risk_dimensions_changed": ["severity"],
            "hazardous_event": "车辆未减速并接近行人",
        "potential_harm": "可能碰撞行人并致伤",
    }]
    template = ROOT / "references" / "HARA_Template_AI_20260327.xlsx"

    result = score_structured_scenarios(
        state,
        _scoring_service(policy),
        TemplateASILService(str(template)),
        FTTIService(),
    )

    assert result.stage is WorkflowStage.SAFETY_GOALS
    assert len(result.risk_results) == 1
    risk = result.risk_results[0]
    assert risk.asil.value == TemplateASILService(str(template)).determine(
        risk.severity.value, risk.exposure.value, risk.controllability.value,
    )
    assert risk.asil.sources[0].location == "ASIL_Table"
    assert risk.asil.status is ReviewStatus.PENDING
    assert risk.exposure_tf == "F"
    assert len(risk.exposure.sources) == 2
    assert risk.exposure.sources[1].source_type == "input_template"
    assert result.pending_reviews

    result.functions = [{"function_id": "FUN-1", "name": "输出制动扭矩"}]
    result = aggregate_safety_goals(result, SafetyGoalCatalogService(profile))
    assert result.stage is WorkflowStage.QUALITY_GATE
    assert len(result.safety_goals) == 1
    assert result.safety_goals[0].sg_id == "SG_AVP_02"
    assert result.safety_goals[0].status is ReviewStatus.PENDING
    assert result.risk_results[0].safety_goal_id == "SG_AVP_02"


def test_quality_gate_advances_only_without_pending_work():
    clean = HARAState(run_id="clean-gate", stage=WorkflowStage.QUALITY_GATE)
    assert pass_quality_gate(clean).stage is WorkflowStage.RENDER

    pending = HARAState(run_id="pending-gate", stage=WorkflowStage.QUALITY_GATE)
    pending.pending_reviews.append({"field": "ftti", "reason": "待批准"})
    assert pass_quality_gate(pending).stage is WorkflowStage.QUALITY_GATE
    assert pending.audit_trail[-1]["event"] == "quality_gate_blocked"

    rejected = HARAState(run_id="rejected-gate", stage=WorkflowStage.QUALITY_GATE)
    rejected.safety_goals = [SafetyGoal(
        "SG-1", "avoid harm", "safe stop", "A", None, status=ReviewStatus.REJECTED,
    )]
    assert not rejected.can_publish
    assert pass_quality_gate(rejected).stage is WorkflowStage.QUALITY_GATE
    assert rejected.audit_trail[-1]["rejected_safety_goals"] == 1


def test_empty_risk_dimensions_is_not_scored():
    class NeverScoring:
        def score(self, scenario, hazard_event):
            raise AssertionError("ineligible scenario must not be scored")

    class UnusedASIL:
        source = "unused"

    state = HARAState(run_id="empty-dimensions", stage=WorkflowStage.SCORING)
    state.scenarios = [ScenarioCandidate("SCN-1", "Parking", "parking", "detail", {})]
    state.malfunctions = [{"malfunction_id": "MF-1"}]
    state.item_definition["scenario_assessments"] = [{
        "malfunction_id": "MF-1", "scenario_id": "SCN-1",
        "physically_feasible": True, "functionally_relevant": True,
        "causally_relevant": True, "risk_dimensions_changed": [],
    }]
    result = score_structured_scenarios(state, NeverScoring(), UnusedASIL(), object())
    assert result.risk_results == []


@pytest.mark.parametrize("value", [0, -1, float("nan"), float("inf")])
def test_ftti_rejects_invalid_action_deceleration(value):
    with pytest.raises(ValueError, match="有限正数"):
        FTTIService(value)


def test_renderer_unique_index_rejects_duplicate_identity():
    with pytest.raises(ValueError, match="不唯一"):
        HARAExcelRenderer._unique_index(
            [{"id": "X"}, {"id": "X"}], lambda item: item["id"], "Risk",
        )


def test_downstream_preflight_rejects_duplicate_malfunction_ids():
    profile = load_domain_profile("avp", require_approved=False)
    state = HARAState(run_id="duplicate-malfunctions")
    state.functions = [
        {"function_id": "F01", "name": "输出制动扭矩"},
        {"function_id": "F02", "name": "输出转向角"},
    ]
    state.malfunctions = [
        {"malfunction_id": "MF_01", "function_id": "F01"},
        {"malfunction_id": "MF_01", "function_id": "F02"},
    ]
    with pytest.raises(ValueError, match="Malfunction ID不唯一"):
        DownstreamPreflightService(SafetyGoalCatalogService(profile)).validate(state)


def test_pending_guideword_flows_to_pending_reviews_and_blocks_release():
    class Client:
        def complete_json(self, request):
            return LLMResponse(data={"assessments": [
                {"guideword": "loss", "applicable": True, "rationale": "request lost",
                 "confidence": 0.8, "status": "FINALIZED"},
                {"guideword": "too late", "applicable": True,
                 "confidence": 0.7, "status": "FINALIZED"},
            ]}, model="fake-model")

    function = FunctionDefinition(
        "FUN-1", "brake", "torque request",
        sources=[SourceRef("item_definition", "item.docx", "p1", "brake")],
    )
    state = assess_guidewords(
        HARAState(run_id="pending-guideword-governance"),
        GuidewordApplicabilityAgent(Client()),
        [function], ["loss", "too late"],
    )

    audit = next(
        item for item in state.audit_trail
        if item["event"] == "guideword_applicability_assessed"
    )
    assert audit["coverage_count"] == 2
    assert audit["finalized_count"] == 1
    assert audit["pending_count"] == 1
    assert audit["parse_error_count"] == 1
    assert any(
        item["field"] == "guideword_assessments"
        and item["issue_type"] == "semantic_incomplete"
        for item in state.pending_reviews
    )
    assert not state.can_publish

    state.stage = WorkflowStage.QUALITY_GATE
    blocked = pass_quality_gate(state)
    assert blocked.stage is WorkflowStage.QUALITY_GATE
    assert blocked.audit_trail[-1]["event"] == "quality_gate_blocked"
    assert blocked.audit_trail[-1]["pending_reviews"] >= 1


def test_excel_renderer_preserves_template_and_overwrites_legacy_asil_formula():
    profile = load_domain_profile("avp", require_approved=False)
    policy = AVPDomainPolicy(profile)
    facts = policy.candidate("near_pedestrian", ego_speed_kph=5.0)
    facts.update({
        **next(item for item in profile.section("scenario_policy")["driver_context_candidates"]
               if item["context_id"] == "driver_outside_remote_monitoring"),
        "exposure_tf": "F",
    })
    state = HARAState(run_id="render-test", stage=WorkflowStage.SCORING)
    state.scenarios = [ScenarioCandidate(
        "SCN-1", "Parking", "停车场内低速泊车", "车辆以5km/h接近行人", facts,
    )]
    state.functions = [{"function_id": "FUN-1", "name": "输出制动扭矩", "output": "目标制动扭矩请求"}]
    state.malfunctions = [{
        "malfunction_id": "MF-1", "function_id": "FUN-1", "guideword": "loss",
        "description": "制动请求丢失", "functional_effect": "车辆不能按需减速",
        "vehicle_level_hazard": "车辆继续运动并接近行人",
        "causal_chain": ["制动请求丢失", "车辆未减速"],
    }]
    state.item_definition["scenario_assessments"] = [{
        "malfunction_id": "MF-1", "scenario_id": "SCN-1",
        "physically_feasible": True, "functionally_relevant": True,
        "causally_relevant": True, "risk_dimensions_changed": ["severity"],
        "hazardous_event": "车辆未减速并接近行人",
        "potential_harm": "可能碰撞行人并致伤",
    }]
    template = ROOT / "references" / "HARA_Template_AI_20260327.xlsx"
    state = score_structured_scenarios(
        state, _scoring_service(policy), TemplateASILService(str(template)), FTTIService(),
    )
    state = aggregate_safety_goals(state, SafetyGoalCatalogService(profile))
    risk = state.risk_results[0]
    for evidence in (risk.severity, risk.exposure, risk.controllability, risk.asil):
        evidence.status = ReviewStatus.FINALIZED
    risk.ftti_seconds.value = 0.5
    risk.ftti_seconds.status = ReviewStatus.FINALIZED
    state.safety_goals[0].status = ReviewStatus.FINALIZED
    state.safety_goals[0].ftti_seconds = 0.5
    state.pending_reviews.clear()
    state.stage = WorkflowStage.RENDER
    output = ROOT / "output" / ".test-report" / "HARA_Agent_Render_Test.xlsx"

    HARAExcelRenderer().render(state, template, output)

    from openpyxl import load_workbook
    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        assert "ASIL_Table" in workbook.sheetnames
        assert workbook["05_HARA"]["Q6"].value == risk.asil.value
        assert workbook["05_HARA"]["A6"].value == "HARA_001"
        assert workbook["06_Safety Goal"]["C4"].value == "SG_AVP_02"
    finally:
        workbook.close()

    from zipfile import ZipFile
    with ZipFile(template) as source, ZipFile(output) as rendered:
        assert set(source.namelist()) == set(rendered.namelist())
        assert rendered.testzip() is None
        target_sheets = {"xl/worksheets/sheet5.xml", "xl/worksheets/sheet6.xml"}
        assert all(
            source.read(name) == rendered.read(name)
            for name in source.namelist() if name not in target_sheets
        )
        assert b"x14:dataValidations" in rendered.read("xl/worksheets/sheet5.xml")
        assert b"x14:dataValidations" in rendered.read("xl/worksheets/sheet6.xml")


def test_excel_static_merge_is_hierarchical_and_ignores_hara_id():
    rows = [
        ["HARA_001", "F1", "O1", "loss", "M1", "H1", "S1"],
        ["HARA_002", "F1", "O1", "loss", "M1", "H1", "S1"],
        ["HARA_003", "F1", "O1", "more", "M2", "H2", "S2"],
    ]
    merges = HARAExcelRenderer._merge_ranges(rows)
    assert "B6:B8" in merges
    assert "C6:C8" in merges
    assert "D6:D7" in merges
    assert "E6:E7" in merges
    assert "F6:F7" in merges
    assert "G6:G7" in merges


def test_draft_renderer_keeps_pending_state_and_adds_visible_watermarks():
    template = ROOT / "references" / "HARA_Template_AI_20260327.xlsx"
    output = ROOT / "output" / ".test-report" / "HARA_Agent_Draft_Test.xlsx"
    state = HARAState(run_id="draft-render", stage=WorkflowStage.QUALITY_GATE)
    state.pending_reviews.append({"field": "ftti", "reason": "待工程批准"})

    with pytest.raises(ValueError, match="禁止生成正式报告"):
        HARAExcelRenderer().render(state, template, output)

    HARAExcelRenderer().render(state, template, output, draft=True)

    assert state.stage is WorkflowStage.QUALITY_GATE
    assert not state.can_publish
    from zipfile import ZipFile
    with ZipFile(output) as package:
        hara_xml = package.read("xl/worksheets/sheet5.xml").decode("utf-8")
        sg_xml = package.read("xl/worksheets/sheet6.xml").decode("utf-8")
        assert "DRAFT — NOT FOR RELEASE" in hara_xml
        assert "DRAFT — NOT FOR RELEASE" in sg_xml


def test_template_input_reader_loads_guidewords_and_dimensions_without_cartesian_product():
    template = ROOT / "references" / "HARA_Template_AI_20260327.xlsx"
    inputs = TemplateInputReader().read(template)

    assert len(inputs.guidewords) == TemplateInputReader.EXPECTED_GUIDEWORD_COUNT
    assert inputs.guidewords[0] == "unintended"
    assert inputs.guidewords[-1] == "as well as"
    assert "Operating scenarios" in inputs.scenario_dimensions
    assert any("Parking Lot" in value for value in inputs.scenario_dimensions["Operating scenarios"])
    assert inputs.scoring_standards.counts == {
        "severity": 4,
        "exposure_duration": 5,
        "exposure_frequency": 5,
        "controllability": 4,
    }
    assert inputs.scoring_standards.reference("exposure", "E4", "T").criterion.startswith(">10")
    assert "almost every drive" in inputs.scoring_standards.reference(
        "exposure", "E4", "F"
    ).criterion.lower()
    assert not hasattr(inputs, "scenario_combinations")


def test_avp_candidate_service_uses_risk_categories_not_template_cartesian_product():
    profile = load_domain_profile("avp", require_approved=False)
    service = AVPScenarioCandidateService(AVPDomainPolicy(profile))

    candidates, audit = service.generate()

    assert len(candidates) == 28
    assert audit["combination_strategy"] == "domain_risk_categories_no_cartesian_product"
    assert audit["speed_source_status"] == "MIGRATION_FALLBACK"
    assert all(item.status is ReviewStatus.PENDING for item in candidates)
    assert len({item.facts["scenario_variant"] for item in candidates}) == 6
    assert {item.facts["driver_position"] for item in candidates} == {"inside", "outside"}
    assert any("行人" in item.situational_detailing for item in candidates)
    assert any("30" in str(item.facts.get("target_speed_kph")) for item in candidates)
    assert all(item.scenario_contract_version == SCENARIO_CONTRACT_VERSION for item in candidates)
    assert len({item.scenario_id for item in candidates}) == len(candidates)
    assert len({item.semantic_fingerprint for item in candidates}) == len(candidates)


def test_atomic_scenario_expansion_is_stable_distinct_and_provenanced():
    profile = load_domain_profile("avp", require_approved=False)
    service = AVPScenarioCandidateService(AVPDomainPolicy(profile))
    first, _ = service.generate(ego_speed_kph=5.0)
    second, _ = service.generate(ego_speed_kph=5.0)
    selected = [item for item in first if item.source_scenario_id == "near_vehicle_10"]

    assert [item.atomic_variant for item in selected] == [
        "front_vehicle", "front_vehicle", "rear_vehicle", "rear_vehicle",
    ]
    assert {item.facts["collision_geometry"] for item in selected} == {
        "frontal vehicle", "rear vehicle",
    }
    assert [item.scenario_id for item in first] == [item.scenario_id for item in second]
    assert [item.semantic_fingerprint for item in first] == [
        item.semantic_fingerprint for item in second
    ]


def test_scenario_fact_consistency_preflight_uses_declared_relative_speed_basis():
    with pytest.raises(ScenarioFactConsistencyError, match="field=relative_speed_kph"):
        AVPScenarioCandidateService._validate_fact_consistency(
            "source", "variant", {
                "ego_speed_kph": 5.0, "target_speed_kph": 0.0,
                "relative_speed_kph": 0.0, "relative_speed_mode": "ego",
            },
        )


def test_scenario_structured_speed_is_not_overwritten_by_standstill_label():
    profile = load_domain_profile("avp", require_approved=False)
    candidates, _ = AVPScenarioCandidateService(AVPDomainPolicy(profile)).generate(
        ego_speed_kph=5.0,
    )
    standstill = next(
        item for item in candidates if item.source_scenario_id == "controlled_standstill"
    )
    assert standstill.facts["ego_speed_kph"] == 5.0
    assert standstill.facts["relative_speed_kph"] == 0.0
    assert standstill.facts["relative_speed_mode"] == "zero"


def test_checkpoint_invalidates_legacy_scenario_outputs_but_preserves_malfunctions():
    value = HARAState(run_id="legacy", stage=WorkflowStage.SCORING).to_dict()
    value.pop("scenario_contract_version")
    value["malfunctions"] = [{"malfunction_id": "MF-KEEP"}]
    value["scenarios"] = [ScenarioCandidate("SCN-OLD", "Parking", "old", "old").to_dict()]
    value["item_definition"]["scenario_assessments"] = [{"scenario_id": "SCN-OLD"}]

    restored = HARAState.from_dict(value)

    assert restored.stage is WorkflowStage.MALFUNCTIONS
    assert restored.malfunctions == [{"malfunction_id": "MF-KEEP"}]
    assert restored.scenarios == []
    assert "scenario_assessments" not in restored.item_definition
    assert restored.scenario_contract_version == SCENARIO_CONTRACT_VERSION
    assert restored.audit_trail[-1]["event"] == "scenario_checkpoint_invalidated"


def test_scenario_smoke_wires_runtime_policy_and_enters_scenario_agent(monkeypatch):
    smoke_dir = ROOT / "runtime" / ".test-scenario-smoke-wiring"
    checkpoint = HARAState(run_id="smoke-wiring", stage=WorkflowStage.MALFUNCTIONS)
    checkpoint.functions = [{
        "function_id": "FUN-1", "name": "输出制动扭矩", "output": "制动请求",
    }]
    checkpoint.malfunctions = [{
        "malfunction_id": "MF-FUN-1-001", "model_local_id": "MF_01",
        "function_id": "FUN-1", "guideword": "loss", "description": "制动请求丢失",
        "functional_effect": "车辆不能减速", "vehicle_level_hazard": "车辆继续运动",
        "causal_chain": ["请求丢失", "车辆未减速"], "status": "PENDING", "confidence": 0.8,
    }]
    CheckpointRepository(smoke_dir).save(checkpoint)

    received_policies = []
    real_candidate_service = scenario_smoke.AVPScenarioCandidateService

    class RecordingCandidateService:
        def __init__(self, policy):
            received_policies.append(policy)
            self.delegate = real_candidate_service(policy)

        def generate(self, **kwargs):
            return self.delegate.generate(**kwargs)

    calls = []

    class FakeClient:
        def complete_json(self, request):
            calls.append(request)
            assessments = _scenario_response_for_prompt(request)
            for item in assessments:
                item.update({
                    "physically_feasible": False,
                    "functionally_relevant": False,
                    "causally_relevant": False,
                    "breakpoint": "I_TO_H",
                    "risk_dimension_changes": [],
                    "causal_chain": {
                        "m_to_b": {"claim": "supported", "basis_type": "DIRECT_FACT", "evidence_refs": ["MF.description"]},
                        "b_to_i": {"claim": "supported", "basis_type": "DIRECT_FACT", "evidence_refs": ["MF.functional_effect"]},
                        "i_to_h": {"claim": "unsupported", "basis_type": "ASSUMPTION", "evidence_refs": []},
                    },
                    "hazardous_event": "",
                    "potential_harm": "",
                })
            return LLMResponse(data={"assessments": assessments}, model="fake-model")

    runtime = default_domain_registry().create("avp", require_approved=False)
    monkeypatch.setattr(scenario_smoke, "AVPScenarioCandidateService", RecordingCandidateService)
    monkeypatch.setattr(scenario_smoke, "create_llm_client", lambda config: FakeClient())
    monkeypatch.setenv("HARA_LLM_BASE_URL", "https://llm.example/v1")
    monkeypatch.setenv("HARA_LLM_MODEL", "fake-model")
    monkeypatch.setenv("HARA_LLM_API_KEY", "fake-key")

    result = scenario_smoke.main([
        "--domain", "avp", "--template", str(TEMPLATE),
        "--run-dir", str(smoke_dir), "--run-id", "smoke-wiring",
        "--malfunction-count", "1", "--scenario-count", "3",
    ])

    assert result == 0
    assert len(calls) >= 1
    assert calls[0].task == "assess_scenario_feasibility"
    assert sum(call.metadata["scenario_count"] for call in calls) == 3
    assert len(received_policies) == 1
    assert type(received_policies[0]) is type(runtime.policy)
    assert received_policies[0].profile == runtime.profile


def test_application_generates_audited_draft_at_pending_quality_gate():
    guidewords = [
        "unintended", "always active", "loss", "too large", "too small", "too early",
        "too late", "too fast", "too slow", "too long", "too short", "incomplete",
        "different to", "as well as",
    ]
    class ApplicationClient:
        def complete_json(self, request):
            if request.task == "extract_core_item_artifacts":
                data = {"item_definition": {
                    "system_description": "AVP低速泊车系统",
                    "item_boundary": "从泊车请求到车辆运动控制输出",
                    "operating_modes": ["搜索车位", "泊入"],
                    "odd": {
                        "locations": ["停车场"], "road_types": ["停车场通道"],
                        "weather_conditions": ["正常天气"], "road_surfaces": ["铺装路面"],
                        "speed_range_kph": [0, 5],
                    },
                    "performance_parameters": [],
                    "driver_contexts": [{
                        "context_id": "driver_outside_remote_monitoring",
                        "driver_position": "outside",
                        "driver_state": "驾驶员不在驾驶位，仅能远程监控或发出停止请求",
                        "direct_vehicle_control": False,
                        "intervention_channels": ["remote_stop_request"],
                    }],
                    "exposure_inputs": [], "source_location": "paragraph 1",
                    "source_excerpt": "AVP低速泊车", "confidence": 0.9, "status": "PENDING",
                }, "functions": [{
                    "function_id": "FUN-1", "name": "输出制动扭矩", "output": "目标制动扭矩请求",
                    "description": "控制车辆低速制动", "preconditions": ["AVP激活"],
                    "triggers": ["规划请求减速"], "odd_constraints": ["停车场"],
                    "fallback_behavior": "请求停车", "consequences": ["车辆不能按需减速"],
                    "source_location": "paragraph 1", "source_excerpt": "AVP制动控制",
                    "confidence": 0.9, "status": "PENDING",
                }]}
            elif request.task == "supplement_project_evidence":
                data = {
                    "performance_parameters": [],
                    "driver_contexts": [{
                        "context_id": "driver_outside_remote_monitoring",
                        "driver_position": "outside",
                        "driver_state": "remote monitoring",
                        "direct_vehicle_control": False,
                        "intervention_channels": ["remote_stop_request"],
                        "source_location": "paragraph 1",
                        "source_excerpt": "remote monitoring",
                    }],
                    "exposure_inputs": [],
                }
            elif request.task == "assess_guideword_applicability":
                data = {"assessments": [{
                    "guideword": word,
                    "applicable": word == "loss",
                    "rationale": "制动请求可能丢失" if word == "loss" else "该偏差不适用于离散制动请求",
                    "confidence": 0.85,
                        "status": "PENDING",
                    } for word in guidewords]}
            elif request.task == "derive_malfunctions_and_hazards":
                data = {"candidates": [{
                    "malfunction_id": "MF-1", "guideword": "loss", "description": "制动请求丢失",
                    "functional_effect": "车辆不能按需减速",
                    "vehicle_level_hazard": "车辆继续运动并接近障碍物",
                    "causal_chain": ["制动请求丢失", "车辆未减速", "车辆接近障碍物"],
                    "source_location": "paragraph 1", "source_excerpt": "AVP制动控制",
                    "confidence": 0.88, "status": "PENDING",
                }]}
            elif request.task == "assess_scenario_feasibility":
                payload_text = request.user_prompt.split("Scenarios=", 1)[1].split(
                    "\n这里只评估", 1,
                )[0]
                scenario_ids = [item["scenario_id"] for item in json.loads(payload_text)]
                data = {"assessments": [_v9_fixture({
                    "scenario_id": scenario_id,
                    "physically_feasible": True,
                    "functionally_relevant": True,
                    "causally_relevant": True,
                    "risk_dimensions_changed": ["collision_object", "severity"],
                    "rationale": "目标物和运动关系会改变风险结果",
                    "hazardous_event": f"制动请求丢失时车辆接近{scenario_id}",
                    "potential_harm": "可能发生碰撞并造成人员受伤",
                    "confidence": 0.82,
                    "status": "PENDING",
                }) for scenario_id in scenario_ids]}
            elif request.task.startswith("extract_targeted_project_facts:"):
                fact_types = {
                    "speed": ["speed.search", "speed.control", "speed.parking"],
                    "performance": [
                        "performance.normal_braking", "performance.emergency_braking",
                        "performance.brake_response", "performance.steering_error",
                        "performance.steering_response",
                    ],
                    "driver": ["driver.inside", "driver.outside"],
                }[request.task.rsplit(":", 1)[1]]
                data = {"results": [
                    {"fact_type": fact_type, "status": "NOT_FOUND"}
                    for fact_type in fact_types
                ]}
            else:
                raise AssertionError(request.task)
            return LLMResponse(data=data, model="fake-model")

    output = ROOT / "output" / ".test-report" / "HARA_Application_Draft.xlsx"
    config = RunConfig(
        item_path=ROOT / "input" / "ItemDef.docx",
        template_path=ROOT / "references" / "HARA_Template_AI_20260327.xlsx",
        output_path=output,
        run_dir=ROOT / "runtime" / ".test-agent-application",
        domain="avp",
        allow_draft=True,
        run_id="application-draft-test",
        operating_mode="parking",
        allow_aggregate_speed_fallback=True,
    )

    result = HARAApplication(config, ApplicationClient()).run()

    assert result.interrupted
    assert result.reason == "draft_report_generated_pending_review"
    assert result.state.stage is WorkflowStage.QUALITY_GATE
    candidate_event = next(
        item for item in result.state.audit_trail if item["event"] == "scenario_candidates_prepared"
    )
    assert len(result.state.risk_results) == candidate_event["atomic_candidate_count"]
    assert len(result.state.safety_goals) > 0
    assert not result.state.can_publish
    assert output.is_file()
    assert any(item["event"] == "draft_excel_report_rendered" for item in result.state.audit_trail)
    assert candidate_event["ego_speed_kph"] == 5.0
    assert candidate_event["speed_source_status"] == "EXPLICIT_AGGREGATE_FALLBACK"
    guideword_event = next(
        item for item in result.state.audit_trail
        if item["event"] == "guideword_applicability_assessed"
    )
    assert guideword_event["complete_count"] == len(guidewords)
    assert guideword_event["incomplete_count"] == 0
    assert guideword_event["review_pending_count"] == len(guidewords)
    assert any(
        item.get("issue_type") == "pending_review"
        for item in result.state.pending_reviews
    )
    malfunction_summary = next(
        item for item in result.state.audit_trail
        if item["event"] == "malfunction_derivation_summary"
    )
    assert malfunction_summary["llm_called"] == 1
    assert malfunction_summary["candidate_count"] == 1
