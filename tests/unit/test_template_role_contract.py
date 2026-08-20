from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.cell_range import CellRange

from hara_agent.contracts import (
    REQUIRED_TEMPLATE_ROLES, TemplateDiagnosticCode, TemplateRole,
    TemplateRoleConfirmation,
)
from hara_agent.template import (
    TemplateRoleAmbiguityError, TemplateRoleCompiler, TemplateRoleManifestStore,
    TemplateWorkbookScanner,
)


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "references" / "HARA_Template_AI_20260327.xlsx"


@pytest.fixture
def template_workspace() -> Path:
    path = ROOT / "tests" / ".template-role-tmp"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _copy_template(tmp_path: Path, name: str) -> Path:
    target = tmp_path / name
    shutil.copy2(TEMPLATE, target)
    return target


def _mutate(path: Path, callback) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        callback(workbook)
        workbook.save(path)
    finally:
        workbook.close()


def test_current_template_discovers_all_canonical_roles():
    contract = TemplateRoleCompiler().compile(TEMPLATE, use_manifest=False)

    assert contract.ready is True
    assert {item.role for item in contract.role_bindings} == set(REQUIRED_TEMPLATE_ROLES)
    assert contract.binding(TemplateRole.GUIDEWORD_TABLE).region == "A2:B16"
    assert contract.binding(TemplateRole.HARA_OUTPUT_TABLE).sheet == "05_HARA"

    method = TemplateRoleCompiler().compile_method_foundation(
        TEMPLATE, use_manifest=False
    )
    assert method.metadata["engineering_rules_compiled"] is False
    assert method.severity.roles == (
        TemplateRole.SEVERITY_LEVELS,
        TemplateRole.SEVERITY_RULES,
    )


def test_severity_sheet_rename_keeps_same_source_regions(template_workspace: Path):
    template = _copy_template(template_workspace, "severity-renamed.xlsx")
    _mutate(template, lambda wb: setattr(wb["Severity"], "title", "Impact Consequences"))

    contract = TemplateRoleCompiler().compile(template, use_manifest=False)

    assert contract.binding(TemplateRole.SEVERITY_LEVELS).sheet == "Impact Consequences"
    assert contract.binding(TemplateRole.SEVERITY_LEVELS).region == "B7:F10"
    assert contract.binding(TemplateRole.SEVERITY_RULES).sheet == "Impact Consequences"


def test_asil_sheet_rename_is_discovered_structurally(template_workspace: Path):
    template = _copy_template(template_workspace, "asil-renamed.xlsx")
    _mutate(template, lambda wb: setattr(wb["ASIL_Table"], "title", "Risk Classification"))

    contract = TemplateRoleCompiler().compile(template, use_manifest=False)

    binding = contract.binding(TemplateRole.ASIL_MATRIX)
    assert binding.sheet == "Risk Classification"
    assert binding.detection_method == "structural_matrix_signature"


def test_sheet_order_does_not_change_role_sources(template_workspace: Path):
    template = _copy_template(template_workspace, "sheet-order.xlsx")

    def reverse_sheets(workbook):
        workbook._sheets = list(reversed(workbook._sheets))

    _mutate(template, reverse_sheets)
    contract = TemplateRoleCompiler().compile(template, use_manifest=False)

    assert contract.binding(TemplateRole.SCENARIO_MODEL).sheet == "Scenarios_Library"
    assert contract.binding(TemplateRole.ASIL_MATRIX).sheet == "ASIL_Table"


def test_severity_region_can_move_down_without_python_change(template_workspace: Path):
    template = _copy_template(template_workspace, "severity-shifted.xlsx")

    def shift_region(workbook):
        sheet = workbook["Severity"]
        merged = [str(item) for item in sheet.merged_cells.ranges]
        for region in merged:
            sheet.unmerge_cells(region)
        sheet.insert_rows(1, amount=8)
        for region in merged:
            shifted = CellRange(region)
            shifted.shift(row_shift=8)
            sheet.merge_cells(str(shifted))

    _mutate(template, shift_region)

    contract = TemplateRoleCompiler().compile(template, use_manifest=False)

    assert contract.binding(TemplateRole.SEVERITY_LEVELS).region == "B15:F18"
    assert contract.binding(TemplateRole.SEVERITY_RULES).region.startswith("I11:")


def test_two_isomorphic_severity_regions_fail_closed(template_workspace: Path):
    template = _copy_template(template_workspace, "ambiguous-severity.xlsx")

    def duplicate(workbook):
        clone = workbook.copy_worksheet(workbook["Severity"])
        clone.title = "Impact Rules Alternative"

    _mutate(template, duplicate)

    with pytest.raises(TemplateRoleAmbiguityError) as exc_info:
        TemplateRoleCompiler().compile(template, use_manifest=False)

    assert exc_info.value.role in {
        TemplateRole.SEVERITY_LEVELS,
        TemplateRole.SEVERITY_RULES,
    }
    assert len(exc_info.value.candidates) == 2


def test_workflow_coordinate_mismatches_and_malformed_text_are_diagnostic():
    contract = TemplateRoleCompiler().compile(TEMPLATE, use_manifest=False)
    mismatches = [
        item for item in contract.diagnostics
        if item.code is TemplateDiagnosticCode.WORKFLOW_COORDINATE_MISMATCH
    ]

    assert {
        (item.details["field"], item.details["declared_column"], item.details["report_column"])
        for item in mismatches
    } == {
        ("situational_detailing", "I", "H"),
        ("hazardous_event", "H", "I"),
        ("safety_goal", "T", "S"),
        ("safe_state", "U", "T"),
    }
    assert any(
        item.code is TemplateDiagnosticCode.STRUCTURE_INCOMPLETE
        and "C0,11,C2" in item.message
        for item in contract.diagnostics
    )


def test_guideword_cardinality_is_template_data_in_new_compiler(template_workspace: Path):
    template = _copy_template(template_workspace, "guideword-15.xlsx")

    def add_guideword(workbook):
        sheet = workbook["04_HAZOP"]
        sheet["A17"] = "intermittent"
        sheet["B17"] = "Action alternates between available and unavailable."

    _mutate(template, add_guideword)
    contract = TemplateRoleCompiler().compile(template, use_manifest=False)

    assert contract.binding(TemplateRole.GUIDEWORD_TABLE).region == "A2:B17"
    assert contract.ready is True


def test_generated_manifest_is_bound_to_template_hash(template_workspace: Path):
    store = TemplateRoleManifestStore(template_workspace / "manifests")
    compiler = TemplateRoleCompiler(manifest_store=store)
    original = compiler.compile(TEMPLATE)

    assert store.path_for(original.template_hash).is_file()
    assert store.load(original.template_hash) == original

    changed = _copy_template(template_workspace, "changed-template.xlsx")
    _mutate(changed, lambda wb: setattr(wb["Severity"], "title", "Impact Consequences"))
    changed_contract = compiler.compile(changed)

    assert changed_contract.template_hash != original.template_hash
    assert store.path_for(changed_contract.template_hash).is_file()


def test_human_ambiguity_confirmation_is_hash_bound_and_reused(template_workspace: Path):
    template = _copy_template(template_workspace, "confirmed-ambiguity.xlsx")

    def duplicate(workbook):
        clone = workbook.copy_worksheet(workbook["Severity"])
        clone.title = "Impact Rules Alternative"

    _mutate(template, duplicate)
    snapshot = TemplateWorkbookScanner().scan(template)
    confirmation = TemplateRoleConfirmation(
        template_hash=snapshot.source_hash,
        selected_regions={
            TemplateRole.SEVERITY_LEVELS.value: {
                "sheet": "Severity", "region": "B7:F10"
            },
            TemplateRole.SEVERITY_RULES.value: {
                "sheet": "Severity", "region": "I3:I26"
            },
        },
        confirmed_by="test-reviewer",
        confirmed_at="2026-08-20T00:00:00Z",
        rationale="The original region is the approved method source.",
    )
    store = TemplateRoleManifestStore(template_workspace / "confirmed-manifests")
    compiler = TemplateRoleCompiler(manifest_store=store)

    confirmed = compiler.compile(template, confirmation=confirmation)
    cached = compiler.compile(template)

    assert confirmed == cached
    assert confirmed.binding(TemplateRole.SEVERITY_LEVELS).detection_method == (
        "human_confirmation"
    )

    _mutate(template, lambda wb: setattr(wb["Severity"], "title", "Approved Impact Rules"))
    with pytest.raises(TemplateRoleAmbiguityError):
        compiler.compile(template)
