from dataclasses import replace
from pathlib import Path
import re
from zipfile import ZipFile

from openpyxl import load_workbook

from hara_agent.services.reporting import HARAExcelRenderer
from hara_agent.template import TemplateRoleCompiler
from hara_agent.workflow import HARAState, WorkflowStage


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "references" / "HARA_Template_AI_20260327.xlsx"


def test_renderer_uses_compiled_report_sheet_names():
    test_dir = ROOT / "output" / ".test-report"
    test_dir.mkdir(parents=True, exist_ok=True)
    template = test_dir / "renamed-template.xlsx"
    workbook = load_workbook(TEMPLATE)
    workbook["05_HARA"].title = "Risk Register"
    workbook["06_Safety Goal"].title = "Safety Objectives"
    workbook.save(template)
    workbook.close()

    method = TemplateRoleCompiler().compile_method(TEMPLATE)
    contract = replace(
        method.report_contract,
        hara_fields=tuple(
            replace(item, sheet="Risk Register")
            for item in method.report_contract.hara_fields
        ),
        safety_goal_fields=tuple(
            replace(item, sheet="Safety Objectives")
            for item in method.report_contract.safety_goal_fields
        ),
    )
    output = test_dir / "renamed-rendered.xlsx"
    state = HARAState(run_id="dynamic-report-contract", stage=WorkflowStage.RENDER)
    HARAExcelRenderer(contract).render(state, template, output, draft=True)

    rendered = load_workbook(output, read_only=True)
    try:
        assert rendered["Risk Register"]["A3"].value.startswith("DRAFT")
        assert rendered["Safety Objectives"]["A2"].value.startswith("DRAFT")
    finally:
        rendered.close()


def test_renderer_rejects_report_contract_with_duplicate_output_column():
    method = TemplateRoleCompiler().compile_method(TEMPLATE)
    mappings = list(method.report_contract.hara_fields)
    mappings[1] = replace(mappings[1], column_index=mappings[0].column_index)
    contract = replace(method.report_contract, hara_fields=tuple(mappings))

    import pytest
    with pytest.raises(ValueError, match="Duplicate HARA output column"):
        HARAExcelRenderer(contract).render(
            HARAState(run_id="invalid-report-contract"),
            TEMPLATE,
            ROOT / "output" / ".test-report" / "invalid.xlsx",
            draft=True,
        )


def test_renderer_preserves_ignorable_namespace_declarations_and_drops_calc_chain():
    method = TemplateRoleCompiler().compile_method(TEMPLATE)
    output = ROOT / "output" / ".test-report" / "namespace-safe.xlsx"

    HARAExcelRenderer(method.report_contract).render(
        HARAState(run_id="namespace-safe"), TEMPLATE, output, draft=True,
    )

    with ZipFile(output, "r") as package:
        assert "xl/calcChain.xml" not in package.namelist()
        relationships = package.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        content_types = package.read("[Content_Types].xml").decode("utf-8")
        assert "calcChain" not in relationships
        assert "calcChain" not in content_types
        for name in ("xl/worksheets/sheet5.xml", "xl/worksheets/sheet6.xml"):
            xml = package.read(name).decode("utf-8")
            root = re.search(r"<worksheet\b[^>]*>", xml)
            assert root is not None
            assert 'mc:Ignorable="x14ac xr xr2 xr3"' in root.group(0)
            for prefix in ("x14ac", "xr", "xr2", "xr3"):
                assert f"xmlns:{prefix}=" in root.group(0)
            assert "ns2:uid=" not in root.group(0)
            assert "xr:uid=" in root.group(0)


def test_renderer_keeps_inapplicable_and_nonhazardous_guideword_rows_as_na():
    method = TemplateRoleCompiler().compile_method(TEMPLATE)
    output = ROOT / "output" / ".test-report" / "guideword-na-rows.xlsx"
    state = HARAState(run_id="guideword-na-rows", stage=WorkflowStage.RENDER)
    state.functions = [{
        "function_id": "F01",
        "name": "Parking motion control",
        "output": "Vehicle motion request",
    }]
    state.guideword_assessments = [
        {
            "function_id": "F01",
            "guideword": "reverse",
            "applicable": False,
            "disposition": "NOT_APPLICABLE",
            "rationale": "The output has no reversible direction semantic.",
        },
        {
            "function_id": "F01",
            "guideword": "different to",
            "applicable": True,
            "disposition": "NO_CREDIBLE_HAZARD",
            "rationale": "The identified difference cannot change vehicle behavior.",
        },
    ]

    HARAExcelRenderer(method.report_contract).render(
        state, TEMPLATE, output, draft=True,
    )

    mappings = {
        item.canonical_field: item.column_index
        for item in method.report_contract.hara_fields
    }
    start_row = max(
        row for item in method.report_contract.hara_fields for row in item.header_rows
    ) + 1
    # Read-only worksheets may retain their package handle on Windows even
    # after close(); this test deletes the workbook in its finally block.
    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        sheet = workbook[method.report_contract.hara_fields[0].sheet]
        assert sheet.cell(start_row, mappings["guideword"]).value == "reverse"
        assert sheet.cell(start_row, mappings["malfunction"]).value.startswith("N/A：")
        assert "语义不适用" in sheet.cell(start_row, mappings["remark"]).value
        assert sheet.cell(start_row + 1, mappings["guideword"]).value == "different to"
        assert "可信车辆级危害" in sheet.cell(start_row + 1, mappings["remark"]).value
        for field in (
            "hazard", "scenario", "hazardous_event", "severity",
            "exposure", "controllability", "asil", "safety_goal", "safe_state",
        ):
            if field in mappings:
                assert sheet.cell(start_row, mappings[field]).value is None
                assert sheet.cell(start_row + 1, mappings[field]).value is None
    finally:
        workbook.close()
        output.unlink()
