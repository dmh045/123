from __future__ import annotations

import shutil
from collections import Counter
from pathlib import Path

import pytest
from openpyxl import load_workbook

from hara_agent.contracts import (
    CompileStatus, CompilerDiagnosticCode, NormativeStrength,
    ParseStatus, RangePredicate,
)
from hara_agent.template import TemplateRoleCompiler


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "references" / "HARA_Template_AI_20260327.xlsx"


@pytest.fixture
def compiler_workspace() -> Path:
    path = ROOT / "tests" / ".template-role-tmp" / "full-compiler"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _copy(workspace: Path, name: str) -> Path:
    target = workspace / name
    shutil.copy2(TEMPLATE, target)
    return target


def _mutate(path: Path, callback) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        callback(workbook)
        workbook.save(path)
    finally:
        workbook.close()


def _compile(path: Path):
    return TemplateRoleCompiler().compile_method(path, use_manifest=False)


def test_current_template_compiles_full_method_contract():
    method = _compile(TEMPLATE)

    assert method.compile_status is CompileStatus.READY_WITH_WARNINGS
    assert method.engineering_rules_compiled is True
    assert len(method.guidewords.guidewords) == 14
    assert len(method.scenario_model.dimensions) == 5
    assert len(method.severity.rules) == 24
    assert len(method.exposure.duration_rules) == 8
    assert len(method.exposure.frequency_rules) == 8
    assert [item.level for item in method.exposure.scale] == ["E0", "E1", "E2", "E3", "E4"]
    assert len(method.exposure.examples) == 67
    assert len(method.exposure.situation_mappings) == 376
    assert len(method.controllability.criteria) == 4
    assert len(method.controllability.examples) == 11
    assert len(method.asil.mappings) == 80
    assert len(method.required_fact_specs) == 9
    assert len(method.report_contract.hara_fields) == 21
    assert len(method.report_contract.safety_goal_fields) == 6
    assert method.scenario_model.source_type == "METHOD_SCENARIO_ONTOLOGY"
    assert {item.activity for item in method.workflow.steps} >= {
        "FUNCTION_EXTRACTION", "GUIDEWORD_APPLICATION", "MALFUNCTION_DERIVATION",
        "HAZARD_DERIVATION", "SCENARIO_CONSTRUCTION", "SCENARIO_DETAILING",
        "HAZARDOUS_EVENT_DERIVATION", "SEVERITY_ASSESSMENT",
        "EXPOSURE_ASSESSMENT", "CONTROLLABILITY_ASSESSMENT", "ASIL_LOOKUP",
        "SAFETY_GOAL_DERIVATION", "SAFE_STATE_DERIVATION", "REPORTING",
    }


def test_rule_classification_does_not_promote_examples():
    method = _compile(TEMPLATE)
    strengths = Counter(rule.normative_strength for rule in method.all_rules())

    assert strengths[NormativeStrength.NORMATIVE] == 24
    assert strengths[NormativeStrength.CRITERION] == 20
    assert strengths[NormativeStrength.EXAMPLE] == 78
    assert all(not rule.executable for rule in method.exposure.examples)
    assert all(not rule.executable for rule in method.controllability.examples)


def test_severity_rule_mutation_changes_contract_without_python_change(
    compiler_workspace: Path,
):
    template = _copy(compiler_workspace, "severity-result-mutated.xlsx")
    _mutate(template, lambda wb: setattr(wb["Severity"]["I20"], "value", "Car-pedestrian (v:15-30): S3"))

    original = _compile(TEMPLATE)
    changed = _compile(template)
    original_rule = next(rule for rule in original.severity.rules if "Car-pedestrian (v:15-30)" in rule.raw_text)
    changed_rule = next(rule for rule in changed.severity.rules if "Car-pedestrian (v:15-30)" in rule.raw_text)

    assert original_rule.result == "S2"
    assert changed_rule.result == "S3"


def test_exposure_threshold_mutation_changes_range(compiler_workspace: Path):
    template = _copy(compiler_workspace, "duration-mutated.xlsx")
    _mutate(template, lambda wb: setattr(wb["Exposure"]["D9"], "value", "<2 % of average operating time"))

    method = _compile(template)
    rule = next(
        item for item in method.exposure.duration_rules
        if item.assessment_method == "T" and item.result == "E2"
    )
    predicate = next(item for item in rule.predicates if isinstance(item, RangePredicate))

    assert predicate.upper == 2.0
    assert any(
        item.code is CompilerDiagnosticCode.RULE_RANGE_OVERLAP
        for item in method.diagnostics
    )


def test_guideword_mutation_changes_cardinality(compiler_workspace: Path):
    template = _copy(compiler_workspace, "guideword-15-full.xlsx")

    def mutate(workbook):
        sheet = workbook["04_HAZOP"]
        sheet["A17"] = "intermittent"
        sheet["B17"] = "Action alternates between available and unavailable."

    _mutate(template, mutate)
    assert len(_compile(template).guidewords.guidewords) == 15


def test_scenario_dimension_mutation_is_discovered(compiler_workspace: Path):
    template = _copy(compiler_workspace, "scenario-dimension.xlsx")

    def mutate(workbook):
        sheet = workbook["Scenarios_Library"]
        sheet["G2"] = "Traffic Density"
        sheet["G3"] = "Low"
        sheet["G4"] = "High"

    _mutate(template, mutate)
    method = _compile(template)

    assert len(method.scenario_model.dimensions) == 6
    assert method.scenario_model.dimensions[-1].canonical_name == "TRAFFIC_DENSITY"


def test_output_column_mutation_changes_only_report_mapping(compiler_workspace: Path):
    template = _copy(compiler_workspace, "output-columns.xlsx")

    def mutate(workbook):
        sheet = workbook["05_HARA"]
        sheet["Q4"], sheet["S5"] = sheet["S5"].value, sheet["Q4"].value

    _mutate(template, mutate)
    original = _compile(TEMPLATE)
    changed = _compile(template)
    fields = {item.canonical_field: item.column for item in changed.report_contract.hara_fields}

    assert fields["asil"] == "S"
    assert fields["safety_goal"] == "Q"
    assert len(changed.severity.rules) == len(original.severity.rules)
    assert len(changed.asil.mappings) == len(original.asil.mappings)


def test_exposure_example_cannot_override_normative_rules(compiler_workspace: Path):
    template = _copy(compiler_workspace, "exposure-example.xlsx")
    _mutate(template, lambda wb: setattr(wb["Exposure"]["C21"], "value", "- Parking example -> E1"))
    method = _compile(template)

    example = next(item for item in method.exposure.examples if "Parking example" in item.raw_text)
    assert example.result == "E1"
    assert example.executable is False
    assert example.parse_status is ParseStatus.NON_EXECUTABLE
    assert all("Parking example" not in item.raw_text for item in method.exposure.duration_rules)
    assert all("Parking example" not in item.raw_text for item in method.exposure.frequency_rules)


def test_controllability_example_cannot_become_keyword_rule(compiler_workspace: Path):
    template = _copy(compiler_workspace, "control-example.xlsx")

    def mutate(workbook):
        sheet = workbook["Controllability"]
        sheet["B25"] = "Example for brake failure added by test"
        sheet["D25"] = "Steer and stop"

    _mutate(template, mutate)
    method = _compile(template)
    example = next(item for item in method.controllability.examples if "added by test" in item.raw_text)

    assert example.result == "C1"
    assert example.executable is False
    assert len(method.controllability.criteria) == 4


def test_conflicting_normative_severity_rules_fail_closed(compiler_workspace: Path):
    template = _copy(compiler_workspace, "severity-conflict.xlsx")
    _mutate(template, lambda wb: setattr(wb["Severity"]["I27"], "value", "Rear-end collision (v<15): S3"))
    method = _compile(template)

    assert method.compile_status is CompileStatus.NOT_READY
    assert method.engineering_rules_compiled is False
    assert any(
        item.code is CompilerDiagnosticCode.RULE_CONFLICT and item.blocking
        for item in method.diagnostics
    )


def test_ambiguous_severity_results_are_preserved():
    method = _compile(TEMPLATE)
    rule = next(item for item in method.severity.rules if "S0 (S1)" in item.raw_text)

    assert rule.result is None
    assert rule.alternatives == ("S0", "S1")
    assert rule.parse_status is ParseStatus.NEEDS_REVIEW


def test_normative_rules_round_trip_to_workbook_sources():
    method = _compile(TEMPLATE)
    samples = (
        method.severity.rules[0],
        next(item for item in method.exposure.duration_rules if item.executable),
        next(item for item in method.controllability.criteria if item.executable),
    )
    workbook = load_workbook(TEMPLATE, read_only=False, data_only=False)
    try:
        for rule in samples:
            assert rule.source_refs
            source = rule.source_refs[0]
            assert source.template_hash == method.metadata["template_hash"]
            assert source.sheet in workbook.sheetnames
            assert source.raw_text
        mapping = method.asil.mappings[37]
        assert mapping.result in mapping.source_ref.raw_text
        guideword = method.guidewords.guidewords[3]
        assert guideword.name in guideword.source_ref.raw_text
    finally:
        workbook.close()


def test_every_executable_normative_rule_has_source_ref():
    method = _compile(TEMPLATE)
    assert all(
        rule.source_refs
        for rule in method.all_rules()
        if rule.executable and rule.normative_strength in {
            NormativeStrength.NORMATIVE, NormativeStrength.CRITERION
        }
    )


def test_golden_contract_is_a_hash_validated_regression_fixture():
    import json

    compiled = _compile(TEMPLATE)
    golden_path = ROOT / "tests" / "fixtures" / "method_contract" / (
        f"{compiled.metadata['template_hash']}.method-contract.golden.json"
    )
    golden = json.loads(golden_path.read_text(encoding="utf-8"))

    assert golden == compiled.audit_snapshot()


def test_new_compiler_is_domain_independent():
    source_paths = tuple((ROOT / "src" / "hara_agent" / "template").glob("*.py")) + tuple(
        (ROOT / "src" / "hara_agent" / "contracts").glob("method*.py")
    )
    source = "\n".join(path.read_text(encoding="utf-8") for path in source_paths)

    for forbidden in (
        "hara_agent.domains", "DomainRegistry", "DomainProfile", "AVPDomainPolicy",
        "DomainScoringService", "SafetyGoalCatalogService", "SG_AVP_", "profile.json",
    ):
        assert forbidden not in source
