from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parent.parent / "output" / "HARA_V13_输出效果对比汇报.docx"
BLUE = "2E74B5"
DARK = "1F4D78"
LIGHT = "E8EEF5"
GRAY = "F2F4F7"
RED = "9B1C1C"


def set_font(run, size=11, bold=False, color="000000"):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for side, val in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                node = OxmlElement(f"w:{side}")
                node.set(qn("w:w"), str(val))
                node.set(qn("w:type"), "dxa")
                mar.append(node)
            tc_pr.append(mar)


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True, color=DARK)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size=16 if level == 1 else 13, bold=True, color=BLUE if level == 1 else DARK)
    return p


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for level, size, color in ((1, 16, BLUE), (2, 13, BLUE), (3, 12, DARK)):
        st = doc.styles[f"Heading {level}"]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        st.paragraph_format.space_after = Pt(6)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("HARA V13 · 输出效果评估")
    set_font(r, size=9, color="666666")
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("内部技术汇报 | 2026-08-04")
    set_font(r, size=9, color="777777")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("HARA V13 输出效果对比汇报")
    set_font(r, size=23, bold=True, color="000000")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("基于标准AVP HARA/Safety Goal输出与当前自动化结果的内容差距分析")
    set_font(r, size=12, color="555555")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.12)
    r = p.add_run("结论：")
    set_font(r, bold=True, color=RED)
    r = p.add_run("当前工程已具备完整自动化流程，但内容仍偏向“通用HAZOP排列组合”，尚未达到“AVP车辆运动控制风险分析”的专业粒度。下一阶段应减少无效数量，提升功能识别、场景定量和Safety Goal收敛能力。")
    set_font(r)

    add_heading(doc, "1. 输出效果概览", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["对比项", "标准输出", "当前输出"]
    for i, text in enumerate(headers):
        shade(table.rows[0].cells[i], LIGHT)
        p = table.rows[0].cells[i].paragraphs[0]
        r = p.add_run(text)
        set_font(r, bold=True, color=DARK)
    data = [
        ("分析规模", "57条HARA记录", "2,352条HARA记录"),
        ("主要功能", "7类车辆控制功能", "14条需求句被识别为功能"),
        ("失效形式", "26类物理失效", "159种文本组合"),
        ("Safety Goal", "8条系统级目标", "547条场景级目标"),
        ("风险定量", "速度/距离/TTC/碰撞速度", "主要依赖通用场景规则"),
    ]
    for a, b, c in data:
        cells = table.add_row().cells
        for i, text in enumerate((a, b, c)):
            p = cells[i].paragraphs[0]
            r = p.add_run(text)
            set_font(r, bold=(i == 0))
    set_table_geometry(table, [1800, 3780, 3780])

    add_heading(doc, "2. 主要内容差距", 1)
    add_bullet(doc, "功能层级：标准输出围绕制动、驱动、转向、驻车、激活/退出等车辆接口；当前输出混入大量需求描述和章节文本。", "功能层级：")
    add_bullet(doc, "失效模式：标准输出采用丢失、过大、过小、反向、固定值、非预期激活等物理失效；当前输出机械套用14个引导词。", "失效模式：")
    add_bullet(doc, "场景选择：标准输出关注泊入方向、目标物、距离和相对速度；当前输出对通用道路场景做大规模全排列。", "场景选择：")
    add_bullet(doc, "S/E/C依据：标准输出计算碰撞时间、碰撞速度并考虑驾驶员不在车内；当前规则仍偏普通驾驶场景。", "S/E/C依据：")
    add_bullet(doc, "Safety Goal：标准输出按车辆级危害聚类；当前按每个非QM场景生成，数量过多且难以评审。", "Safety Goal：")
    add_bullet(doc, "Safe State：当前AVP结果错误调用Wiper参考规则，属于必须修复的子系统路由问题。", "Safe State：")

    doc.add_page_break()
    add_heading(doc, "3. 建议补充方向", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, text in enumerate(("优先级", "补充能力", "预期效果")):
        shade(table.rows[0].cells[i], LIGHT)
        p = table.rows[0].cells[i].paragraphs[0]
        r = p.add_run(text)
        set_font(r, bold=True, color=DARK)
    roadmap = [
        ("P0", "功能归一化与AVP参考库路由", "只分析正确的车辆级功能，消除Wiper等错误结果"),
        ("P0", "不适用组合终止逻辑", "不再对无效组合继续生成场景、评分和Safety Goal"),
        ("P0", "Safety Goal聚类与去重", "由数百条场景目标收敛到可评审的系统级目标"),
        ("P1", "AVP专属场景库", "覆盖前/后/侧目标物、泊入方向、驾驶员位置和ODD"),
        ("P1", "运动学与碰撞计算", "输出TTC、碰撞速度、停车距离及故障后移动距离"),
        ("P1", "AVP专属S/E/C与FTTI", "形成可追溯、可审核的评分依据和响应时间要求"),
    ]
    for idx, (priority, capability, effect) in enumerate(roadmap):
        cells = table.add_row().cells
        if priority == "P0":
            shade(cells[0], "FCE8E6")
        for i, text in enumerate((priority, capability, effect)):
            p = cells[i].paragraphs[0]
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            set_font(r, bold=(i == 0), color=RED if priority == "P0" and i == 0 else "000000")
    set_table_geometry(table, [1100, 3100, 5160])

    add_heading(doc, "4. 改进后的目标输出", 1)
    add_bullet(doc, "功能：从需求长句转换为制动、驱动、转向、驻车、功能状态、感知与规划等原子功能。")
    add_bullet(doc, "HARA：每条记录包含相关场景、目标物、速度、距离、TTC、碰撞速度及可控措施。")
    add_bullet(doc, "Safety Goal：按车辆级危害收敛，并输出Max ASIL、Min FTTI和AVP专属Safe State。")
    add_bullet(doc, "结果规模：减少无效全排列，输出数量可控、内容可解释、结论可追溯。")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("建议实施顺序：先完成P0内容质量修复，再建设P1定量分析能力。")
    set_font(r, bold=True, color=DARK)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
