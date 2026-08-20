#!/usr/bin/env python3
"""
Malfunction逻辑检查器
用于检查功能与引导词组合的合理性
"""

import re
import json
import sys
import os
from typing import Dict, List, Any, Tuple, Set
import logging
import argparse
from collections import Counter

# 设置日志
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class MalfunctionChecker:
    """Malfunction逻辑检查器"""
    
    def __init__(self):
        """初始化检查器"""
        # 不合理的组合模式
        self.invalid_patterns = [
            # 功能与引导词语义冲突
            (r"关断.*(always active|too long|too early)", "功能为关断时，不应与持续激活、过长、过早等引导词组合"),
            (r"开启.*(loss|too short|too late)", "功能为开启时，不应与丢失、过短、过晚等引导词组合"),
            (r"保持.*(loss|different to)", "功能为保持时，不应与丢失、不同等引导词组合"),
            (r"增加.*(loss|too small|too short)", "功能为增加时，不应与丢失、过小、过短等引导词组合"),
            (r"减少.*(too large|too long)", "功能为减少时，不应与过大、过长等引导词组合"),
            
            # 引导词之间的逻辑冲突
            (r"(too early.*too late|too late.*too early)", "过早和过晚不能同时存在"),
            (r"(too large.*too small|too small.*too large)", "过大和过小不能同时存在"),
            (r"(too fast.*too slow|too slow.*too fast)", "过快和过慢不能同时存在"),
            (r"(too long.*too short|too short.*too long)", "过长和过短不能同时存在"),
            
            # 物理不可能的组合
            (r"位置.*(too large|too small)", "位置参数通常不与大小引导词组合"),
            (r"时间.*(too large|too small)", "时间参数通常不与大小引导词组合"),
            (r"速度.*(too early|too late)", "速度通常不与时间引导词组合"),
            (r"温度.*(too fast|too slow)", "温度通常不与速度引导词组合"),
        ]
        
        # 功能类型分类
        self.function_categories = {
            # Put the more specific categories first.  The old fallback treated
            # torque/control outputs as "其他类", which effectively made every
            # guideword applicable and recreated a Cartesian product.
            "执行器输出类": ["扭矩", "制动", "驱动", "转向", "减速度", "加速度", "转角"],
            "提示类": ["报警", "提示", "显示", "告警", "警告"],
            "开关类": ["开启", "关闭", "退出", "关断", "启动", "停止", "激活", "禁用", "使能", "禁止"],
            "调节类": ["增加", "减少", "调节", "调整", "改变", "修改", "优化"],
            "保持类": ["保持", "维持", "稳定", "恒定", "固定"],
            "位置类": ["移动", "定位", "放置", "位置", "坐标", "方向"],
            "时间类": ["定时", "延时", "周期", "频率", "时间", "时刻"],
            "速度类": ["加速", "减速", "速度", "速率", "快速", "慢速"],
            "温度类": ["加热", "冷却", "温度", "热", "冷"],
            "压力类": ["加压", "减压", "压力", "压强"],
            "信号类": ["信号", "通信", "传输", "接收", "发送", "广播"],
            "状态类": ["状态", "模式", "条件", "情况", "情形"],
        }
        
        # 14个标准HAZOP引导词（必须全部使用）
        self.standard_guidewords = [
            "unintended",
            "always active",
            "loss",
            "too large",
            "too small",
            "too early",
            "too late",
            "too fast",
            "too slow",
            "too long",
            "too short",
            "incomplete",
            "different to",
            "as well as"
        ]
        
        # 引导词分类
        self.guideword_categories = {
            "存在性": ["unintended", "always active", "loss", "as well as"],
            "量值": ["too large", "too small"],
            "时间": ["too early", "too late", "too long", "too short"],
            "速度": ["too fast", "too slow"],
            "完整性": ["incomplete"],
            "差异性": ["different to"],
        }
        
        # 合理的组合规则
        self.valid_combinations = {
            "执行器输出类": [
                "unintended", "always active", "loss", "too large", "too small",
                "too early", "too late", "too fast", "too slow", "incomplete",
                "different to"
            ],
            "提示类": ["unintended", "loss", "too early", "too late", "incomplete", "different to"],
            "开关类": ["unintended", "always active", "loss", "too early", "too late", "different to"],
            "调节类": ["too large", "too small", "too fast", "too slow", "different to"],
            "保持类": ["unintended", "always active", "loss", "different to"],
            "位置类": ["too large", "too small", "too early", "too late", "different to"],
            "时间类": ["too early", "too late", "too long", "too short", "too fast", "too slow"],
            "速度类": ["too fast", "too slow", "too early", "too late"],
            "温度类": ["too large", "too small", "too fast", "too slow"],
            "压力类": ["too large", "too small", "too fast", "too slow"],
            "信号类": ["unintended", "loss", "incomplete", "different to"],
            "状态类": ["unintended", "always active", "loss", "different to"],
        }
    
    def check_function_guideword_combo(self, function: str, guideword: str) -> Dict[str, Any]:
        """
        检查功能与引导词组合的合理性
        
        Args:
            function: 功能描述
            guideword: 引导词
            
        Returns:
            检查结果，包含是否合理、原因和建议
        """
        result = {
            "is_valid": True,
            "function": function,
            "guideword": guideword,
            "reasons": [],
            "suggestions": [],
            "function_category": None,
            "guideword_category": None
        }
        
        # 识别功能类别
        function_category = self._categorize_function(function)
        result["function_category"] = function_category
        
        # 识别引导词类别
        guideword_category = self._categorize_guideword(guideword)
        result["guideword_category"] = guideword_category
        
        # 检查1：基本合理性检查
        if not function or not guideword:
            result["is_valid"] = False
            result["reasons"].append("功能或引导词为空")
            return result
        
        # 检查2：检查不合理的模式
        combo_text = f"{function} {guideword}"
        for pattern, reason in self.invalid_patterns:
            if re.search(pattern, combo_text, re.IGNORECASE):
                result["is_valid"] = False
                result["reasons"].append(f"匹配到不合理模式: {reason}")
        
        # 检查3：检查类别匹配性
        if function_category and guideword_category:
            # 检查是否在有效组合中
            if function_category in self.valid_combinations:
                valid_guidewords = self.valid_combinations[function_category]
                if guideword not in valid_guidewords:
                    result["is_valid"] = False
                    result["reasons"].append(f"功能类别'{function_category}'通常不与引导词'{guideword}'组合")
                    result["suggestions"].append(f"建议使用: {', '.join(valid_guidewords)}")
            elif function_category == "其他类":
                # Unknown functions are still assessed, but only guidewords with
                # broadly valid action semantics may proceed automatically.
                generic_guidewords = {
                    "unintended", "always active", "loss", "too early", "too late",
                    "incomplete", "different to"
                }
                if guideword not in generic_guidewords:
                    result["is_valid"] = False
                    result["reasons"].append(
                        f"未识别功能类别，不能自动确认引导词'{guideword}'适用"
                    )
                    result["suggestions"].append("补充功能类型或由人工确认适用性")
            
            # 检查类别兼容性
            category_compatibility = self._check_category_compatibility(function_category, guideword_category)
            if not category_compatibility["compatible"]:
                result["is_valid"] = False
                result["reasons"].append(category_compatibility["reason"])
        
        # 检查4：语义合理性检查
        semantic_check = self._check_semantic_compatibility(function, guideword)
        if not semantic_check["compatible"]:
            result["is_valid"] = False
            result["reasons"].append(semantic_check["reason"])
        
        # 如果检查通过，添加积极反馈
        if result["is_valid"]:
            result["reasons"].append(f"功能'{function}'与引导词'{guideword}'组合合理")
        
        return result
    
    def generate_malfunction_description(self, function: str, guideword: str, 
                                        guideword_desc: str = "") -> str:
        """
        生成Malfunction描述
        
        Args:
            function: 功能描述
            guideword: 引导词
            guideword_desc: 引导词解释
            
        Returns:
            Malfunction描述
        """
        # 基础组合
        if guideword.lower() == "unintended":
            return f"非预期触发{function}"
        elif guideword.lower() == "always active":
            return f"{function}持续激活"
        elif guideword.lower() == "loss":
            return f"{function}丢失"
        elif guideword.lower() == "too large":
            return f"{function}过大"
        elif guideword.lower() == "too small":
            return f"{function}过小"
        elif guideword.lower() == "too early":
            return f"{function}过早"
        elif guideword.lower() == "too late":
            return f"{function}过晚"
        elif guideword.lower() == "too fast":
            return f"{function}过快"
        elif guideword.lower() == "too slow":
            return f"{function}过慢"
        elif guideword.lower() == "too long":
            return f"{function}过长"
        elif guideword.lower() == "too short":
            return f"{function}过短"
        elif guideword.lower() == "incomplete":
            return f"{function}不完整"
        elif guideword.lower() == "different to":
            return f"{function}不同"
        elif guideword.lower() == "as well as":
            return f"{function}同时发生其他动作"
        else:
            return f"{guideword} {function}"
    
    def check_multiple_combinations(self, functions: List[str], guidewords: List[Dict[str, str]]) -> List[Dict[str, Any]]:
        """
        检查多个功能与引导词的组合
        
        Args:
            functions: 功能列表
            guidewords: 引导词列表，每个元素包含guideword和description
            
        Returns:
            检查结果列表
        """
        results = []
        
        for function in functions:
            for guideword_info in guidewords:
                guideword = guideword_info["guideword"]
                guideword_desc = guideword_info.get("description", "")
                
                # 检查组合合理性
                check_result = self.check_function_guideword_combo(function, guideword)
                
                # 生成Malfunction描述
                if check_result["is_valid"]:
                    malfunction_desc = self.generate_malfunction_description(function, guideword, guideword_desc)
                else:
                    malfunction_desc = f"不合理组合: {function} + {guideword}"
                
                result = {
                    "function": function,
                    "guideword": guideword,
                    "guideword_description": guideword_desc,
                    "check_result": check_result,
                    "malfunction_description": malfunction_desc,
                    "should_include": check_result["is_valid"]
                }
                
                results.append(result)
        
        return results
    
    def _categorize_function(self, function: str) -> str:
        """
        将功能分类
        
        Args:
            function: 功能描述
            
        Returns:
            功能类别
        """
        function_lower = function.lower()
        
        for category, keywords in self.function_categories.items():
            for keyword in keywords:
                if keyword in function_lower:
                    return category
        
        # 默认类别
        return "其他类"
    
    def _categorize_guideword(self, guideword: str) -> str:
        """
        将引导词分类
        
        Args:
            guideword: 引导词
            
        Returns:
            引导词类别
        """
        guideword_lower = guideword.lower()
        
        for category, guidewords in self.guideword_categories.items():
            if guideword_lower in [gw.lower() for gw in guidewords]:
                return category
        
        # 默认类别
        return "其他类"
    
    def _check_category_compatibility(self, function_category: str, guideword_category: str) -> Dict[str, Any]:
        """
        检查功能类别与引导词类别的兼容性
        
        Args:
            function_category: 功能类别
            guideword_category: 引导词类别
            
        Returns:
            兼容性检查结果
        """
        result = {
            "compatible": True,
            "reason": ""
        }
        
        # 不兼容的组合
        incompatible_pairs = [
            ("开关类", "量值"),  # 开关通常不与量值组合
            ("位置类", "速度"),  # 位置通常不与速度组合
            ("时间类", "量值"),  # 时间通常不与量值组合
            ("温度类", "时间"),  # 温度通常不与时间组合（除非是变化率）
            ("压力类", "完整性"),  # 压力通常不与完整性组合
        ]
        
        for func_cat, guide_cat in incompatible_pairs:
            if function_category == func_cat and guideword_category == guide_cat:
                result["compatible"] = False
                result["reason"] = f"功能类别'{function_category}'与引导词类别'{guideword_category}'通常不兼容"
                return result
        
        return result
    
    def _check_semantic_compatibility(self, function: str, guideword: str) -> Dict[str, Any]:
        """
        检查语义兼容性
        
        Args:
            function: 功能描述
            guideword: 引导词
            
        Returns:
            语义兼容性检查结果
        """
        result = {
            "compatible": True,
            "reason": ""
        }
        
        function_lower = function.lower()
        guideword_lower = guideword.lower()
        
        # 特定语义检查规则
        semantic_rules = [
            # 功能包含"关闭"但引导词是"持续激活" - 不合理
            ("关闭", "always active", "关闭功能不应持续激活"),
            ("停止", "always active", "停止功能不应持续激活"),
            ("禁用", "always active", "禁用功能不应持续激活"),
            
            # 功能包含"开启"但引导词是"丢失" - 不合理
            ("开启", "loss", "开启功能不应丢失"),
            ("启动", "loss", "启动功能不应丢失"),
            ("激活", "loss", "激活功能不应丢失"),
            
            # 功能包含"保持"但引导词是"不同" - 可能不合理
            ("保持", "different to", "保持功能变为不同可能不合理"),
            ("恒定", "different to", "恒定功能变为不同可能不合理"),
            
            # 功能包含"增加"但引导词是"过小" - 不合理
            ("增加", "too small", "增加功能不应过小"),
            ("增大", "too small", "增大功能不应过小"),
            
            # 功能包含"减少"但引导词是"过大" - 不合理
            ("减少", "too large", "减少功能不应过大"),
            ("减小", "too large", "减小功能不应过大"),
        ]
        
        for func_keyword, guide_keyword, reason in semantic_rules:
            if func_keyword in function_lower and guide_keyword in guideword_lower:
                result["compatible"] = False
                result["reason"] = reason
                return result
        
        return result
    
    def get_standard_guidewords(self) -> List[str]:
        """
        获取14个标准HAZOP引导词
        
        Returns:
            14个标准引导词列表
        """
        return self.standard_guidewords.copy()
    
    def check_guideword_completeness(self, used_guidewords: List[str]) -> Dict[str, Any]:
        """
        检查HARA报告中是否使用了所有14个引导词
        
        Args:
            used_guidewords: 实际使用的引导词列表
            
        Returns:
            完整性检查结果，包含缺失的引导词和建议
        """
        # 标准化已使用的引导词
        used_normalized = set()
        for gw in used_guidewords:
            if gw:
                gw_lower = gw.lower().strip()
                used_normalized.add(gw_lower)
        
        # 标准化标准引导词
        standard_normalized = set(gw.lower() for gw in self.standard_guidewords)
        
        # 找出缺失的引导词
        missing = standard_normalized - used_normalized
        used = used_normalized & standard_normalized
        
        result = {
            "complete": len(missing) == 0,
            "total_standard": len(self.standard_guidewords),
            "used_count": len(used),
            "missing_count": len(missing),
            "used_guidewords": sorted(list(used)),
            "missing_guidewords": sorted(list(missing)),
            "message": "",
            "suggestion": ""
        }
        
        if result["complete"]:
            result["message"] = f"✓ 所有14个标准HAZOP引导词均已使用 ({len(used)}/{len(self.standard_guidewords)})"
        else:
            result["message"] = f"✗ 缺少 {len(missing)} 个标准引导词 ({len(used)}/{len(self.standard_guidewords)})"
            result["suggestion"] = f"建议添加以下引导词的HARA分析: {', '.join(sorted(missing))}"
        
        return result
    
    def check_full_combination_completeness(self, hara_records: List[Dict[str, Any]],
                                            function_column: str = "function",
                                            guideword_column: str = "guideword") -> Dict[str, Any]:
        """
        检查功能和14个引导词的全组合是否完整
        
        **重要**: 本检查要求每个功能必须与所有14个引导词进行组合分析。
        如果不满足全组合要求，需要回到HARA-engine继续补充。
        
        Args:
            hara_records: HARA记录列表
            function_column: 功能列名
            guideword_column: 引导词列名
            
        Returns:
            全组合完整性检查结果
        """
        # 收集所有功能
        all_functions = set()
        for record in hara_records:
            func = record.get(function_column, "")
            if func:
                all_functions.add(func.strip())
        
        # 收集每个功能使用的引导词
        function_guidewords = {}
        for record in hara_records:
            func = record.get(function_column, "")
            gw = record.get(guideword_column, "")
            if func and gw:
                func = func.strip()
                if func not in function_guidewords:
                    function_guidewords[func] = set()
                function_guidewords[func].add(gw.lower().strip())
        
        # 检查每个功能是否使用了所有14个引导词
        full_combination_required = True  # 必须全组合
        incomplete_functions = []
        missing_combinations = []
        
        for func in all_functions:
            func_gws = function_guidewords.get(func, set())
            standard_gws = set(gw.lower() for gw in self.standard_guidewords)
            missing_gws = standard_gws - func_gws
            
            if missing_gws:
                incomplete_functions.append({
                    "function": func,
                    "used_count": len(func_gws),
                    "required_count": 14,
                    "missing_guidewords": sorted(list(missing_gws))
                })
                for gw in missing_gws:
                    missing_combinations.append({
                        "function": func,
                        "guideword": gw
                    })
        
        # 计算全组合率
        if all_functions:
            expected_combinations = len(all_functions) * 14
            actual_combinations = sum(len(gws) for gws in function_guidewords.values())
            full_combination_rate = round(actual_combinations / expected_combinations * 100, 2)
        else:
            expected_combinations = 0
            actual_combinations = 0
            full_combination_rate = 0.0
        
        result = {
            "full_combination_required": full_combination_required,
            "total_functions": len(all_functions),
            "complete_functions": len(all_functions) - len(incomplete_functions),
            "incomplete_functions_count": len(incomplete_functions),
            "missing_count": len(missing_combinations),
            "expected_combinations": expected_combinations,
            "actual_combinations": actual_combinations,
            "full_combination_rate": full_combination_rate,
            "incomplete_functions": incomplete_functions,
            "missing_combinations": missing_combinations,
            "complete": len(incomplete_functions) == 0,
            "message": "",
            "suggestion": ""
        }
        
        if result["complete"]:
            result["message"] = f"✓ 功能和14个引导词全组合完整 ({len(all_functions)}个功能 × 14个引导词 = {expected_combinations}组合)"
        else:
            result["message"] = f"✗ 功能和引导词未全组合: {len(incomplete_functions)}个功能缺少引导词 ({full_combination_rate}%)"
            result["suggestion"] = (
                f"【必须回到HARA-engine继续补充】\n"
                f"缺少 {len(missing_combinations)} 个功能-引导词组合:\n"
                f"提示: 建议添加缺失的功能-引导词组合到HARA分析中"
            )
        
        return result
    
    def check_excel_full_combination_completeness(self, excel_path: str,
                                                  sheet_name: str = "05_HARA",
                                                  function_column: int = 2,
                                                  guideword_column: int = 4) -> Dict[str, Any]:
        """
        检查Excel中HARA报告的功能-引导词全组合完整性
        
        **重要**: 检查每个功能是否与所有14个标准引导词进行了组合分析。
        如果不满足全组合要求，必须回到HARA-engine继续补充。
        
        Args:
            excel_path: Excel文件路径
            sheet_name: HARA sheet名称
            function_column: 功能列号（B列=2）
            guideword_column: 引导词列号（D列=4）
            
        Returns:
            全组合完整性检查结果
        """
        try:
            import openpyxl
            
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            ws = wb[sheet_name]
            
            # 收集HARA记录
            hara_records = []
            for row in range(6, ws.max_row + 1):
                func = ws.cell(row, function_column).value
                gw = ws.cell(row, guideword_column).value
                if func:
                    hara_records.append({
                        "function": func,
                        "guideword": gw if gw else ""
                    })
            
            wb.close()
            
            return self.check_full_combination_completeness(hara_records)
            
        except Exception as e:
            return {
                "complete": False,
                "full_combination_required": True,
                "error": str(e),
                "message": f"检查失败: {e}",
                "suggestion": "【必须回到HARA-engine继续补充】请修复Excel文件后重新检查"
            }
    
    def check_hara_report_completeness(self, hara_records: List[Dict[str, Any]], 
                                        guideword_column: str = "guideword") -> Dict[str, Any]:
        """
        检查HARA报告的引导词完整性
        
        Args:
            hara_records: HARA记录列表（每个记录是一个字典）
            guideword_column: 引导词列名（默认为"guideword"）
            
        Returns:
            完整性检查结果
        """
        # 从HARA记录中提取使用的引导词
        used_guidewords = []
        for record in hara_records:
            if guideword_column in record:
                gw = record[guideword_column]
                if isinstance(gw, list):
                    used_guidewords.extend(gw)
                else:
                    used_guidewords.append(gw)
        
        return self.check_guideword_completeness(used_guidewords)
    
    def check_excel_hara_completeness(self, excel_path: str, 
                                      sheet_name: str = "05_HARA",
                                      guideword_column: int = 4) -> Dict[str, Any]:
        """
        检查Excel中HARA报告的引导词完整性
        
        Args:
            excel_path: Excel文件路径
            sheet_name: HARA sheet名称
            guideword_column: 引导词列号（D列=4）
            
        Returns:
            完整性检查结果
        """
        try:
            import openpyxl
            
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            ws = wb[sheet_name]
            
            used_guidewords = []
            
            # 从第6行开始扫描（第5行是表头）
            for row in range(6, ws.max_row + 1):
                gw = ws.cell(row, guideword_column).value
                if gw:
                    used_guidewords.append(str(gw).strip())
            
            wb.close()
            
            return self.check_guideword_completeness(used_guidewords)
            
        except Exception as e:
            return {
                "complete": False,
                "error": str(e),
                "message": f"检查失败: {e}"
            }
    
    def check(self, function: str, guideword: str) -> Dict[str, Any]:
        """
        检查功能与引导词组合并返回标准化检查结果

        Returns:
            - 检查通过: {"status": "pass", "errors": []}
            - 检查失败: {"status": "fail", "errors": [...]}
        """
        result = self.check_function_guideword_combo(function, guideword)
        errors = []

        if not result["is_valid"]:
            errors.append({
                "function": function,
                "guideword": guideword,
                "reasons": result.get("reasons", []),
                "suggestions": result.get("suggestions", [])
            })

        if errors:
            return {"status": "fail", "errors": errors}
        return {"status": "pass", "errors": []}

    def filter_valid_combinations(self, combinations: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """
        过滤出有效的组合

        Args:
            combinations: 组合列表

        Returns:
            有效组合列表
        """
        return [combo for combo in combinations if combo["check_result"]["is_valid"]]
    
    def export_check_results(self, results: List[Dict[str, Any]], output_path: str) -> bool:
        """
        导出检查结果
        
        Args:
            results: 检查结果列表
            output_path: 输出文件路径
            
        Returns:
            是否导出成功
        """
        try:
            # 统计信息
            total_combinations = len(results)
            valid_combinations = len([r for r in results if r["check_result"]["is_valid"]])
            invalid_combinations = total_combinations - valid_combinations
            
            summary = {
                "total_combinations": total_combinations,
                "valid_combinations": valid_combinations,
                "invalid_combinations": invalid_combinations,
                "valid_percentage": round(valid_combinations / total_combinations * 100, 2) if total_combinations > 0 else 0,
                "results": results
            }
            
            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(summary, f, ensure_ascii=False, indent=2)
            
            logger.info(f"检查结果已导出到: {output_path}")
            logger.info(f"总组合数: {total_combinations}, 有效组合: {valid_combinations}, 无效组合: {invalid_combinations}")
            
            return True
            
        except Exception as e:
            logger.error(f"导出检查结果失败: {e}")
            return False


def test_malfunction_checker(input_path: str, verbose: bool = False) -> Dict[str, Any]:
    """
    测试 Malfunction 检查器 - 从 JSON 文件加载数据并检查
    
    Args:
        input_path: malfunction_output.json 文件路径
        verbose: 是否显示详细输出
        
    Returns:
        检查结果字典 {"passed": bool, "errors": [...], "summary": {...}}
    """
    checker = MalfunctionChecker()
    errors = []
    
    if not os.path.exists(input_path):
        return {
            "passed": False,
            "errors": [{"type": "file_not_found", "message": f"文件不存在：{input_path}"}]
        }
    
    try:
        with open(input_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
    except json.JSONDecodeError as e:
        return {
            "passed": False,
            "errors": [{"type": "json_parse_error", "message": f"JSON 解析失败：{e}"}]
        }
    
    functions = data.get("functions", [])
    guidewords_info = data.get("function_guidewords", {})
    
    if not functions:
        return {
            "passed": False,
            "errors": [{"type": "no_data", "message": "未找到 functions 数据"}]
        }
    
    guidewords = []
    for gw in checker.get_standard_guidewords():
        guidewords.append({"guideword": gw, "description": ""})
    
    results = checker.check_multiple_combinations(functions, guidewords)

    # All 14 guidewords must be assessed, but only applicable, non-duplicate
    # malfunctions may be passed downstream.  This separates audit
    # completeness from Cartesian-product generation.
    actual_malfunctions = data.get("function_malfunctions", {})
    assessments = data.get("function_guideword_assessments", {})
    invalid_results = []
    valid_statuses = {"applicable", "not_applicable", "duplicate"}

    if assessments:
        standard_set = set(checker.get_standard_guidewords())
        for function in functions:
            items = assessments.get(function, []) or []
            by_guideword = {item.get("guideword"): item for item in items}
            missing = sorted(standard_set - set(by_guideword))
            if missing:
                errors.append({
                    "type": "guideword_assessment_incomplete",
                    "function": function,
                    "missing_guidewords": missing,
                })

            for guideword, assessment in by_guideword.items():
                status = assessment.get("status")
                if status not in valid_statuses:
                    errors.append({
                        "type": "invalid_analysis_status",
                        "function": function,
                        "guideword": guideword,
                        "status": status,
                    })
                if status == "not_applicable" and not assessment.get("reason"):
                    errors.append({
                        "type": "missing_not_applicable_reason",
                        "function": function,
                        "guideword": guideword,
                    })
                if status == "duplicate" and not assessment.get("covered_by"):
                    errors.append({
                        "type": "missing_duplicate_target",
                        "function": function,
                        "guideword": guideword,
                    })

            expected_downstream = {
                gw for gw, item in by_guideword.items()
                if item.get("status") == "applicable"
            }
            actual_downstream = {
                item.get("guideword") for item in actual_malfunctions.get(function, []) or []
                if item.get("analysis_status", "applicable") == "applicable"
            }
            if actual_downstream != expected_downstream:
                errors.append({
                    "type": "downstream_applicability_mismatch",
                    "function": function,
                    "missing": sorted(expected_downstream - actual_downstream),
                    "unexpected": sorted(actual_downstream - expected_downstream),
                })
    else:
        # Backward-compatible validation for legacy outputs.
        invalid_results = [r for r in results if not r["check_result"]["is_valid"]]
        for result in invalid_results:
            errors.append({
                "type": "invalid_combination",
                "function": result["function"],
                "guideword": result["guideword"],
                "reasons": result["check_result"]["reasons"],
                "suggestions": result["check_result"]["suggestions"]
            })
    
    guideword_completeness = checker.check_guideword_completeness(
        [gw["guideword"] for gw in guidewords]
    )
    
    if not guideword_completeness.get("complete", False):
        errors.append({
            "type": "guideword_incomplete",
            "missing_count": guideword_completeness.get("missing_count", 0),
            "missing_guidewords": guideword_completeness.get("missing_guidewords", []),
            "message": guideword_completeness.get("message", "")
        })
    
    passed = len(errors) == 0
    
    summary = {
        "total_functions": len(functions),
        "total_combinations": len(results),
        "valid_combinations": sum(len(v or []) for v in actual_malfunctions.values()),
        "invalid_combinations": len(invalid_results),
        "assessment_statuses": dict(Counter(
            item.get("status", "missing")
            for items in assessments.values() for item in (items or [])
        )) if assessments else {},
        "guideword_completeness": guideword_completeness
    }
    
    if verbose:
        print(f"\nMalfunction 检查结果:")
        print(f"  功能数量：{len(functions)}")
        print(f"  总组合数：{len(results)}")
        print(f"  下游适用组合：{summary['valid_combinations']}")
        print(f"  无效组合：{len(invalid_results)}")
        if invalid_results:
            print(f"\n  无效组合详情:")
            for err in invalid_results[:10]:
                print(f"    - {err['function']} + {err['guideword']}: {', '.join(err['check_result']['reasons'][:2])}")
    
    return {
        "passed": passed,
        "errors": errors,
        "summary": summary
    }


def main():
    parser = argparse.ArgumentParser(description="Malfunction Checker - 功能异常组合合理性检查")
    parser.add_argument("--input", "-i", type=str, required=True, help="malfunction_output.json 路径")
    parser.add_argument("--output", "-o", type=str, help="输出报告路径 (可选)")
    parser.add_argument("--verbose", "-v", action="store_true", help="详细输出")
    args = parser.parse_args()
    
    result = test_malfunction_checker(args.input, verbose=args.verbose)
    
    if args.output:
        with open(args.output, 'w', encoding='utf-8') as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
        print(f"检查结果已保存到：{args.output}")
    else:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    
    sys.exit(0 if result.get("passed") else 1)


if __name__ == "__main__":
    main()
