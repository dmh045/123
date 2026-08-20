"""ASIL determination backed by the input workbook's ``ASIL_Table`` sheet."""

from pathlib import Path
import re
from typing import Dict, Optional, Tuple

from openpyxl import load_workbook


DEFAULT_TEMPLATE_PATH = (
    Path(__file__).resolve().parents[1]
    / "references"
    / "HARA_Template_AI_20260327.xlsx"
)


class TemplateASILMatrix:
    """Load and query the S/E/C to ASIL matrix from a HARA template."""

    SHEET_NAME = "ASIL_Table"
    EXPECTED_S = tuple(f"S{i}" for i in range(4))
    EXPECTED_E = tuple(f"E{i}" for i in range(5))
    EXPECTED_C = tuple(f"C{i}" for i in range(4))
    VALID_RESULTS = {"QM", "A", "B", "C", "D"}

    def __init__(self, template_path: Optional[str] = None):
        path = Path(template_path) if template_path else DEFAULT_TEMPLATE_PATH
        self.template_path = path.expanduser().resolve()
        self.matrix = self._load()
        self.source = f"{self.template_path}#{self.SHEET_NAME}"

    @staticmethod
    def _code(value: object, prefix: str, maximum: int) -> Optional[str]:
        text = str(value or "").strip().upper()
        match = re.fullmatch(rf"{prefix}([0-{maximum}])", text)
        return match.group(0) if match else None

    @classmethod
    def _normalize_result(cls, value: object) -> str:
        text = str(value or "").strip().upper().replace("ASIL ", "")
        # The template expresses the S0/E0/C0 axes as NA, while HARA output
        # represents those non-ASIL combinations as QM.
        if text in {"NA", "N/A", "-"}:
            return "QM"
        if text not in cls.VALID_RESULTS:
            raise ValueError(f"ASIL_Table contains unsupported result: {value!r}")
        return text

    def _load(self) -> Dict[Tuple[str, str, str], str]:
        if not self.template_path.is_file():
            raise FileNotFoundError(f"ASIL template not found: {self.template_path}")

        workbook = load_workbook(self.template_path, read_only=True, data_only=False)
        try:
            if self.SHEET_NAME not in workbook.sheetnames:
                raise ValueError(
                    f"Template {self.template_path} is missing required sheet "
                    f"{self.SHEET_NAME!r}"
                )
            sheet = workbook[self.SHEET_NAME]

            header_row = None
            c_columns: Dict[str, int] = {}
            for row in range(1, min(sheet.max_row, 30) + 1):
                found: Dict[str, int] = {}
                for col in range(1, sheet.max_column + 1):
                    code = self._code(sheet.cell(row, col).value, "C", 3)
                    if code:
                        found[code] = col
                if set(found) == set(self.EXPECTED_C):
                    header_row = row
                    c_columns = found
                    break
            if header_row is None:
                raise ValueError("ASIL_Table does not contain a complete C0-C3 header")

            matrix: Dict[Tuple[str, str, str], str] = {}
            current_s: Optional[str] = None
            for row in range(header_row + 1, sheet.max_row + 1):
                row_s = None
                row_e = None
                for col in range(1, min(c_columns.values())):
                    value = sheet.cell(row, col).value
                    row_s = row_s or self._code(value, "S", 3)
                    row_e = row_e or self._code(value, "E", 4)
                if row_s:
                    current_s = row_s
                if not current_s or not row_e:
                    continue
                for c_code, col in c_columns.items():
                    matrix[(current_s, row_e, c_code)] = self._normalize_result(
                        sheet.cell(row, col).value
                    )

            expected = {
                (s_code, e_code, c_code)
                for s_code in self.EXPECTED_S
                for e_code in self.EXPECTED_E
                for c_code in self.EXPECTED_C
            }
            missing = sorted(expected - set(matrix))
            extra = sorted(set(matrix) - expected)
            if missing or extra:
                raise ValueError(
                    "ASIL_Table matrix is incomplete or ambiguous: "
                    f"expected={len(expected)}, loaded={len(matrix)}, "
                    f"missing={missing[:5]}, extra={extra[:5]}"
                )
            return matrix
        finally:
            workbook.close()

    def determine(self, severity: str, exposure: str, controllability: str) -> str:
        key = (
            str(severity or "").strip().upper(),
            str(exposure or "").strip().upper(),
            str(controllability or "").strip().upper(),
        )
        if key not in self.matrix:
            raise ValueError(
                "Invalid S/E/C combination for ASIL_Table lookup: "
                f"severity={severity!r}, exposure={exposure!r}, "
                f"controllability={controllability!r}"
            )
        return self.matrix[key]
