import json
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent


class ReportTextClarityTests(unittest.TestCase):
    template = ROOT / "references" / "HARA_Template_AI_20260327.xlsx"
    report = ROOT / "output" / "p018_iav_style_hara_report.xlsx"
    phase2 = ROOT / "output" / "p016_scen_hazevent_output.json"
    phase3 = ROOT / "output" / "p016_scoring_sg_output.json"

    @classmethod
    def setUpClass(cls):
        missing = [path for path in (cls.report, cls.phase2, cls.phase3) if not path.is_file()]
        if missing:
            raise unittest.SkipTest(
                "legacy generated fixtures are absent: " + ", ".join(str(path) for path in missing)
            )
        # Random cell access in openpyxl read-only mode reparses the sheet for
        # every lookup and becomes quadratic once all risk scenarios are exported.
        cls.wb = load_workbook(cls.report, data_only=True, read_only=False)
        cls.ws = cls.wb["05_HARA"]
        cls.rows = [
            row for row in range(6, cls.ws.max_row + 1)
            if cls.ws.cell(row, 1).value not in (None, "")
        ]

    def test_template_structure_and_headers_are_unchanged(self):
        template_wb = load_workbook(self.template, data_only=True, read_only=True)
        self.assertEqual(template_wb.sheetnames, self.wb.sheetnames)
        template_ws = template_wb["05_HARA"]
        for row in (4, 5):
            expected = [template_ws.cell(row, col).value for col in range(1, 22)]
            actual = [self.ws.cell(row, col).value for col in range(1, 22)]
            self.assertEqual(expected, actual)

    def test_display_text_is_concise_and_has_no_template_artifacts(self):
        limits = {7: 55, 8: 80, 9: 60, 11: 45, 14: 40, 16: 40, 21: 40}
        for row in self.rows:
            for col, limit in limits.items():
                value = str(self.ws.cell(row, col).value or "")
                self.assertLessEqual(len(value), limit, (row, col, value))
                self.assertNotIn("\n", value)
            self.assertNotIn("<", str(self.ws.cell(row, 9).value))
            self.assertNotIn(">", str(self.ws.cell(row, 9).value))
            self.assertNotIn("危害事件:", str(self.ws.cell(row, 11).value))
            self.assertNotIn("Event:", str(self.ws.cell(row, 21).value))

    def test_key_scenario_facts_and_score_conclusions_remain_visible(self):
        for row in self.rows:
            scene = str(self.ws.cell(row, 7).value or "")
            detail = str(self.ws.cell(row, 8).value or "")
            self.assertIn("停车场", scene)
            self.assertIn("驾驶员", scene)
            self.assertNotIn("km/h", scene)
            self.assertNotIn("距离=", scene)
            for token in ("自车", "目标", "相对", "距离", "坡度", "路面"):
                self.assertIn(token, detail, (row, token, detail))
            for score_col, reason_col in ((10, 11), (12, 14), (15, 16)):
                score = str(self.ws.cell(row, score_col).value or "")
                reason = str(self.ws.cell(row, reason_col).value or "")
                self.assertIn(score, reason)
                self.assertIn("→", reason)

    def test_repeated_static_columns_are_merged_but_dynamic_columns_are_not(self):
        vertical_merges = [
            cell_range for cell_range in self.ws.merged_cells.ranges
            if cell_range.min_col == cell_range.max_col and cell_range.min_row >= 6
        ]
        merged_columns = {cell_range.min_col for cell_range in vertical_merges}
        for col in (2, 3, 4, 5, 6):
            self.assertIn(col, merged_columns)
        for col in range(8, 22):
            self.assertNotIn(col, merged_columns)

    def test_all_selected_risk_scenarios_are_exported(self):
        phase3 = json.loads(self.phase3.read_text(encoding="utf-8"))
        expected = phase3["metadata"]["total_report_risk_classes"]
        self.assertEqual(expected, len(self.rows))
        ids = [str(self.ws.cell(row, 1).value or "") for row in self.rows]
        self.assertTrue(any(value.endswith(".2") for value in ids))
        self.assertTrue(any(value.endswith(".3") for value in ids))

    def test_scenario_and_scores_stay_joined_by_scenario_id(self):
        phase3 = json.loads(self.phase3.read_text(encoding="utf-8"))
        selected_count = sum(
            len(ids)
            for functions in phase3["report_scenario_ids"].values()
            for malfunctions in functions.values()
            for ids in malfunctions.values()
        )
        self.assertEqual(selected_count, len(self.rows))

    def test_full_evidence_is_preserved_in_json(self):
        phase2 = json.loads(self.phase2.read_text(encoding="utf-8"))
        phase3 = json.loads(self.phase3.read_text(encoding="utf-8"))
        self.assertTrue(any(
            len(str(item.get("situational_detailing", ""))) > 250
            for item in phase2["scenario_catalog"].values()
        ))

        severity_reason = ""
        for functions in phase3["severity_results"].values():
            for malfunctions in functions.values():
                for items in malfunctions.values():
                    if items:
                        severity_reason = str(items[0].get("reasoning", ""))
                        break
                if severity_reason:
                    break
            if severity_reason:
                break
        self.assertGreater(len(severity_reason), 250)


if __name__ == "__main__":
    unittest.main()
