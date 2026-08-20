from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook


class TemplateASILService:
    """Load and query the authoritative ASIL matrix from the run template."""

    SHEET_NAME = "ASIL_Table"
    EXPECTED_S = tuple(f"S{i}" for i in range(4))
    EXPECTED_E = tuple(f"E{i}" for i in range(5))
    EXPECTED_C = tuple(f"C{i}" for i in range(4))
    VALID_RESULTS = {"QM", "A", "B", "C", "D"}

    def __init__(self, template_path: str):
        if not template_path:
            raise ValueError("必须显式提供本次运行的Excel模板，禁止使用默认ASIL矩阵")
        self.template_path = Path(template_path).expanduser().resolve()
        self.matrix = self._load()
        self.source = f"{self.template_path}#{self.SHEET_NAME}"

    @staticmethod
    def _code(value: object, prefix: str, maximum: int) -> Optional[str]:
        text = str(value or "").strip().upper()
        match = re.fullmatch(rf"{prefix}([0-{maximum}])", text)
        return match.group(0) if match else None

    @classmethod
    def _normalize_result(cls, value: object) -> str:
        text = str(value or "").strip().upper().replace("ASIL ", "")
        if text in {"NA", "N/A", "-"}:
            return "QM"
        if text not in cls.VALID_RESULTS:
            raise ValueError(f"ASIL_Table包含不支持的结果: {value!r}")
        return text

    def _load(self) -> dict[tuple[str, str, str], str]:
        if not self.template_path.is_file():
            raise FileNotFoundError(f"ASIL模板不存在: {self.template_path}")
        workbook = load_workbook(self.template_path, read_only=True, data_only=False)
        try:
            if self.SHEET_NAME not in workbook.sheetnames:
                raise ValueError(f"模板缺少必需Sheet: {self.SHEET_NAME}")
            sheet = workbook[self.SHEET_NAME]
            header_row = None
            c_columns: dict[str, int] = {}
            for row in range(1, min(sheet.max_row, 30) + 1):
                found = {}
                for col in range(1, sheet.max_column + 1):
                    code = self._code(sheet.cell(row, col).value, "C", 3)
                    if code:
                        found[code] = col
                if set(found) == set(self.EXPECTED_C):
                    header_row, c_columns = row, found
                    break
            if header_row is None:
                raise ValueError("ASIL_Table缺少完整C0-C3表头")

            matrix = {}
            current_s = None
            for row in range(header_row + 1, sheet.max_row + 1):
                row_s = row_e = None
                for col in range(1, min(c_columns.values())):
                    value = sheet.cell(row, col).value
                    row_s = row_s or self._code(value, "S", 3)
                    row_e = row_e or self._code(value, "E", 4)
                current_s = row_s or current_s
                if not current_s or not row_e:
                    continue
                for c_code, col in c_columns.items():
                    matrix[(current_s, row_e, c_code)] = self._normalize_result(
                        sheet.cell(row, col).value
                    )

            expected = {
                (s_code, e_code, c_code)
                for s_code in self.EXPECTED_S
                for e_code in self.EXPECTED_E
                for c_code in self.EXPECTED_C
            }
            missing = sorted(expected - set(matrix))
            extra = sorted(set(matrix) - expected)
            if missing or extra:
                raise ValueError(
                    "ASIL_Table矩阵不完整或存在歧义: "
                    f"expected={len(expected)}, loaded={len(matrix)}, "
                    f"missing={missing[:5]}, extra={extra[:5]}"
                )
            return matrix
        finally:
            workbook.close()

    def determine(self, severity: str, exposure: str, controllability: str) -> str:
        key = tuple(str(value or "").strip().upper() for value in (
            severity, exposure, controllability
        ))
        if key not in self.matrix:
            raise ValueError(
                "Invalid S/E/C combination / 无效S/E/C组合: "
                f"severity={severity!r}, exposure={exposure!r}, controllability={controllability!r}"
            )
        return self.matrix[key]
