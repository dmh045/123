from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .template_scoring import TemplateScoringStandardReader, TemplateScoringStandards


@dataclass(frozen=True)
class TemplateInputs:
    guidewords: list[str]
    guideword_descriptions: dict[str, str]
    scenario_dimensions: dict[str, list[str]]
    scoring_standards: TemplateScoringStandards
    source_path: Path


class TemplateInputReader:
    """Read versioned method inputs from the current run template.

    Guideword values are template-driven. Their current cardinality is an
    explicit HARA method-contract invariant, not an AVP/domain business rule.
    """

    HAZOP_SHEET = "04_HAZOP"
    SCENARIO_SHEET = "Scenarios_Library"
    EXPECTED_GUIDEWORD_COUNT = 14

    def read(self, template_path: str | Path) -> TemplateInputs:
        path = Path(template_path).expanduser().resolve()
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for name in (self.HAZOP_SHEET, self.SCENARIO_SHEET):
                if name not in workbook.sheetnames:
                    raise ValueError(f"模板缺少分析输入Sheet: {name}")
            guidewords, descriptions = self._guidewords(workbook[self.HAZOP_SHEET])
            dimensions = self._scenario_dimensions(workbook[self.SCENARIO_SHEET])
            scoring_standards = TemplateScoringStandardReader().read(workbook, path)
        finally:
            workbook.close()
        return TemplateInputs(guidewords, descriptions, dimensions, scoring_standards, path)

    @staticmethod
    def _guidewords(sheet) -> tuple[list[str], dict[str, str]]:
        guidewords = []
        descriptions = {}
        for row in range(3, sheet.max_row + 1):
            guideword = str(sheet.cell(row, 1).value or "").strip()
            if not guideword:
                continue
            if guideword in descriptions:
                raise ValueError(f"04_HAZOP包含重复Guideword: {guideword}")
            guidewords.append(guideword)
            descriptions[guideword] = str(sheet.cell(row, 2).value or "").strip()
        if len(guidewords) != TemplateInputReader.EXPECTED_GUIDEWORD_COUNT:
            raise ValueError(
                "04_HAZOP Guideword数量不符合当前方法契约: "
                f"expected={TemplateInputReader.EXPECTED_GUIDEWORD_COUNT}, "
                f"actual={len(guidewords)}"
            )
        return guidewords, descriptions

    @staticmethod
    def _scenario_dimensions(sheet) -> dict[str, list[str]]:
        headers = {}
        for column in range(1, sheet.max_column + 1):
            header = str(sheet.cell(2, column).value or "").strip()
            if header:
                headers[column] = header
        dimensions = {header: [] for header in headers.values()}
        for column, header in headers.items():
            for row in range(3, sheet.max_row + 1):
                value = str(sheet.cell(row, column).value or "").strip()
                if not value or value.lower().startswith("all "):
                    continue
                if value not in dimensions[header]:
                    dimensions[header].append(value)
        if not dimensions:
            raise ValueError("Scenarios_Library未提供任何场景维度")
        return dimensions
