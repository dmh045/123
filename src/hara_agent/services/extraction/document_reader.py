from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path


@dataclass(frozen=True)
class DocumentBlock:
    block_id: str
    kind: str
    location: str
    text: str


@dataclass
class DocumentArtifact:
    source_path: Path
    source_id: str
    blocks: list[DocumentBlock] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(block.text for block in self.blocks if block.text)


class DocumentReader:
    """Read source documents into traceable blocks without semantic guessing."""

    def read(self, path: str | Path) -> DocumentArtifact:
        source = Path(path).expanduser().resolve()
        if not source.is_file():
            raise FileNotFoundError(f"Item Definition不存在: {source}")
        suffix = source.suffix.lower()
        if suffix == ".docx":
            blocks = self._read_docx(source)
        elif suffix == ".pdf":
            blocks = self._read_pdf(source)
        else:
            raise ValueError(f"不支持的Item Definition格式: {suffix}")
        if not blocks:
            raise ValueError(f"文档未提取到有效文本: {source}")
        return DocumentArtifact(source, source.name, blocks)

    @staticmethod
    def _read_docx(path: Path) -> list[DocumentBlock]:
        from docx import Document

        document = Document(str(path))
        blocks = []
        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if text:
                blocks.append(DocumentBlock(
                    f"P-{index:04d}", "paragraph", f"paragraph[{index}]", text,
                ))
        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                values = [cell.text.strip() for cell in row.cells]
                text = " | ".join(value for value in values if value)
                if text:
                    blocks.append(DocumentBlock(
                        f"T-{table_index:03d}-R-{row_index:04d}", "table_row",
                        f"table[{table_index}].row[{row_index}]", text,
                    ))
        return blocks

    @staticmethod
    def _read_pdf(path: Path) -> list[DocumentBlock]:
        from PyPDF2 import PdfReader

        reader = PdfReader(str(path))
        blocks = []
        for index, page in enumerate(reader.pages, start=1):
            text = str(page.extract_text() or "").strip()
            if text:
                blocks.append(DocumentBlock(
                    f"PAGE-{index:04d}", "page", f"page[{index}]", text,
                ))
        return blocks

