from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


class HARATemplateContract:
    HARA_SHEET = "05_HARA"
    SG_SHEET = "06_Safety Goal"
    ASIL_SHEET = "ASIL_Table"
    HARA_START_ROW = 6
    SG_START_ROW = 4
    HARA_COLUMNS = 21

    def validate(self, template_path: str | Path) -> Path:
        path = Path(template_path).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(f"HARA模板不存在: {path}")
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            required = {self.HARA_SHEET, self.SG_SHEET, self.ASIL_SHEET}
            missing = sorted(required - set(workbook.sheetnames))
            if missing:
                raise ValueError(f"HARA模板缺少必需Sheet: {missing}")
            hara = workbook[self.HARA_SHEET]
            expected = {
                1: "hara-id", 2: "function", 3: "output", 4: "guide-word",
                5: "malfunction", 6: "hazard", 17: "asil",
                18: "sg-id", 19: "safety goal", 20: "safe state", 21: "remark",
            }
            for column, token in expected.items():
                text = " ".join(str(hara.cell(row, column).value or "") for row in (4, 5)).lower()
                if token not in text:
                    raise ValueError(f"05_HARA模板列契约不匹配: column={column}, expected={token!r}")
            sg = workbook[self.SG_SHEET]
            sg_headers = [str(sg.cell(3, col).value or "").replace("\n", " ").lower() for col in range(1, 7)]
            for token in ("hz-id", "hazard", "sg-id", "safety goal", "safety state", "asil"):
                if not any(token in header for header in sg_headers):
                    raise ValueError(f"06_Safety Goal模板列契约缺失: {token!r}")
        finally:
            workbook.close()
        return path
