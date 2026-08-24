from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Iterable

from openpyxl.utils import get_column_letter, range_boundaries

from hara_agent.contracts import (
    ASILMapping, ASILMatrix, CategoricalPredicate, CompileStatus, CompiledRule,
    CompilerDiagnostic, CompilerDiagnosticCode, CompilerDiagnosticSeverity,
    ControllabilityContract, DerivationMethod, ExposureContract, ExposureEntry,
    FactOrigin, FactType, Guideword, GuidewordContract, MethodAssumption,
    MethodContract, NormativeStrength, ParseStatus, Predicate, PredicateOperator,
    RangePredicate, ReportContract, ReportFieldMapping, RequiredFactSpec,
    RoleBinding, RuleType, ScaleLevel, ScenarioDimension, ScenarioModel,
    SeverityContract, SeverityScale, SourceRef, TemplateDiagnosticCode,
    TemplateRole, TemplateRoleContract,
    WorkflowContract, WorkflowStep, unique_sources,
)

from .scanner import CellSnapshot, SheetSnapshot, WorkbookSnapshot, normalize_template_text


_RATING_RE = re.compile(r"^[SEC][0-9]$", re.IGNORECASE)
_EXPOSURE_RE = re.compile(r"^E[0-4]$", re.IGNORECASE)


@dataclass(frozen=True)
class _BoundRegion:
    snapshot: WorkbookSnapshot
    binding: RoleBinding

    @property
    def sheet(self) -> SheetSnapshot:
        return self.snapshot.sheet(self.binding.sheet)

    @property
    def bounds(self) -> tuple[int, int, int, int]:
        return range_boundaries(self.binding.region)

    @property
    def cells(self) -> tuple[CellSnapshot, ...]:
        min_column, min_row, max_column, max_row = self.bounds
        return tuple(
            cell for cell in self.sheet.cells
            if min_row <= cell.row <= max_row and min_column <= cell.column <= max_column
        )

    def row(self, row: int) -> tuple[CellSnapshot, ...]:
        return tuple(sorted((c for c in self.cells if c.row == row), key=lambda c: c.column))

    def cell(self, row: int, column: int) -> CellSnapshot | None:
        return next((c for c in self.cells if c.row == row and c.column == column), None)

    def source(self, cells: Iterable[CellSnapshot]) -> SourceRef:
        selected = tuple(sorted(cells, key=lambda c: (c.row, c.column)))
        if not selected:
            raise ValueError(f"Cannot create SourceRef for empty {self.binding.role.value} source")
        min_row = min(c.row for c in selected)
        max_row = max(c.row for c in selected)
        min_column = min(c.column for c in selected)
        max_column = max(c.column for c in selected)
        region = (
            f"{get_column_letter(min_column)}{min_row}:"
            f"{get_column_letter(max_column)}{max_row}"
            if len(selected) > 1
            else selected[0].coordinate
        )
        raw = " | ".join(f"{c.coordinate}={c.raw_value}" for c in selected)
        return SourceRef.create(
            workbook=self.snapshot.source_path.name,
            template_hash=self.snapshot.source_hash,
            sheet=self.binding.sheet,
            range=region,
            raw_text=raw,
        )


class FullTemplateCompiler:
    """Compile already-discovered workbook roles into an executable method IR.

    Parsers are selected by canonical role and consume only the role-bound
    region. Sheet names and workbook hashes never select engineering behavior.
    """

    def compile(
        self, snapshot: WorkbookSnapshot, role_contract: TemplateRoleContract
    ) -> MethodContract:
        if snapshot.source_hash != role_contract.template_hash:
            raise ValueError("TemplateRoleContract hash does not match workbook snapshot")
        if not role_contract.ready:
            raise ValueError("TemplateRoleContract is not ready for full method compilation")

        diagnostics = self._role_diagnostics(role_contract)
        report_contract, report_diags = self._compile_report(snapshot, role_contract)
        diagnostics.extend(report_diags)
        workflow = self._compile_workflow(snapshot, role_contract, report_contract)
        guidewords = self._compile_guidewords(snapshot, role_contract)
        scenario_model = self._compile_scenarios(snapshot, role_contract)
        severity, severity_diags = self._compile_severity(snapshot, role_contract)
        diagnostics.extend(severity_diags)
        exposure, exposure_diags = self._compile_exposure(snapshot, role_contract)
        diagnostics.extend(exposure_diags)
        controllability, controllability_diags = self._compile_controllability(
            snapshot, role_contract
        )
        diagnostics.extend(controllability_diags)
        asil, asil_diags = self._compile_asil(snapshot, role_contract)
        diagnostics.extend(asil_diags)
        safety_goal = self._compile_derivation_method(
            snapshot, role_contract, report_contract, TemplateRole.SAFETY_GOAL_METHOD
        )
        safe_state = self._compile_derivation_method(
            snapshot, role_contract, report_contract, TemplateRole.SAFE_STATE_METHOD
        )

        all_rules = (
            severity.rules + exposure.duration_rules + exposure.frequency_rules
            + exposure.examples + controllability.criteria + controllability.examples
            + controllability.references + safety_goal.instructions + safe_state.instructions
        )
        for rule in all_rules:
            if (
                rule.executable
                and rule.normative_strength in {NormativeStrength.NORMATIVE, NormativeStrength.CRITERION}
                and not rule.source_refs
            ):
                diagnostics.append(
                    CompilerDiagnostic(
                        CompilerDiagnosticSeverity.ERROR,
                        CompilerDiagnosticCode.SOURCE_REF_MISSING,
                        f"Executable rule {rule.rule_id} has no workbook SourceRef.",
                        role=rule.role,
                        blocking=True,
                    )
                )

        required_facts = self._derive_required_facts(all_rules, workflow, exposure)
        core_complete = all(
            (
                guidewords.guidewords,
                scenario_model.dimensions,
                severity.scale.levels,
                severity.rules,
                exposure.duration_rules,
                exposure.frequency_rules,
                controllability.scale,
                controllability.criteria,
                asil.mappings,
                report_contract.hara_fields,
                report_contract.safety_goal_fields,
            )
        )
        if not core_complete:
            diagnostics.append(
                CompilerDiagnostic(
                    CompilerDiagnosticSeverity.ERROR,
                    CompilerDiagnosticCode.RULE_PARSE_FAILED,
                    "One or more required MethodContract sections compiled empty.",
                    blocking=True,
                )
            )
        blocking = any(item.blocking for item in diagnostics)
        status = (
            CompileStatus.NOT_READY if blocking
            else CompileStatus.READY_WITH_WARNINGS
            if any(item.severity is CompilerDiagnosticSeverity.WARNING for item in diagnostics)
            else CompileStatus.READY
        )
        source_items: list[SourceRef] = []
        source_items.extend(step.source_ref for step in workflow.steps)
        source_items.extend(item.source_ref for item in guidewords.guidewords)
        source_items.extend(item.source_ref for item in scenario_model.dimensions)
        source_items.extend(item.source_ref for item in severity.scale.levels)
        source_items.extend(source for rule in all_rules for source in rule.source_refs)
        source_items.extend(item.source_ref for item in exposure.situation_mappings)
        source_items.extend(item.source_ref for item in asil.mappings)
        source_items.extend(item.source_ref for item in report_contract.hara_fields)
        source_items.extend(item.source_ref for item in report_contract.safety_goal_fields)

        return MethodContract(
            metadata={
                "template_id": role_contract.template_id,
                "template_hash": role_contract.template_hash,
                "role_contract_version": role_contract.contract_version,
                "compiler_architecture": "RoleContract -> Rule Parser -> Rule IR -> MethodContract",
                "engineering_rules_compiled": core_complete and not blocking,
            },
            workflow=workflow,
            guidewords=guidewords,
            scenario_model=scenario_model,
            severity=severity,
            exposure=exposure,
            controllability=controllability,
            asil=asil,
            safety_goal_method=safety_goal,
            safe_state_method=safe_state,
            report_contract=report_contract,
            required_fact_specs=required_facts,
            diagnostics=tuple(diagnostics),
            sources=unique_sources(source_items),
            compile_status=status,
            engineering_rules_compiled=core_complete and not blocking,
        )

    @staticmethod
    def _region(
        snapshot: WorkbookSnapshot, role_contract: TemplateRoleContract, role: TemplateRole
    ) -> _BoundRegion:
        return _BoundRegion(snapshot, role_contract.binding(role))

    def _role_diagnostics(
        self, role_contract: TemplateRoleContract
    ) -> list[CompilerDiagnostic]:
        result: list[CompilerDiagnostic] = []
        for item in role_contract.diagnostics:
            code = (
                CompilerDiagnosticCode.WORKFLOW_COORDINATE_MISMATCH
                if item.code is TemplateDiagnosticCode.WORKFLOW_COORDINATE_MISMATCH
                else CompilerDiagnosticCode.MALFORMED_TEMPLATE_TEXT
                if item.code is TemplateDiagnosticCode.STRUCTURE_INCOMPLETE
                else None
            )
            if code is None:
                continue
            result.append(
                CompilerDiagnostic(
                    CompilerDiagnosticSeverity.WARNING,
                    code,
                    item.message,
                    role=item.role,
                    details=item.details,
                )
            )
        return result

    def _compile_workflow(
        self,
        snapshot: WorkbookSnapshot,
        role_contract: TemplateRoleContract,
        report: ReportContract,
    ) -> WorkflowContract:
        region = self._region(snapshot, role_contract, TemplateRole.WORKFLOW)
        header_row = min(c.row for c in region.cells if "activity" in c.normalized_value or "活动" in c.normalized_value)
        header = {c.normalized_value: c.column for c in region.row(header_row)}
        activity_col = next(v for k, v in header.items() if "activity" in k or "活动" in k)
        input_col = next((v for k, v in header.items() if k in {"input", "输入"}), activity_col - 1)
        description_col = next((v for k, v in header.items() if "description" in k or "描述" in k), activity_col + 1)
        output_col = next((v for k, v in header.items() if k in {"output", "输出"}), max(header.values()))
        activity_map = (
            (("引导词", "guideword"), "GUIDEWORD_APPLICATION"),
            (("hazardous event", "危害事件"), "HAZARDOUS_EVENT_DERIVATION"),
            (("hazard", "危害"), "HAZARD_DERIVATION"),
            (("malfunction",), "MALFUNCTION_DERIVATION"),
            (("细化", "detailing"), "SCENARIO_DETAILING"),
            (("组合场景", "scenario"), "SCENARIO_CONSTRUCTION"),
            (("严重度", "severity"), "SEVERITY_ASSESSMENT"),
            (("暴露度", "exposure"), "EXPOSURE_ASSESSMENT"),
            (("可控度", "controllability"), "CONTROLLABILITY_ASSESSMENT"),
            (("asil",), "ASIL_LOOKUP"),
            (("安全目标", "safety goal"), "SAFETY_GOAL_DERIVATION"),
            (("安全状态", "safe state"), "SAFE_STATE_DERIVATION"),
            (("输出", "功能", "function"), "FUNCTION_EXTRACTION"),
        )
        steps: list[WorkflowStep] = []
        for row in range(header_row + 1, region.bounds[3] + 1):
            activity = region.cell(row, activity_col)
            if activity is None:
                continue
            normalized = activity.normalized_value
            canonical = next(
                (value for aliases, value in activity_map if any(alias in normalized for alias in aliases)),
                "WORKFLOW_INSTRUCTION",
            )
            input_cell = region.cell(row, input_col)
            description_cell = region.cell(row, description_col)
            output_cell = region.cell(row, output_col)
            row_cells = region.row(row)
            steps.append(
                WorkflowStep(
                    activity=canonical,
                    inputs=(str(input_cell.raw_value),) if input_cell else (),
                    outputs=(str(output_cell.raw_value),) if output_cell else (),
                    description=str(description_cell.raw_value) if description_cell else "",
                    source_ref=region.source(row_cells),
                )
            )
        report_sources = tuple(item.source_ref for item in report.hara_fields)
        if report_sources:
            steps.append(
                WorkflowStep(
                    activity="REPORTING",
                    inputs=("canonical HARA records",),
                    outputs=("template-bound report fields",),
                    description="Render canonical fields through ReportContract mappings.",
                    source_ref=report_sources[0],
                )
            )
        return WorkflowContract(tuple(steps), region.binding)

    def _compile_guidewords(
        self, snapshot: WorkbookSnapshot, role_contract: TemplateRoleContract
    ) -> GuidewordContract:
        region = self._region(snapshot, role_contract, TemplateRole.GUIDEWORD_TABLE)
        rows = sorted({cell.row for cell in region.cells})
        header_row = rows[0]
        guidewords: list[Guideword] = []
        for order, row in enumerate((r for r in rows if r > header_row), start=1):
            cells = region.row(row)
            if len(cells) < 2:
                continue
            guidewords.append(
                Guideword(
                    guideword_id=f"GW-{order:03d}",
                    name=str(cells[0].raw_value).strip(),
                    description=str(cells[1].raw_value).strip(),
                    order=order,
                    source_ref=region.source(cells[:2]),
                )
            )
        return GuidewordContract(tuple(guidewords), region.binding)

    def _compile_scenarios(
        self, snapshot: WorkbookSnapshot, role_contract: TemplateRoleContract
    ) -> ScenarioModel:
        region = self._region(snapshot, role_contract, TemplateRole.SCENARIO_MODEL)
        min_column, min_row, max_column, max_row = region.bounds
        header_row = min_row
        aliases = {
            "operating scenario": "OPERATING_SCENARIO",
            "vehicle state": "VEHICLE_STATE",
            "vehicle speed": "VEHICLE_SPEED",
            "weather": "WEATHER",
            "road surface": "ROAD_SURFACE",
            "traffic density": "TRAFFIC_DENSITY",
        }
        dimensions: list[ScenarioDimension] = []
        constraints: list[str] = []
        for column in range(min_column, max_column + 1):
            header = region.cell(header_row, column)
            if header is None:
                continue
            normalized = header.normalized_value
            canonical = next((value for key, value in aliases.items() if key in normalized), None)
            if canonical is None:
                canonical = re.sub(r"[^A-Z0-9]+", "_", str(header.raw_value).upper()).strip("_")
            values: list[str] = []
            source_cells = [header]
            for row in range(header_row + 1, max_row + 1):
                cell = region.cell(row, column)
                if cell is None:
                    continue
                source_cells.append(cell)
                text = str(cell.raw_value).strip()
                if normalize_template_text(text).startswith("all "):
                    constraints.append(text)
                else:
                    values.append(text)
            dimensions.append(
                ScenarioDimension(
                    dimension_id=f"DIM-{len(dimensions) + 1:03d}",
                    canonical_name=canonical,
                    display_name=str(header.raw_value).strip(),
                    values=tuple(values),
                    unit="km/h" if "speed" in normalized else "",
                    source_ref=region.source(source_cells),
                )
            )
        return ScenarioModel(tuple(dimensions), tuple(constraints), region.binding)

    def _compile_severity(
        self, snapshot: WorkbookSnapshot, role_contract: TemplateRoleContract
    ) -> tuple[SeverityContract, list[CompilerDiagnostic]]:
        diagnostics: list[CompilerDiagnostic] = []
        level_region = self._region(snapshot, role_contract, TemplateRole.SEVERITY_LEVELS)
        levels = self._compile_scale(level_region, "S")
        scale = SeverityScale(tuple(levels), level_region.binding)
        rule_region = self._region(snapshot, role_contract, TemplateRole.SEVERITY_RULES)
        rules: list[CompiledRule] = []
        unresolved_speed_sources: list[SourceRef] = []
        ambiguous_sources: list[SourceRef] = []
        for cell in rule_region.cells:
            raw = str(cell.raw_value).strip()
            if not re.search(r"\bS[0-3]\b", raw, re.IGNORECASE) or ":" not in raw:
                continue
            source = rule_region.source((cell,))
            condition, _separator, result_text = raw.rpartition(":")
            result_values = tuple(dict.fromkeys(re.findall(r"\bS[0-3]\b", result_text.upper())))
            if not result_values:
                continue
            subject = condition.split("(", 1)[0].strip()
            predicates: list[Predicate | RangePredicate | CategoricalPredicate] = []
            subject_lower = subject.casefold()
            if "rear" in subject_lower:
                predicates.append(
                    CategoricalPredicate(FactType.COLLISION_TYPE, ("REAR_END",), source=source)
                )
            elif "frontal" in subject_lower:
                predicates.append(
                    CategoricalPredicate(FactType.COLLISION_TYPE, ("FRONTAL",), source=source)
                )
            elif "side" in subject_lower:
                predicates.append(
                    CategoricalPredicate(FactType.COLLISION_TYPE, ("SIDE",), source=source)
                )
            elif "pedestrian" in subject_lower:
                predicates.extend(
                    (
                        CategoricalPredicate(
                            FactType.COLLISION_TYPE, ("VEHICLE_TO_ROAD_USER",), source=source
                        ),
                        CategoricalPredicate(
                            FactType.ROAD_USER_TYPE, ("PEDESTRIAN",), source=source
                        ),
                    )
                )
            elif "cyclist" in subject_lower:
                predicates.extend(
                    (
                        CategoricalPredicate(
                            FactType.COLLISION_TYPE, ("VEHICLE_TO_ROAD_USER",), source=source
                        ),
                        CategoricalPredicate(
                            FactType.ROAD_USER_TYPE, ("CYCLIST",), source=source
                        ),
                    )
                )
            speed_text_match = re.search(r"\(([^)]*\bv\b[^)]*)\)", condition, re.IGNORECASE)
            if speed_text_match:
                speed_text = speed_text_match.group(1)
                speed_predicate = self._parse_numeric_range(
                    speed_text, FactType.SPEED_UNSPECIFIED, "km/h", source
                )
                if speed_predicate is not None:
                    predicates.append(speed_predicate)
                    unresolved_speed_sources.append(source)
            alternatives = result_values if len(result_values) > 1 else ()
            if alternatives:
                ambiguous_sources.append(source)
            rules.append(
                CompiledRule(
                    rule_id=f"TPL-SEV-{len(rules) + 1:04d}",
                    rule_type=RuleType.NORMATIVE_RULE,
                    role=TemplateRole.SEVERITY_RULES,
                    predicates=tuple(predicates),
                    result=result_values[0] if len(result_values) == 1 else None,
                    alternatives=alternatives,
                    priority=len(rules) + 1,
                    normative_strength=NormativeStrength.NORMATIVE,
                    source_refs=(source,),
                    raw_text=raw,
                    parser="severity-collision-rule-v1",
                    parse_status=(
                        ParseStatus.NEEDS_REVIEW
                        if alternatives or speed_text_match or any(
                            isinstance(item, RangePredicate) and self._range_needs_review(item)
                            for item in predicates
                        )
                        else ParseStatus.COMPILED
                    ),
                    diagnostics=tuple(
                        item for item in (
                            "speed variable semantics unresolved" if speed_text_match else "",
                            "result alternatives preserved" if alternatives else "",
                        ) if item
                    ),
                    resolution_requirement=(
                        "Template must specify the condition selecting among severity alternatives."
                        if alternatives else
                        "Template must define whether v is ego, relative, impact speed, or delta-V."
                    ),
                )
            )
        if unresolved_speed_sources:
            diagnostics.append(
                CompilerDiagnostic(
                    CompilerDiagnosticSeverity.WARNING,
                    CompilerDiagnosticCode.UNRESOLVED_RULE_VARIABLE,
                    "Severity rules use 'v' without defining ego, relative, impact speed, or delta-V; compiled as SPEED_UNSPECIFIED.",
                    role=TemplateRole.SEVERITY_RULES,
                    source_refs=unique_sources(unresolved_speed_sources),
                    details={"field": FactType.SPEED_UNSPECIFIED.value},
                )
            )
        if ambiguous_sources:
            diagnostics.append(
                CompilerDiagnostic(
                    CompilerDiagnosticSeverity.WARNING,
                    CompilerDiagnosticCode.AMBIGUOUS_RESULT,
                    "Severity alternatives were preserved and must fail closed until a selection condition is available.",
                    role=TemplateRole.SEVERITY_RULES,
                    source_refs=unique_sources(ambiguous_sources),
                )
            )
        diagnostics.extend(self._detect_rule_conflicts(rules, TemplateRole.SEVERITY_RULES))
        diagnostics.extend(self._detect_range_coverage(rules, FactType.SPEED_UNSPECIFIED))
        return SeverityContract(scale, tuple(rules), tuple(diagnostics)), diagnostics

    def _compile_scale(self, region: _BoundRegion, prefix: str) -> list[ScaleLevel]:
        expected = re.compile(rf"^{re.escape(prefix)}\d+$", re.IGNORECASE)
        header_cells = [c for c in region.cells if expected.fullmatch(c.normalized_value.upper())]
        if not header_cells:
            return []
        header_row = min(c.row for c in header_cells)
        levels: list[ScaleLevel] = []
        for header in sorted((c for c in header_cells if c.row == header_row), key=lambda c: c.column):
            description = ""
            criterion_parts: list[str] = []
            evidence_parts: list[str] = []
            source_cells = [header]
            for row in range(header_row + 1, region.bounds[3] + 1):
                value = region.cell(row, header.column)
                label = region.cell(row, region.bounds[0])
                if value is None:
                    continue
                source_cells.append(value)
                label_text = label.normalized_value if label else ""
                if "description" in label_text:
                    description = str(value.raw_value).strip()
                elif "example" in label_text:
                    evidence_parts.append(str(value.raw_value).strip())
                else:
                    criterion_parts.append(str(value.raw_value).strip())
            levels.append(
                ScaleLevel(
                    level=header.normalized_value.upper(),
                    description=description,
                    criterion="\n".join(criterion_parts),
                    evidence="\n".join(evidence_parts),
                    source_ref=region.source(source_cells),
                )
            )
        return levels

    def _parse_numeric_range(
        self,
        raw: str,
        field: FactType,
        unit: str,
        source: SourceRef,
    ) -> RangePredicate | None:
        text = normalize_template_text(raw).replace("≤", "<=").replace("≥", ">=")
        text = re.sub(r"(?<=\d)\s*-\s*(?=\d)", " to ", text)
        numbers = [float(item.replace(",", ".")) for item in re.findall(r"-?\d+(?:[.,]\d+)?", text)]
        if not numbers:
            return None
        lowered = text.casefold()
        if "more than" in lowered and len(numbers) == 1:
            return RangePredicate(field, numbers[0], None, False, None, unit, source)
        if "less than" in lowered and len(numbers) == 1:
            return RangePredicate(field, None, numbers[0], None, False, unit, source)
        if "<=" in text and len(numbers) == 1:
            return RangePredicate(field, None, numbers[0], None, True, unit, source)
        if ">=" in text and len(numbers) == 1:
            return RangePredicate(field, numbers[0], None, True, None, unit, source)
        if "<" in text and len(numbers) == 1:
            return RangePredicate(field, None, numbers[0], None, False, unit, source)
        if ">" in text and len(numbers) == 1:
            return RangePredicate(field, numbers[0], None, False, None, unit, source)
        if len(numbers) >= 2:
            lower_inclusive: bool | None = None
            upper_inclusive: bool | None = None
            compact = text.replace(" ", "")
            if re.search(rf"{numbers[0]:g}<=", compact):
                lower_inclusive = True
            elif re.search(rf"{numbers[0]:g}<", compact):
                lower_inclusive = False
            if re.search(rf"<={numbers[1]:g}", compact):
                upper_inclusive = True
            elif re.search(rf"<{numbers[1]:g}", compact):
                upper_inclusive = False
            return RangePredicate(
                field, min(numbers[0], numbers[1]), max(numbers[0], numbers[1]),
                lower_inclusive, upper_inclusive, unit, source
            )
        return None

    @staticmethod
    def _range_needs_review(predicate: RangePredicate) -> bool:
        return (
            predicate.lower is not None and predicate.lower_inclusive is None
        ) or (
            predicate.upper is not None and predicate.upper_inclusive is None
        )

    def _compile_exposure(
        self, snapshot: WorkbookSnapshot, role_contract: TemplateRoleContract
    ) -> tuple[ExposureContract, list[CompilerDiagnostic]]:
        diagnostics: list[CompilerDiagnostic] = []
        duration_region = self._region(
            snapshot, role_contract, TemplateRole.EXPOSURE_DURATION_RULES
        )
        frequency_region = self._region(
            snapshot, role_contract, TemplateRole.EXPOSURE_FREQUENCY_RULES
        )
        duration_rules = self._compile_duration_rules(duration_region)
        frequency_rules = self._compile_frequency_rules(frequency_region)
        example_region = self._region(snapshot, role_contract, TemplateRole.EXPOSURE_EXAMPLES)
        examples = self._compile_exposure_examples(example_region)
        mappings: list[ExposureEntry] = []
        for role in (
            TemplateRole.VDA702_SUMMARY,
            TemplateRole.VDA702_FULL,
            TemplateRole.SITUATION_CATALOG,
        ):
            mappings.extend(self._compile_exposure_catalog(self._region(snapshot, role_contract, role)))
        diagnostics.append(
            CompilerDiagnostic(
                CompilerDiagnosticSeverity.WARNING,
                CompilerDiagnosticCode.EXPOSURE_METHOD_SELECTION_UNRESOLVED,
                "Template provides duration (T), frequency (F), and situation classifications but does not define a universal method-selection precedence.",
                role=TemplateRole.EXPOSURE_DURATION_RULES,
                source_refs=tuple(
                    source for rule in duration_rules[:1] + frequency_rules[:1]
                    for source in rule.source_refs
                ),
            )
        )
        diagnostics.extend(self._detect_range_coverage(duration_rules, FactType.DURATION_PERCENT))
        ambiguous_duration = [
            rule for rule in duration_rules
            if any(
                isinstance(item, RangePredicate) and self._range_needs_review(item)
                for item in rule.predicates
            )
        ]
        if ambiguous_duration:
            diagnostics.append(
                CompilerDiagnostic(
                    CompilerDiagnosticSeverity.WARNING,
                    CompilerDiagnosticCode.SEMANTIC_VARIABLE_UNRESOLVED,
                    "Exposure duration ranges using natural-language 'to' preserve unresolved boundary inclusivity.",
                    role=TemplateRole.EXPOSURE_DURATION_RULES,
                    source_refs=tuple(
                        source for rule in ambiguous_duration for source in rule.source_refs
                    ),
                )
            )
        scale = self._exposure_scale(duration_region, frequency_region)
        workflow_region = self._region(snapshot, role_contract, TemplateRole.WORKFLOW)
        e0_cell = next(
            (
                cell for cell in workflow_region.cells
                if "e0" in cell.normalized_value and (
                    "表示" in cell.normalized_value or "means" in cell.normalized_value
                )
            ),
            None,
        )
        if e0_cell is not None and not any(item.level == "E0" for item in scale):
            scale.insert(
                0,
                ScaleLevel(
                    "E0", str(e0_cell.raw_value), str(e0_cell.raw_value), "",
                    workflow_region.source((e0_cell,)),
                ),
            )
        contract = ExposureContract(
            scale=tuple(scale),
            duration_rules=tuple(duration_rules),
            frequency_rules=tuple(frequency_rules),
            situation_mappings=tuple(mappings),
            examples=tuple(examples),
            diagnostics=tuple(diagnostics),
        )
        return contract, diagnostics

    def _rating_columns(self, region: _BoundRegion, row: int, prefix: str) -> dict[int, str]:
        result: dict[int, str] = {}
        for candidate_row in range(row - 1, max(region.bounds[1] - 1, row - 7), -1):
            found = {
                cell.column: cell.normalized_value.upper()
                for cell in region.row(candidate_row)
                if re.fullmatch(rf"{prefix}\d", cell.normalized_value.upper())
            }
            if len(found) >= 3:
                return found
        return result

    def _compile_duration_rules(self, region: _BoundRegion) -> list[CompiledRule]:
        rules: list[CompiledRule] = []
        for row in sorted({c.row for c in region.cells}):
            row_cells = region.row(row)
            if not any("duration (%" in c.normalized_value for c in row_cells):
                continue
            rating_columns = self._rating_columns(region, row, "E")
            section_text = " ".join(
                c.normalized_value
                for previous in range(max(region.bounds[1], row - 6), row)
                for c in region.row(previous)
            )
            method = "T:T&B" if "t&b" in section_text else "T"
            for column, rating in sorted(rating_columns.items()):
                cell = region.cell(row, column)
                if cell is None or cell.normalized_value in {"", "-"}:
                    continue
                source = region.source((cell,))
                predicate = self._parse_numeric_range(
                    str(cell.raw_value), FactType.DURATION_PERCENT, "%", source
                )
                rules.append(
                    CompiledRule(
                        rule_id=f"TPL-EXP-DUR-{len(rules) + 1:04d}",
                        rule_type=RuleType.CRITERION,
                        role=TemplateRole.EXPOSURE_DURATION_RULES,
                        predicates=(predicate,) if predicate else (),
                        result=rating,
                        alternatives=(),
                        priority=len(rules) + 1,
                        normative_strength=NormativeStrength.CRITERION,
                        source_refs=(source,),
                        raw_text=str(cell.raw_value),
                        parser="exposure-duration-range-v1",
                        parse_status=(
                            ParseStatus.NEEDS_REVIEW
                            if predicate is None or self._range_needs_review(predicate)
                            else ParseStatus.COMPILED
                        ),
                        diagnostics=(
                            ("non-numeric duration criterion preserved",)
                            if predicate is None else
                            ("range boundary inclusivity unresolved",)
                            if self._range_needs_review(predicate) else ()
                        ),
                        resolution_requirement=(
                            "Define the meaning of 'Not specified' for deterministic evaluation."
                            if predicate is None else ""
                        ),
                        assessment_method=method,
                        executable=predicate is not None,
                    )
                )
        return rules

    def _compile_frequency_rules(self, region: _BoundRegion) -> list[CompiledRule]:
        categories = (
            ("less often than once a year", "LESS_THAN_ANNUAL"),
            ("few times a year", "FEW_PER_YEAR"),
            ("once a month or more often", "MONTHLY_OR_MORE"),
            ("almost every drive", "ALMOST_EVERY_DRIVE"),
        )
        rules: list[CompiledRule] = []
        for row in sorted({c.row for c in region.cells}):
            row_cells = region.row(row)
            if not any("frequency of situation" in c.normalized_value for c in row_cells):
                continue
            rating_columns = self._rating_columns(region, row, "E")
            section_text = " ".join(
                c.normalized_value
                for previous in range(max(region.bounds[1], row - 6), row)
                for c in region.row(previous)
            )
            method = "F:T&B" if "t&b" in section_text else "F"
            for column, rating in sorted(rating_columns.items()):
                cell = region.cell(row, column)
                if cell is None:
                    continue
                category = next(
                    (value for phrase, value in categories if phrase in cell.normalized_value), None
                )
                source = region.source((cell,))
                rules.append(
                    CompiledRule(
                        rule_id=f"TPL-EXP-FREQ-{len(rules) + 1:04d}",
                        rule_type=RuleType.CRITERION,
                        role=TemplateRole.EXPOSURE_FREQUENCY_RULES,
                        predicates=(
                            Predicate(
                                FactType.OCCURRENCE_FREQUENCY,
                                PredicateOperator.EQ,
                                category,
                                source=source,
                            ),
                        ) if category else (),
                        result=rating,
                        alternatives=(),
                        priority=len(rules) + 1,
                        normative_strength=NormativeStrength.CRITERION,
                        source_refs=(source,),
                        raw_text=str(cell.raw_value),
                        parser="exposure-frequency-ordinal-v1",
                        parse_status=ParseStatus.COMPILED if category else ParseStatus.NEEDS_REVIEW,
                        diagnostics=() if category else ("ordinal frequency category unresolved",),
                        assessment_method=method,
                        executable=category is not None,
                    )
                )
        return rules

    def _compile_exposure_examples(self, region: _BoundRegion) -> list[CompiledRule]:
        examples: list[CompiledRule] = []
        active_ratings: dict[int, str] = {}
        in_examples = False
        for row in sorted({c.row for c in region.cells}):
            row_cells = region.row(row)
            row_text = " | ".join(c.normalized_value for c in row_cells)
            direct_ratings = {
                c.column: c.normalized_value.upper()
                for c in row_cells if re.fullmatch(r"E[1-4]", c.normalized_value.upper())
            }
            if len(direct_ratings) >= 3:
                active_ratings = direct_ratings
                in_examples = False
                continue
            label_cells = [c for c in row_cells if c.column < min(active_ratings, default=10**6)]
            if any("examples for" in c.normalized_value for c in label_cells):
                in_examples = True
            elif any(
                token in row_text for token in (
                    "ref_", "class of probability", "description", "duration (%",
                    "frequency of situation", "provide examples",
                )
            ):
                if not any("examples for" in c.normalized_value for c in label_cells):
                    in_examples = False
            if not in_examples or not active_ratings:
                continue
            for column, rating in active_ratings.items():
                cell = region.cell(row, column)
                if cell is None or cell.normalized_value in {"-", ""}:
                    continue
                source = region.source((cell,))
                examples.append(
                    CompiledRule(
                        rule_id=f"TPL-EXP-EX-{len(examples) + 1:04d}",
                        rule_type=RuleType.EXAMPLE,
                        role=TemplateRole.EXPOSURE_EXAMPLES,
                        predicates=(
                            CategoricalPredicate(
                                FactType.SITUATION_CLASSIFICATION,
                                (str(cell.raw_value).strip(),),
                                matching_semantics="REFERENCE_ONLY_NO_PATTERN_MATCH",
                                source=source,
                            ),
                        ),
                        result=rating,
                        alternatives=(),
                        priority=0,
                        normative_strength=NormativeStrength.EXAMPLE,
                        source_refs=(source,),
                        raw_text=str(cell.raw_value),
                        parser="exposure-example-v1",
                        parse_status=ParseStatus.NON_EXECUTABLE,
                        diagnostics=("example is not a deterministic override",),
                        assessment_method="EXAMPLE",
                        executable=False,
                    )
                )
        return examples

    def _compile_exposure_catalog(self, region: _BoundRegion) -> list[ExposureEntry]:
        entries: list[ExposureEntry] = []
        role = region.binding.role
        for row in sorted({c.row for c in region.cells}):
            cells = region.row(row)
            if not cells:
                continue
            ratings = [
                c for c in cells if _EXPOSURE_RE.fullmatch(c.normalized_value.upper())
            ]
            if not ratings:
                continue
            if role in {TemplateRole.VDA702_SUMMARY, TemplateRole.VDA702_FULL}:
                identifier = next(
                    (
                        c for c in cells
                        if re.fullmatch(r"[A-Za-zÄÖÜäöü]{1,5}\d{2,4}", str(c.raw_value).strip())
                    ),
                    None,
                )
                if identifier is None:
                    continue
                description = next(
                    (c for c in cells if c.column > identifier.column and c.column < ratings[0].column),
                    None,
                )
                if description is None:
                    continue
                time_rating = str(ratings[0].raw_value).upper()
                frequency_rating = str(ratings[1].raw_value).upper() if len(ratings) > 1 else ""
                metadata = {"catalog_id": str(identifier.raw_value).strip()}
                source_cells = cells
            else:
                description = cells[0]
                identifier = next(
                    (c for c in cells if c.column > description.column and c not in ratings), None
                )
                time_rating = str(ratings[0].raw_value).upper()
                frequency_rating = str(ratings[1].raw_value).upper() if len(ratings) > 1 else ""
                metadata = {"catalog_id": str(identifier.raw_value).strip() if identifier else ""}
                source_cells = tuple(c for c in cells if not c.is_formula)
            source = region.source(source_cells)
            entries.append(
                ExposureEntry(
                    entry_id=f"TPL-{role.value}-{len(entries) + 1:04d}",
                    source_role=role,
                    description=str(description.raw_value).strip(),
                    duration_rating=time_rating,
                    frequency_rating=frequency_rating,
                    normative_strength=NormativeStrength.REFERENCE,
                    source_ref=source,
                    metadata=metadata,
                )
            )
        return entries

    def _exposure_scale(
        self, duration_region: _BoundRegion, frequency_region: _BoundRegion
    ) -> list[ScaleLevel]:
        by_level: dict[str, list[CellSnapshot]] = {}
        for region in (duration_region, frequency_region):
            for cell in region.cells:
                level = cell.normalized_value.upper()
                if re.fullmatch(r"E[1-4]", level):
                    by_level.setdefault(level, []).append(cell)
        levels: list[ScaleLevel] = []
        for level in sorted(by_level):
            cells = by_level[level]
            source = duration_region.source((cells[0],))
            levels.append(ScaleLevel(level, "", "", "", source))
        return levels

    def _compile_controllability(
        self, snapshot: WorkbookSnapshot, role_contract: TemplateRoleContract
    ) -> tuple[ControllabilityContract, list[CompilerDiagnostic]]:
        diagnostics: list[CompilerDiagnostic] = []
        level_region = self._region(
            snapshot, role_contract, TemplateRole.CONTROLLABILITY_LEVELS
        )
        scale = self._compile_scale(level_region, "C")
        rule_region = self._region(
            snapshot, role_contract, TemplateRole.CONTROLLABILITY_RULES
        )
        header_cells = [
            c for c in rule_region.cells if re.fullmatch(r"C[0-3]", c.normalized_value.upper())
        ]
        header_row = min((c.row for c in header_cells), default=0)
        columns = {
            c.column: c.normalized_value.upper() for c in header_cells if c.row == header_row
        }
        criteria: list[CompiledRule] = []
        examples: list[CompiledRule] = []
        references: list[CompiledRule] = []
        assumptions: list[MethodAssumption] = []
        criterion_row = next(
            (
                row for row in range(header_row + 1, rule_region.bounds[3] + 1)
                if any("driving factors" in c.normalized_value for c in rule_region.row(row))
            ),
            0,
        )
        if criterion_row:
            for column, level in sorted(columns.items()):
                cell = rule_region.cell(criterion_row, column)
                if cell is None:
                    continue
                source = rule_region.source((cell,))
                predicate = self._parse_numeric_range(
                    str(cell.raw_value), FactType.AVOIDABILITY_PERCENT, "%", source
                )
                criteria.append(
                    CompiledRule(
                        rule_id=f"TPL-CTRL-CRIT-{len(criteria) + 1:04d}",
                        rule_type=RuleType.CRITERION,
                        role=TemplateRole.CONTROLLABILITY_RULES,
                        predicates=(predicate,) if predicate else (),
                        result=level,
                        alternatives=(),
                        priority=len(criteria) + 1,
                        normative_strength=NormativeStrength.CRITERION,
                        source_refs=(source,),
                        raw_text=str(cell.raw_value),
                        parser="controllability-avoidability-v1",
                        parse_status=(
                            ParseStatus.NEEDS_REVIEW
                            if predicate is None or self._range_needs_review(predicate)
                            else ParseStatus.COMPILED
                        ),
                        diagnostics=(
                            ("qualitative criterion preserved",)
                            if predicate is None else
                            ("range boundary inclusivity unresolved",)
                            if self._range_needs_review(predicate) else ()
                        ),
                        executable=predicate is not None,
                    )
                )
        for row in range(criterion_row + 1, rule_region.bounds[3] + 1):
            row_cells = rule_region.row(row)
            if not row_cells:
                continue
            label = row_cells[0]
            if label.normalized_value.startswith("example"):
                result_cell = next(
                    (
                        cell for cell in row_cells
                        if cell.column in columns and cell.normalized_value not in {"", "-"}
                    ),
                    None,
                )
                if result_cell is None:
                    continue
                source = rule_region.source(row_cells)
                examples.append(
                    CompiledRule(
                        rule_id=f"TPL-CTRL-EX-{len(examples) + 1:04d}",
                        rule_type=RuleType.EXAMPLE,
                        role=TemplateRole.CONTROLLABILITY_RULES,
                        predicates=(
                            CategoricalPredicate(
                                FactType.HAZARDOUS_EVENT,
                                (str(label.raw_value).strip(),),
                                matching_semantics="REFERENCE_ONLY_NO_KEYWORD_RULE",
                                source=source,
                            ),
                        ),
                        result=columns[result_cell.column],
                        alternatives=(),
                        priority=0,
                        normative_strength=NormativeStrength.EXAMPLE,
                        source_refs=(source,),
                        raw_text=" | ".join(str(c.raw_value) for c in row_cells),
                        parser="controllability-example-v1",
                        parse_status=ParseStatus.NON_EXECUTABLE,
                        diagnostics=("case-by-case informative example",),
                        executable=False,
                    )
                )
            elif label.normalized_value.startswith("note"):
                source = rule_region.source(row_cells)
                references.append(
                    CompiledRule(
                        rule_id=f"TPL-CTRL-REF-{len(references) + 1:04d}",
                        rule_type=RuleType.REFERENCE,
                        role=TemplateRole.CONTROLLABILITY_RULES,
                        predicates=(), result=None, alternatives=(), priority=0,
                        normative_strength=NormativeStrength.REFERENCE,
                        source_refs=(source,), raw_text=str(label.raw_value),
                        parser="controllability-reference-v1",
                        parse_status=ParseStatus.NON_EXECUTABLE,
                        executable=False,
                    )
                )
        workflow = self._region(snapshot, role_contract, TemplateRole.WORKFLOW)
        reaction = next(
            (c for c in workflow.cells if re.search(r"1[.,]3\s*s", c.normalized_value)), None
        )
        if reaction:
            source = workflow.source((reaction,))
            assumptions.append(
                MethodAssumption(
                    parameter="DRIVER_REACTION_TIME_REFERENCE",
                    value="1.3",
                    unit="s",
                    scope="controllability reasoning reference",
                    normative_strength=NormativeStrength.REFERENCE,
                    source_ref=source,
                )
            )
        diagnostics.extend(self._detect_range_coverage(criteria, FactType.AVOIDABILITY_PERCENT))
        ambiguous_criteria = [
            rule for rule in criteria
            if any(
                isinstance(item, RangePredicate) and self._range_needs_review(item)
                for item in rule.predicates
            )
        ]
        if ambiguous_criteria:
            diagnostics.append(
                CompilerDiagnostic(
                    CompilerDiagnosticSeverity.WARNING,
                    CompilerDiagnosticCode.SEMANTIC_VARIABLE_UNRESOLVED,
                    "Controllability percentage text preserves unresolved boundary inclusivity.",
                    role=TemplateRole.CONTROLLABILITY_RULES,
                    source_refs=tuple(
                        source for rule in ambiguous_criteria for source in rule.source_refs
                    ),
                )
            )
        contract = ControllabilityContract(
            tuple(scale), tuple(criteria), tuple(examples), tuple(references),
            tuple(assumptions), tuple(diagnostics)
        )
        return contract, diagnostics

    def _compile_asil(
        self, snapshot: WorkbookSnapshot, role_contract: TemplateRoleContract
    ) -> tuple[ASILMatrix, list[CompilerDiagnostic]]:
        diagnostics: list[CompilerDiagnostic] = []
        region = self._region(snapshot, role_contract, TemplateRole.ASIL_MATRIX)
        header_row = next(
            (
                row for row in sorted({c.row for c in region.cells})
                if len([
                    c for c in region.row(row)
                    if re.fullmatch(r"C\d+", c.normalized_value.upper())
                ]) >= 2
            ),
            0,
        )
        c_columns = {
            c.column: c.normalized_value.upper()
            for c in region.row(header_row)
            if re.fullmatch(r"C\d+", c.normalized_value.upper())
        }
        severity = ""
        mappings: list[ASILMapping] = []
        severity_levels: list[str] = []
        exposure_levels: list[str] = []
        for row in range(header_row + 1, region.bounds[3] + 1):
            cells = region.row(row)
            severity_cell = next(
                (c for c in cells if re.fullmatch(r"S\d+", c.normalized_value.upper())), None
            )
            if severity_cell:
                severity = severity_cell.normalized_value.upper()
                if severity not in severity_levels:
                    severity_levels.append(severity)
            exposure_cell = next(
                (c for c in cells if re.fullmatch(r"E\d+", c.normalized_value.upper())), None
            )
            if not severity or exposure_cell is None:
                continue
            exposure = exposure_cell.normalized_value.upper()
            if exposure not in exposure_levels:
                exposure_levels.append(exposure)
            for column, controllability in sorted(c_columns.items()):
                result_cell = region.cell(row, column)
                if result_cell is None:
                    continue
                raw_result = result_cell.normalized_value.upper().replace("ASIL ", "")
                if raw_result not in {"NA", "N/A", "QM", "A", "B", "C", "D"}:
                    continue
                source_cells = tuple(
                    c for c in (severity_cell, exposure_cell, result_cell) if c is not None
                )
                mappings.append(
                    ASILMapping(
                        severity=severity,
                        exposure=exposure,
                        controllability=controllability,
                        result=raw_result,
                        source_ref=region.source(source_cells),
                    )
                )
        expected = len(severity_levels) * len(exposure_levels) * len(c_columns)
        if len(mappings) != expected:
            diagnostics.append(
                CompilerDiagnostic(
                    CompilerDiagnosticSeverity.ERROR,
                    CompilerDiagnosticCode.ASIL_MATRIX_INCOMPLETE,
                    f"ASIL matrix has {len(mappings)} mappings; {expected} are required by its discovered axes.",
                    role=TemplateRole.ASIL_MATRIX,
                    source_refs=(region.source(region.cells),),
                    blocking=True,
                )
            )
        na_mappings = [item for item in mappings if item.result in {"NA", "N/A"}]
        if na_mappings:
            diagnostics.append(
                CompilerDiagnostic(
                    CompilerDiagnosticSeverity.WARNING,
                    CompilerDiagnosticCode.ASIL_NA_SEMANTICS_UNRESOLVED,
                    "ASIL matrix contains NA while the workbook also uses QM; no template statement proves they are equivalent, so NA is preserved.",
                    role=TemplateRole.ASIL_MATRIX,
                    source_refs=tuple(item.source_ref for item in na_mappings[:4]),
                )
            )
        matrix = ASILMatrix(
            severity_levels=tuple(severity_levels),
            exposure_levels=tuple(exposure_levels),
            controllability_levels=tuple(c_columns[column] for column in sorted(c_columns)),
            mappings=tuple(mappings),
            na_semantics="PRESERVE_NA_UNTIL_TEMPLATE_SEMANTICS_CONFIRMED",
            source_binding=region.binding,
            diagnostics=tuple(diagnostics),
        )
        return matrix, diagnostics

    _HARA_HEADER_ALIASES: tuple[tuple[str, tuple[str, ...]], ...] = (
        ("severity_rationale", ("rationale severity", "severity evaluation")),
        ("exposure_rationale", ("rationale exposure", "exposure evaluation")),
        ("controllability_rationale", ("rationale controllability", "controllability evaluation")),
        ("scenario_detail", ("situational detailing", "scenario detail")),
        ("scenario", ("situational description",)),
        ("hazardous_event", ("hazardous event", "potential damage")),
        ("hara_id", ("hara-id", "hara id")),
        ("malfunction", ("malfunction",)),
        ("function", ("function",)),
        ("output", ("output",)),
        ("guideword", ("guide-word", "guide word", "guideword")),
        ("hazard", ("hazard",)),
        ("severity", ("severity",)),
        ("exposure_method", ("t/f", "exposure method")),
        ("exposure", ("exposure",)),
        ("controllability", ("controllability",)),
        ("asil", ("asil",)),
        ("sg_id", ("sg-id", "sg id", "sz-id")),
        ("safety_goal", ("safety goal",)),
        ("safe_state", ("safe state", "safety state")),
        ("remark", ("remark",)),
    )

    _SG_HEADER_ALIASES: tuple[tuple[str, tuple[str, ...]], ...] = (
        ("hazard_id", ("hz-id", "hazard id")),
        ("hazard", ("hazard",)),
        ("sg_id", ("sg-id", "sg id", "sz-id")),
        ("safety_goal", ("safety goal",)),
        ("safe_state", ("safe state", "safety state")),
        ("max_asil", ("max. asil", "max asil", "asil")),
    )

    def _compile_report(
        self, snapshot: WorkbookSnapshot, role_contract: TemplateRoleContract
    ) -> tuple[ReportContract, list[CompilerDiagnostic]]:
        diagnostics: list[CompilerDiagnostic] = []
        hara_region = self._region(snapshot, role_contract, TemplateRole.HARA_OUTPUT_TABLE)
        sg_region = self._region(
            snapshot, role_contract, TemplateRole.SAFETY_GOAL_OUTPUT_TABLE
        )
        hara = self._compile_report_fields(hara_region, self._HARA_HEADER_ALIASES)
        safety = self._compile_report_fields(sg_region, self._SG_HEADER_ALIASES)
        for label, fields in (("HARA", hara), ("Safety Goal", safety)):
            names = [item.canonical_field for item in fields]
            duplicates = sorted({name for name in names if names.count(name) > 1})
            if duplicates:
                diagnostics.append(
                    CompilerDiagnostic(
                        CompilerDiagnosticSeverity.ERROR,
                        CompilerDiagnosticCode.REPORT_FIELD_MISMATCH,
                        f"{label} output maps canonical fields more than once: {duplicates}",
                        role=(
                            TemplateRole.HARA_OUTPUT_TABLE
                            if label == "HARA" else TemplateRole.SAFETY_GOAL_OUTPUT_TABLE
                        ),
                        blocking=True,
                    )
                )
        return (
            ReportContract(
                tuple(hara), tuple(safety),
                (hara_region.binding, sg_region.binding), tuple(diagnostics)
            ),
            diagnostics,
        )

    def _compile_report_fields(
        self,
        region: _BoundRegion,
        aliases: tuple[tuple[str, tuple[str, ...]], ...],
    ) -> list[ReportFieldMapping]:
        min_column, min_row, max_column, max_row = region.bounds
        fields: list[ReportFieldMapping] = []
        for column in range(min_column, max_column + 1):
            cells = tuple(
                cell for row in range(min_row, max_row + 1)
                for cell in (region.cell(row, column),) if cell is not None
            )
            if not cells:
                continue
            normalized = " ".join(c.normalized_value for c in cells)
            canonical = None
            if normalized.strip() in {"s", "e", "c"}:
                canonical = {"s": "severity", "e": "exposure", "c": "controllability"}[normalized.strip()]
            else:
                canonical = next(
                    (name for name, variants in aliases if any(alias in normalized for alias in variants)),
                    None,
                )
            if canonical is None:
                continue
            fields.append(
                ReportFieldMapping(
                    canonical_field=canonical,
                    sheet=region.binding.sheet,
                    column=get_column_letter(column),
                    column_index=column,
                    header_rows=tuple(range(min_row, max_row + 1)),
                    source_ref=region.source(cells),
                )
            )
        return fields

    def _compile_derivation_method(
        self,
        snapshot: WorkbookSnapshot,
        role_contract: TemplateRoleContract,
        report: ReportContract,
        role: TemplateRole,
    ) -> DerivationMethod:
        region = self._region(snapshot, role_contract, role)
        cells = region.cells
        raw = "\n".join(str(c.raw_value) for c in cells)
        source = region.source(cells)
        column_to_field = {
            item.column: item.canonical_field for item in report.hara_fields
        }
        referenced_columns = tuple(dict.fromkeys(re.findall(r"([A-Z])列", raw, re.IGNORECASE)))
        inputs = tuple(
            dict.fromkeys(
                column_to_field[column.upper()]
                for column in referenced_columns
                if column.upper() in column_to_field
            )
        )
        description = next(
            (str(c.raw_value) for c in cells if c.column == region.bounds[0] + 4), raw
        )
        instruction = CompiledRule(
            rule_id=("TPL-SG-METHOD-0001" if role is TemplateRole.SAFETY_GOAL_METHOD else "TPL-SS-METHOD-0001"),
            rule_type=RuleType.WORKFLOW_INSTRUCTION,
            role=role,
            predicates=(), result=None, alternatives=(), priority=1,
            normative_strength=NormativeStrength.INSTRUCTION,
            source_refs=(source,), raw_text=raw,
            parser="semantic-derivation-method-v1",
            parse_status=ParseStatus.COMPILED,
            diagnostics=("semantic derivation required; no fixed catalog generated",),
            executable=False,
        )
        qm_handling = next(
            (str(c.raw_value) for c in cells if "qm" in c.normalized_value), ""
        )
        return DerivationMethod(
            inputs_required=inputs,
            derivation_pattern=description,
            qm_handling=qm_handling,
            aggregation_instructions=(
                "Output table requests Max.ASIL; aggregation semantics require downstream contract validation."
                if role is TemplateRole.SAFETY_GOAL_METHOD else ""
            ),
            semantic_derivation_required=True,
            assumptions=(),
            instructions=(instruction,),
            source_binding=region.binding,
        )

    def _derive_required_facts(
        self,
        rules: tuple[CompiledRule, ...],
        workflow: WorkflowContract,
        exposure: ExposureContract,
    ) -> tuple[RequiredFactSpec, ...]:
        by_fact: dict[FactType, list[tuple[CompiledRule, object]]] = {}
        for rule in rules:
            if not rule.executable:
                continue
            for predicate in rule.predicates:
                by_fact.setdefault(predicate.field, []).append((rule, predicate))
        origins = {
            FactType.FUNCTION: FactOrigin.PROJECT_FACT,
            FactType.OUTPUT: FactOrigin.PROJECT_FACT,
            FactType.COLLISION_TYPE: FactOrigin.SCENARIO_FACT,
            FactType.ROAD_USER_TYPE: FactOrigin.SCENARIO_FACT,
            FactType.SPEED_UNSPECIFIED: FactOrigin.SCENARIO_FACT,
            FactType.DURATION_PERCENT: FactOrigin.DERIVED_FACT,
            FactType.OCCURRENCE_FREQUENCY: FactOrigin.HUMAN_EVIDENCE,
            FactType.AVOIDABILITY_PERCENT: FactOrigin.HUMAN_EVIDENCE,
        }
        specs: list[RequiredFactSpec] = []
        for fact, items in sorted(by_fact.items(), key=lambda item: item[0].value):
            source_refs = unique_sources(
                [source for rule, _predicate in items for source in rule.source_refs]
            )
            units = {
                getattr(predicate, "unit", "") for _rule, predicate in items
                if getattr(predicate, "unit", "")
            }
            constraints = tuple(dict.fromkeys(
                self._predicate_constraint(predicate) for _rule, predicate in items
            ))
            specs.append(
                RequiredFactSpec(
                    fact_type=fact,
                    required_for=tuple(dict.fromkeys(rule.rule_id for rule, _ in items)),
                    unit=next(iter(units), ""),
                    constraints=constraints,
                    condition="Required only when evaluating a rule that references this field.",
                    origin=origins.get(fact, FactOrigin.DERIVED_FACT),
                    source_rule_ids=tuple(dict.fromkeys(rule.rule_id for rule, _ in items)),
                    source_refs=source_refs,
                )
            )
        if exposure.duration_rules and exposure.frequency_rules:
            selector_sources = unique_sources([
                source
                for rule in (*exposure.duration_rules, *exposure.frequency_rules)
                for source in rule.source_refs
            ])
            selector_rule_ids = tuple(
                rule.rule_id
                for rule in (*exposure.duration_rules, *exposure.frequency_rules)
            )
            specs.append(RequiredFactSpec(
                fact_type=FactType.EXPOSURE,
                required_for=("EXPOSURE_METHOD_SELECTION",),
                unit="",
                constraints=("EXPOSURE IN ('T', 'F')",),
                condition=(
                    "Required when the compiled MethodContract contains both "
                    "duration and frequency exposure rule families."
                ),
                origin=FactOrigin.PROJECT_FACT,
                source_rule_ids=selector_rule_ids,
                source_refs=selector_sources,
            ))
        for fact, activity in (
            (FactType.FUNCTION, "FUNCTION_EXTRACTION"),
            (FactType.OUTPUT, "FUNCTION_EXTRACTION"),
        ):
            if fact in by_fact:
                continue
            steps = [step for step in workflow.steps if step.activity == activity]
            if not steps:
                continue
            specs.append(
                RequiredFactSpec(
                    fact_type=fact,
                    required_for=(f"WORKFLOW:{activity}",),
                    unit="", constraints=(),
                    condition="Required by the compiled workflow.",
                    origin=FactOrigin.PROJECT_FACT,
                    source_rule_ids=(),
                    source_refs=tuple(step.source_ref for step in steps),
                )
            )
        return tuple(specs)

    @staticmethod
    def _predicate_constraint(predicate: object) -> str:
        if isinstance(predicate, RangePredicate):
            return (
                f"{predicate.field.value}: lower={predicate.lower} ({predicate.lower_inclusive}), "
                f"upper={predicate.upper} ({predicate.upper_inclusive}) {predicate.unit}"
            )
        if isinstance(predicate, CategoricalPredicate):
            return f"{predicate.field.value} IN {predicate.values}"
        if isinstance(predicate, Predicate):
            return f"{predicate.field.value} {predicate.operator.value} {predicate.value}"
        return str(predicate)

    def _detect_rule_conflicts(
        self, rules: list[CompiledRule], role: TemplateRole
    ) -> list[CompilerDiagnostic]:
        by_signature: dict[str, list[CompiledRule]] = {}
        for rule in rules:
            signature = repr(
                (rule.assessment_method,)
                + tuple(self._predicate_constraint(item) for item in rule.predicates)
            )
            by_signature.setdefault(signature, []).append(rule)
        diagnostics: list[CompilerDiagnostic] = []
        for matches in by_signature.values():
            outcomes = {
                (item.result, item.alternatives) for item in matches
            }
            if len(matches) > 1 and len(outcomes) > 1:
                diagnostics.append(
                    CompilerDiagnostic(
                        CompilerDiagnosticSeverity.ERROR,
                        CompilerDiagnosticCode.RULE_CONFLICT,
                        f"Same normative conditions produce conflicting results in {role.value}.",
                        role=role,
                        source_refs=tuple(source for item in matches for source in item.source_refs),
                        details={"rule_ids": [item.rule_id for item in matches]},
                        blocking=True,
                    )
                )
        return diagnostics

    def _detect_range_coverage(
        self, rules: list[CompiledRule], field: FactType
    ) -> list[CompilerDiagnostic]:
        grouped: dict[tuple[str, ...], list[tuple[CompiledRule, RangePredicate]]] = {}
        for rule in rules:
            target = next(
                (item for item in rule.predicates if isinstance(item, RangePredicate) and item.field is field),
                None,
            )
            if target is None:
                continue
            categories = (rule.assessment_method,) + tuple(
                self._predicate_constraint(item)
                for item in rule.predicates if not isinstance(item, RangePredicate)
            )
            grouped.setdefault(categories, []).append((rule, target))
        diagnostics: list[CompilerDiagnostic] = []
        for items in grouped.values():
            ordered = sorted(items, key=lambda item: float("-inf") if item[1].lower is None else item[1].lower)
            for (left_rule, left), (right_rule, right) in zip(ordered, ordered[1:]):
                if left.upper is None or right.lower is None:
                    continue
                if left.upper < right.lower:
                    code = CompilerDiagnosticCode.RULE_RANGE_GAP
                    message = f"Range gap between {left_rule.rule_id} and {right_rule.rule_id}."
                elif left.upper > right.lower:
                    code = CompilerDiagnosticCode.RULE_RANGE_OVERLAP
                    message = f"Range overlap between {left_rule.rule_id} and {right_rule.rule_id}."
                else:
                    continue
                diagnostics.append(
                    CompilerDiagnostic(
                        CompilerDiagnosticSeverity.WARNING,
                        code, message, role=left_rule.role,
                        source_refs=left_rule.source_refs + right_rule.source_refs,
                    )
                )
        return diagnostics
