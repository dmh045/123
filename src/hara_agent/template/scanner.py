from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def normalize_template_text(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    text = text.replace("\u00a0", " ").replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip().casefold()


@dataclass(frozen=True)
class CellSnapshot:
    coordinate: str
    row: int
    column: int
    raw_value: object
    normalized_value: str
    is_formula: bool


@dataclass(frozen=True)
class HeaderCandidate:
    row: int
    coordinates: tuple[str, ...]
    normalized_values: tuple[str, ...]


@dataclass(frozen=True)
class TableSnapshot:
    name: str
    region: str
    headers: tuple[str, ...]


@dataclass(frozen=True)
class NamedRangeSnapshot:
    name: str
    destinations: tuple[tuple[str, str], ...]


@dataclass(frozen=True)
class SheetSnapshot:
    name: str
    max_row: int
    max_column: int
    used_region: str
    merged_cells: tuple[str, ...]
    cells: tuple[CellSnapshot, ...]
    formulas: tuple[str, ...]
    tables: tuple[TableSnapshot, ...]
    header_candidates: tuple[HeaderCandidate, ...]

    def cells_in_row(self, row: int) -> tuple[CellSnapshot, ...]:
        return tuple(cell for cell in self.cells if cell.row == row)

    def cell(self, row: int, column: int) -> CellSnapshot | None:
        return next(
            (cell for cell in self.cells if cell.row == row and cell.column == column), None
        )

    @property
    def normalized_text(self) -> str:
        return " | ".join(cell.normalized_value for cell in self.cells)


@dataclass(frozen=True)
class WorkbookSnapshot:
    source_path: Path
    source_hash: str
    sheets: tuple[SheetSnapshot, ...]
    named_ranges: tuple[NamedRangeSnapshot, ...]
    workbook_metadata: dict[str, str]

    def sheet(self, name: str) -> SheetSnapshot:
        for sheet in self.sheets:
            if sheet.name == name:
                return sheet
        raise KeyError(name)


class TemplateWorkbookScanner:
    """Create a structural workbook snapshot without assigning HARA roles."""

    def scan(self, template_path: str | Path) -> WorkbookSnapshot:
        path = Path(template_path).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(f"Template does not exist: {path}")
        source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            sheets = tuple(self._scan_sheet(sheet) for sheet in workbook.worksheets)
            named_ranges = []
            defined_names = workbook.defined_names
            # openpyxl 3.0 exposes a DefinedNameList, while 3.1 exposes a
            # mapping-like DefinedNameDict.  requirements.txt intentionally
            # supports both versions.
            if hasattr(defined_names, "items"):
                named_name_items = defined_names.items()
            else:
                named_name_items = (
                    (str(item.name), item)
                    for item in getattr(defined_names, "definedName", ())
                )
            for name, defined_name in named_name_items:
                try:
                    destinations = tuple(
                        (str(sheet), str(region)) for sheet, region in defined_name.destinations
                    )
                except (AttributeError, TypeError, ValueError):
                    destinations = ()
                named_ranges.append(NamedRangeSnapshot(str(name), destinations))
            properties = workbook.properties
            metadata = {
                key: str(value)
                for key, value in {
                    "title": properties.title,
                    "subject": properties.subject,
                    "creator": properties.creator,
                    "description": properties.description,
                    "category": properties.category,
                }.items()
                if value not in (None, "")
            }
            return WorkbookSnapshot(
                source_path=path,
                source_hash=source_hash,
                sheets=sheets,
                named_ranges=tuple(named_ranges),
                workbook_metadata=metadata,
            )
        finally:
            workbook.close()

    def _scan_sheet(self, sheet: Any) -> SheetSnapshot:
        cells: list[CellSnapshot] = []
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value in (None, ""):
                    continue
                cells.append(
                    CellSnapshot(
                        coordinate=cell.coordinate,
                        row=cell.row,
                        column=cell.column,
                        raw_value=cell.value,
                        normalized_value=normalize_template_text(cell.value),
                        is_formula=isinstance(cell.value, str) and cell.value.startswith("="),
                    )
                )
        if cells:
            min_row = min(cell.row for cell in cells)
            max_row = max(cell.row for cell in cells)
            min_column = min(cell.column for cell in cells)
            max_column = max(cell.column for cell in cells)
            used_region = (
                f"{get_column_letter(min_column)}{min_row}:"
                f"{get_column_letter(max_column)}{max_row}"
            )
        else:
            used_region = "A1:A1"

        tables = []
        for table in sheet.tables.values():
            headers: tuple[str, ...] = ()
            try:
                headers = tuple(
                    normalize_template_text(item.name) for item in table.tableColumns
                )
            except (AttributeError, TypeError):
                pass
            tables.append(TableSnapshot(str(table.name), str(table.ref), headers))

        headers = []
        rows = sorted({cell.row for cell in cells})
        for row in rows:
            row_cells = tuple(cell for cell in cells if cell.row == row)
            if len(row_cells) < 2:
                continue
            textual = tuple(
                cell for cell in row_cells if len(cell.normalized_value) <= 160
            )
            if len(textual) < 2:
                continue
            headers.append(
                HeaderCandidate(
                    row=row,
                    coordinates=tuple(cell.coordinate for cell in textual),
                    normalized_values=tuple(cell.normalized_value for cell in textual),
                )
            )
        return SheetSnapshot(
            name=str(sheet.title),
            max_row=int(sheet.max_row),
            max_column=int(sheet.max_column),
            used_region=used_region,
            merged_cells=tuple(str(item) for item in sheet.merged_cells.ranges),
            cells=tuple(cells),
            formulas=tuple(cell.coordinate for cell in cells if cell.is_formula),
            tables=tuple(tables),
            header_candidates=tuple(headers),
        )
