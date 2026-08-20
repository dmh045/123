import json
import sys
import unittest
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "scripts"))

from logic_checkers.scoring_checker import ScoringChecker


class P009SafetyGoalAggregationTests(unittest.TestCase):
    scoring = ROOT / "output" / "p009_scoring_sg_output.json"
    report = ROOT / "output" / "HARA_Report_AVP_P009_SG_Aggregated.xlsx"
    temp_root = ROOT / "tmp"

    @classmethod
    def setUpClass(cls):
        missing = [path for path in (cls.scoring, cls.report) if not path.is_file()]
        if missing:
            raise unittest.SkipTest(
                "legacy generated fixtures are absent: " + ", ".join(str(path) for path in missing)
            )

    @staticmethod
    def _leaf_items(value):
        if isinstance(value, list):
            for item in value:
                yield from P009SafetyGoalAggregationTests._leaf_items(item)
        elif isinstance(value, dict):
            if "scenario_id" in value:
                yield value
            else:
                for item in value.values():
                    yield from P009SafetyGoalAggregationTests._leaf_items(item)

    def test_avp_has_seven_stable_top_level_goals(self):
        data = json.loads(self.scoring.read_text(encoding="utf-8"))
        catalog = data["safety_goal_catalog"]
        self.assertEqual(
            {f"SG_AVP_{index:02d}" for index in range(1, 8)}, set(catalog)
        )
        self.assertEqual(7, len({item["safety_goal"] for item in catalog.values()}))
        self.assertTrue(all(item["max_asil"] == "B" for item in catalog.values()))

    def test_event_rows_reference_catalog_without_changing_risk_distribution(self):
        data = json.loads(self.scoring.read_text(encoding="utf-8"))
        catalog = data["safety_goal_catalog"]
        items = list(self._leaf_items(data["safety_goals"]))
        self.assertEqual(70, len(items))
        self.assertEqual(Counter({"B": 62, "QM": 8}), Counter(x["ASIL"] for x in items))
        non_qm = [item for item in items if item["ASIL"] != "QM"]
        qm = [item for item in items if item["ASIL"] == "QM"]
        self.assertTrue(all(item["sg_id"] in catalog for item in non_qm))
        self.assertTrue(all(item["safety_goal"] == "NA" and not item.get("sg_id") for item in qm))

    def test_excel_reuses_same_ids_and_has_one_row_per_top_level_goal(self):
        wb = load_workbook(self.report, data_only=True)
        hara = wb["05_HARA"]
        sg_sheet = wb["06_Safety Goal"]
        hara_ids = {
            str(hara.cell(row, 18).value)
            for row in range(6, hara.max_row + 1)
            if str(hara.cell(row, 17).value) in {"A", "B", "C", "D"}
        }
        sg_rows = [
            [sg_sheet.cell(row, col).value for col in range(1, 7)]
            for row in range(1, sg_sheet.max_row + 1)
            if str(sg_sheet.cell(row, 3).value).startswith("SG_AVP_")
        ]
        self.assertEqual(7, len(sg_rows))
        self.assertEqual(hara_ids, {str(row[2]) for row in sg_rows})
        self.assertTrue(all(str(row[0]).startswith("HZ_AVP_") for row in sg_rows))

    def test_checker_rejects_incorrect_max_asil(self):
        self.temp_root.mkdir(exist_ok=True)
        invalid = self.temp_root / "p009_invalid_max_asil.json"
        try:
            data = json.loads(self.scoring.read_text(encoding="utf-8"))
            data["safety_goal_catalog"]["SG_AVP_01"]["max_asil"] = "D"
            invalid.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            checker = ScoringChecker(str(invalid))
            self.assertTrue(checker.load())
            result = checker.check_all()
            self.assertEqual("fail", result["status"])
            self.assertTrue(any(
                error.get("type") == "safety_goal_max_asil_invalid"
                for error in result["errors"]
            ))
        finally:
            invalid.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
