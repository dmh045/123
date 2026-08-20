import shutil
import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "scripts"))

from asil_matrix import TemplateASILMatrix
from scoring_sg_engine import ASILDeterminer
from logic_checkers.scoring_checker import ScoringChecker


TEMPLATE = ROOT / "references" / "HARA_Template_AI_20260327.xlsx"


class P017ASILMatrixTests(unittest.TestCase):
    def test_reads_complete_matrix_and_normalizes_na_to_qm(self):
        determiner = ASILDeterminer(str(TEMPLATE))

        self.assertEqual(80, len(determiner.matrix))
        self.assertEqual("B", determiner.determine("S1", "E4", "C3"))
        self.assertEqual("D", determiner.determine("S3", "E4", "C3"))
        self.assertEqual("QM", determiner.determine("S0", "E4", "C3"))
        self.assertTrue(determiner.source.endswith("#ASIL_Table"))

    def test_result_follows_modified_template_cell_not_numeric_sum(self):
        changed_template = ROOT / "tmp" / "p017_changed_matrix.xlsx"
        try:
            shutil.copy2(TEMPLATE, changed_template)
            workbook = load_workbook(changed_template)
            # H14 is the template's S1/E4/C3 result (normally B). Deliberately
            # make it differ from the historical sum rule to prove lookup.
            workbook["ASIL_Table"]["H14"] = "D"
            workbook.save(changed_template)
            workbook.close()

            determiner = ASILDeterminer(str(changed_template))
            self.assertEqual("D", determiner.determine("S1", "E4", "C3"))
        finally:
            changed_template.unlink(missing_ok=True)

    def test_missing_asil_table_is_a_hard_error(self):
        changed_template = ROOT / "tmp" / "p017_missing_matrix.xlsx"
        try:
            shutil.copy2(TEMPLATE, changed_template)
            workbook = load_workbook(changed_template)
            del workbook["ASIL_Table"]
            workbook.save(changed_template)
            workbook.close()

            with self.assertRaisesRegex(ValueError, "missing required sheet"):
                TemplateASILMatrix(str(changed_template))
        finally:
            changed_template.unlink(missing_ok=True)

    def test_invalid_sec_value_is_a_hard_error(self):
        determiner = ASILDeterminer(str(TEMPLATE))
        with self.assertRaisesRegex(ValueError, "Invalid S/E/C combination"):
            determiner.determine("S4", "E4", "C3")

    def test_scoring_checker_rejects_result_that_disagrees_with_matrix(self):
        checker = ScoringChecker("unused.json", str(TEMPLATE))
        checker.data = {
            "asil_results": {
                "测试功能": {
                    "测试失效||loss": {
                        "loss": [{
                            "scenario_id": "SCN_TEST",
                            "severity": "S0",
                            "exposure": "E4",
                            "controllability": "C3",
                            "ASIL": "A",
                        }]
                    }
                }
            }
        }

        checker._check_asil_results()
        self.assertEqual("asil_matrix_mismatch", checker.errors[0]["type"])
        self.assertEqual("QM", checker.errors[0]["expected"])


if __name__ == "__main__":
    unittest.main()
